"""
Company HR System - runs on one office PC, used by everyone on the LAN via browser.
Features: employees, departments, attendance, leave management, payroll & payslips.
"""
import os
import re
import time
import threading
import csv
import io
import socket
import secrets
import sqlite3
import calendar
from collections import Counter
from datetime import date, datetime, timedelta
from functools import wraps

from flask import (Flask, g, render_template, request, redirect, url_for, session,
                   flash, abort, send_file, Response)
from werkzeug.security import generate_password_hash, check_password_hash

APP_VERSION = "2.2.0"
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
# Where the data lives. Desktop app sets HR_DATA (e.g. Documents\NB Sons HR Data); otherwise next to app.py.
DATA_DIR = os.environ.get("HR_DATA") or BASE_DIR
os.makedirs(DATA_DIR, exist_ok=True)
DB_PATH = os.environ.get("HR_DB", os.path.join(DATA_DIR, "hr.db"))
KEY_PATH = os.path.join(DATA_DIR, "secret.key")
BACKUP_DIR = os.path.join(DATA_DIR, "backups")


def _load_secret():
    if os.path.exists(KEY_PATH):
        with open(KEY_PATH) as f:
            return f.read().strip()
    key = secrets.token_hex(32)
    with open(KEY_PATH, "w") as f:
        f.write(key)
    return key


app = Flask(__name__)
app.secret_key = _load_secret()
app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(hours=10)
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
MAX_FAILS, LOCK_MINUTES, AUTO_BACKUPS_KEEP = 5, 15, 30
app.config["MAX_CONTENT_LENGTH"] = 200 * 1024 * 1024

ATT_STATUSES = ["present", "late", "half_day", "absent", "leave"]
ATT_CODES = {"present": "P", "late": "L", "half_day": "H", "absent": "A", "leave": "LV"}
WEEKDAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]

# --------------------------------------------------------------------------- database
SCHEMA = """
CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY, value TEXT);
CREATE TABLE IF NOT EXISTS departments(id INTEGER PRIMARY KEY, name TEXT UNIQUE NOT NULL COLLATE NOCASE);
CREATE TABLE IF NOT EXISTS employees(
  id INTEGER PRIMARY KEY, emp_code TEXT UNIQUE NOT NULL, name TEXT NOT NULL, father_name TEXT,
  cnic TEXT, gender TEXT, dob TEXT, email TEXT, phone TEXT, address TEXT,
  department_id INTEGER REFERENCES departments(id) ON DELETE SET NULL,
  designation TEXT, join_date TEXT, employment_type TEXT,
  basic_salary REAL DEFAULT 0, allowances REAL DEFAULT 0,
  bank_account TEXT, emergency_contact TEXT, status TEXT DEFAULT 'active',
  created_at TEXT DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE IF NOT EXISTS users(
  id INTEGER PRIMARY KEY, username TEXT UNIQUE NOT NULL COLLATE NOCASE, password_hash TEXT NOT NULL,
  role TEXT NOT NULL,
  employee_id INTEGER UNIQUE REFERENCES employees(id) ON DELETE CASCADE,
  active INTEGER DEFAULT 1, must_change INTEGER DEFAULT 0, last_login TEXT);
CREATE TABLE IF NOT EXISTS attendance(
  id INTEGER PRIMARY KEY, employee_id INTEGER NOT NULL REFERENCES employees(id) ON DELETE CASCADE,
  date TEXT NOT NULL, status TEXT NOT NULL, check_in TEXT, check_out TEXT, note TEXT,
  UNIQUE(employee_id, date));
CREATE TABLE IF NOT EXISTS recycle_bin(
  id INTEGER PRIMARY KEY, kind TEXT NOT NULL, label TEXT, detail TEXT, deleted_at TEXT, deleted_by TEXT, data TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS zkt_punches(
  badge TEXT NOT NULL, punch_time TEXT NOT NULL, PRIMARY KEY(badge, punch_time));
CREATE TABLE IF NOT EXISTS leave_types(
  id INTEGER PRIMARY KEY, name TEXT UNIQUE NOT NULL COLLATE NOCASE,
  days_per_year INTEGER DEFAULT 0, paid INTEGER DEFAULT 1, encash INTEGER DEFAULT 0);
CREATE TABLE IF NOT EXISTS leaves(
  id INTEGER PRIMARY KEY, employee_id INTEGER NOT NULL REFERENCES employees(id) ON DELETE CASCADE,
  leave_type_id INTEGER NOT NULL REFERENCES leave_types(id),
  start_date TEXT NOT NULL, end_date TEXT NOT NULL, days REAL NOT NULL, reason TEXT,
  status TEXT DEFAULT 'pending', applied_on TEXT, decided_by TEXT, decided_on TEXT, remarks TEXT);
CREATE TABLE IF NOT EXISTS payroll(
  id INTEGER PRIMARY KEY, employee_id INTEGER NOT NULL REFERENCES employees(id) ON DELETE CASCADE,
  month TEXT NOT NULL, basic REAL, allowances REAL, overtime REAL DEFAULT 0, bonus REAL DEFAULT 0,
  working_days INTEGER, present_days REAL, absent_days REAL, late_days INTEGER,
  unpaid_leave_days REAL, deduction_days REAL, absent_deduction REAL DEFAULT 0,
  tax REAL DEFAULT 0, other_deductions REAL DEFAULT 0, net REAL, remarks TEXT,
  status TEXT DEFAULT 'draft', generated_on TEXT, paid_on TEXT,
  UNIQUE(employee_id, month));
CREATE TABLE IF NOT EXISTS roles(
  key TEXT PRIMARY KEY, name TEXT NOT NULL, description TEXT, perms TEXT DEFAULT '', system INTEGER DEFAULT 0);
CREATE TABLE IF NOT EXISTS assets(
  id INTEGER PRIMARY KEY, asset_code TEXT UNIQUE NOT NULL, name TEXT NOT NULL, category TEXT, brand TEXT,
  model TEXT, serial_no TEXT, purchase_date TEXT, cost REAL DEFAULT 0, condition TEXT DEFAULT 'Good',
  status TEXT DEFAULT 'available', location TEXT, notes TEXT, created_at TEXT DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE IF NOT EXISTS asset_assignments(
  id INTEGER PRIMARY KEY, asset_id INTEGER NOT NULL REFERENCES assets(id) ON DELETE CASCADE,
  employee_id INTEGER REFERENCES employees(id) ON DELETE SET NULL, employee_name TEXT,
  assigned_on TEXT, expected_return TEXT, returned_on TEXT, condition_out TEXT, condition_in TEXT,
  notes TEXT, return_notes TEXT, assigned_by TEXT, received_by TEXT);
CREATE TABLE IF NOT EXISTS loans(
  id INTEGER PRIMARY KEY, employee_id INTEGER NOT NULL REFERENCES employees(id) ON DELETE CASCADE,
  amount REAL NOT NULL, installment REAL NOT NULL, issued_on TEXT, start_month TEXT, reason TEXT,
  status TEXT DEFAULT 'active', notes TEXT, created_by TEXT, closed_on TEXT,
  created_at TEXT DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE IF NOT EXISTS loan_payments(
  id INTEGER PRIMARY KEY, loan_id INTEGER NOT NULL REFERENCES loans(id) ON DELETE CASCADE,
  payroll_id INTEGER REFERENCES payroll(id) ON DELETE CASCADE, month TEXT, amount REAL NOT NULL,
  paid_on TEXT, method TEXT DEFAULT 'salary', note TEXT, created_by TEXT);
CREATE INDEX IF NOT EXISTS idx_loan_emp ON loans(employee_id);
CREATE INDEX IF NOT EXISTS idx_lpay_loan ON loan_payments(loan_id);
CREATE TABLE IF NOT EXISTS audit_log(
  id INTEGER PRIMARY KEY, at TEXT NOT NULL, username TEXT, ip TEXT, action TEXT NOT NULL, detail TEXT);
CREATE INDEX IF NOT EXISTS idx_audit_at ON audit_log(at);
CREATE TABLE IF NOT EXISTS holidays(
  id INTEGER PRIMARY KEY, date TEXT UNIQUE NOT NULL, name TEXT NOT NULL, kind TEXT DEFAULT 'public');
CREATE TABLE IF NOT EXISTS increments(
  id INTEGER PRIMARY KEY, employee_id INTEGER NOT NULL REFERENCES employees(id) ON DELETE CASCADE,
  eff_date TEXT NOT NULL, kind TEXT DEFAULT 'increment', old_basic REAL, old_allow REAL, new_basic REAL, new_allow REAL,
  old_desig TEXT, new_desig TEXT, reason TEXT, created_by TEXT, created_at TEXT);
CREATE TABLE IF NOT EXISTS settlements(
  id INTEGER PRIMARY KEY, employee_id INTEGER NOT NULL REFERENCES employees(id) ON DELETE CASCADE,
  exit_date TEXT, exit_type TEXT, reason TEXT, salary_days REAL DEFAULT 0, salary_amount REAL DEFAULT 0,
  encash_days REAL DEFAULT 0, encash_amount REAL DEFAULT 0, loan_recovery REAL DEFAULT 0, asset_deduction REAL DEFAULT 0,
  other_add REAL DEFAULT 0, other_add_note TEXT, other_ded REAL DEFAULT 0, other_ded_note TEXT, net REAL DEFAULT 0,
  assets_json TEXT, remarks TEXT, status TEXT DEFAULT 'draft', created_by TEXT, created_at TEXT, finalized_on TEXT);
CREATE TABLE IF NOT EXISTS emp_documents(
  id INTEGER PRIMARY KEY, employee_id INTEGER NOT NULL REFERENCES employees(id) ON DELETE CASCADE,
  title TEXT NOT NULL, doc_type TEXT, filename TEXT NOT NULL, orig_name TEXT, size INTEGER DEFAULT 0,
  expiry TEXT, notes TEXT, uploaded_by TEXT, uploaded_at TEXT);
CREATE TABLE IF NOT EXISTS warnings(
  id INTEGER PRIMARY KEY, employee_id INTEGER NOT NULL REFERENCES employees(id) ON DELETE CASCADE,
  issued_on TEXT NOT NULL, kind TEXT DEFAULT 'written', subject TEXT NOT NULL, details TEXT, action TEXT,
  reply_by TEXT, status TEXT DEFAULT 'open', response TEXT, issued_by TEXT, created_at TEXT);
CREATE INDEX IF NOT EXISTS idx_inc_emp ON increments(employee_id);
CREATE INDEX IF NOT EXISTS idx_doc_emp ON emp_documents(employee_id);
CREATE INDEX IF NOT EXISTS idx_warn_emp ON warnings(employee_id);
CREATE INDEX IF NOT EXISTS idx_att_date ON attendance(date);
CREATE INDEX IF NOT EXISTS idx_asg_asset ON asset_assignments(asset_id);
CREATE INDEX IF NOT EXISTS idx_asg_emp ON asset_assignments(employee_id);
CREATE INDEX IF NOT EXISTS idx_leave_emp ON leaves(employee_id);
CREATE INDEX IF NOT EXISTS idx_pay_month ON payroll(month);
"""

DEFAULT_SETTINGS = {
    "zkt_on": "0", "zkt_path": "", "zkt_from": "", "zkt_last": "", "zkt_err": "", "zkt_mtime": "",
    "company_name": "My Company",
    "company_address": "",
    "currency": "PKR",
    "office_start": "09:00",
    "office_end": "17:00",
    "grace_minutes": "15",
    "weekly_off": "6",          # comma list of weekday numbers, Monday=0 ... Sunday=6
    "encash_basis": "gross26",  # unused leave payable per day: salary / 26 (company rule); also basic26 / basic30 / gross30
    "lates_per_day": "0",       # e.g. 3 = every 3 lates deduct 1 day salary; 0 = off
    "idle_minutes": "30",       # auto logout after this many minutes without activity
    "notify_off": "",           # comma list of alert types switched off
    "notify_sound": "1",        # play a sound with alerts
    "notify_loud": "1",         # different, louder sound for important alerts
    "click_sound": "1",         # sound on every click (buttons, menu, tabs)
    "click_volume": "high",     # low / medium / high
}


PERMISSIONS = [   # (key, module, what it allows)
    ("employees.view", "Employees", "View employees & profiles"),
    ("employees.manage", "Employees", "Add, edit, import employees & logins"),
    ("employees.delete", "Employees", "Delete employees permanently"),
    ("departments.view", "Departments", "View departments"),
    ("departments.manage", "Departments", "Add, edit, delete departments"),
    ("attendance.view", "Attendance", "View attendance reports"),
    ("attendance.manage", "Attendance", "Mark, edit & import attendance"),
    ("leaves.view", "Leaves", "View all leave requests"),
    ("leaves.approve", "Leaves", "Approve / reject / record leave"),
    ("payroll.view", "Payroll", "View payroll & payslips"),
    ("payroll.manage", "Payroll", "Generate, edit & pay salaries"),
    ("assets.view", "Assets", "View asset register"),
    ("assets.manage", "Assets", "Add assets, assign & return"),
    ("settings", "System", "Company settings & leave types"),
    ("users", "System", "Users, roles & permissions"),
    ("backup", "System", "Download database backup"),
]
ALL_PERMS = [p[0] for p in PERMISSIONS]
DEFAULT_ROLES = [
    ("admin", "Administrator", "Full access to everything", ALL_PERMS),
    ("hr", "HR Manager", "Runs HR day to day", [p for p in ALL_PERMS if p not in ("settings", "users", "backup", "employees.delete")]),
    ("manager", "Manager (view only)", "Can view employees, attendance, leaves and assets – cannot change anything",
     ["employees.view", "departments.view", "attendance.view", "leaves.view", "assets.view"]),
    ("employee", "No access", "Cannot sign in – this system is for HR and managers only", []),
]


def init_db():
    con = sqlite3.connect(DB_PATH)
    con.execute("PRAGMA journal_mode=WAL")
    con.executescript(SCHEMA)
    dcols = {r[1] for r in con.execute("PRAGMA table_info(departments)")}
    for col, typ in (("code", "TEXT"), ("head_id", "INTEGER"), ("description", "TEXT")):
        if col not in dcols:
            con.execute(f"ALTER TABLE departments ADD COLUMN {col} {typ}")
    cols = {r[1] for r in con.execute("PRAGMA table_info(employees)")}
    if "source" not in cols:
        con.execute("ALTER TABLE employees ADD COLUMN source TEXT DEFAULT 'manual'")
        con.execute("UPDATE employees SET source='manual' WHERE source IS NULL")
    pcols = {r[1] for r in con.execute("PRAGMA table_info(payroll)")}
    if "loan_deduction" not in pcols:
        con.execute("ALTER TABLE payroll ADD COLUMN loan_deduction REAL DEFAULT 0")
    if "suggested_deduction" not in pcols:          # deductions are suggested, HR applies them by hand
        con.execute("ALTER TABLE payroll ADD COLUMN suggested_deduction REAL DEFAULT 0")
        con.execute("ALTER TABLE payroll ADD COLUMN ded_status TEXT")
    if "manager_id" not in cols:
        con.execute("ALTER TABLE employees ADD COLUMN manager_id INTEGER")
    usql = con.execute("SELECT sql FROM sqlite_master WHERE type='table' AND name='users'").fetchone()[0]
    if "CHECK(role" in usql:                       # allow custom roles (old databases had a fixed list)
        con.execute("PRAGMA foreign_keys=OFF")
        con.executescript("""
          CREATE TABLE users_new(
            id INTEGER PRIMARY KEY, username TEXT UNIQUE NOT NULL COLLATE NOCASE, password_hash TEXT NOT NULL,
            role TEXT NOT NULL, employee_id INTEGER UNIQUE REFERENCES employees(id) ON DELETE CASCADE,
            active INTEGER DEFAULT 1, must_change INTEGER DEFAULT 0, last_login TEXT);
          INSERT INTO users_new SELECT id, username, password_hash, role, employee_id, active, must_change, last_login FROM users;
          DROP TABLE users;
          ALTER TABLE users_new RENAME TO users;""")
        con.execute("PRAGMA foreign_keys=ON")
    ucols = {r[1] for r in con.execute("PRAGMA table_info(users)")}
    for col, typ in (("failed_count", "INTEGER DEFAULT 0"), ("locked_until", "TEXT"),
                     ("perm_add", "TEXT DEFAULT ''"), ("perm_remove", "TEXT DEFAULT ''")):
        if col not in ucols:
            con.execute(f"ALTER TABLE users ADD COLUMN {col} {typ}")
    for key, name, desc, perms in DEFAULT_ROLES:
        con.execute("INSERT OR IGNORE INTO roles(key,name,description,perms,system) VALUES(?,?,?,?,1)",
                    (key, name, desc, ",".join(perms)))
    con.execute("""UPDATE roles SET name='No access', description='Cannot sign in – this system is for HR and managers only'
                   WHERE key='employee' AND name='Employee'""")
    if not con.execute("SELECT 1 FROM settings WHERE key='encash_rule_26'").fetchone():
        con.execute("UPDATE settings SET value='gross26' WHERE key='encash_basis' AND value='basic30'")
        con.execute("INSERT OR IGNORE INTO settings(key,value) VALUES('encash_rule_26','1')")
    for k, v in DEFAULT_SETTINGS.items():
        con.execute("INSERT OR IGNORE INTO settings(key,value) VALUES(?,?)", (k, v))
    ltcols = {r[1] for r in con.execute("PRAGMA table_info(leave_types)")}
    if "encash" not in ltcols:                     # leave encashment: Annual 12 + Sick 8 are paid out if unused
        con.execute("ALTER TABLE leave_types ADD COLUMN encash INTEGER DEFAULT 0")
        con.execute("UPDATE leave_types SET encash=1 WHERE name IN ('Annual','Sick')")
        con.execute("UPDATE leave_types SET days_per_year=12 WHERE name='Annual' AND days_per_year=14")
    if not con.execute("SELECT 1 FROM leave_types").fetchone():
        con.executemany("INSERT INTO leave_types(name,days_per_year,paid,encash) VALUES(?,?,?,?)",
                        [("Annual", 12, 1, 1), ("Casual", 10, 1, 0), ("Sick", 8, 1, 1), ("Unpaid", 0, 0, 0)])
    if not con.execute("SELECT 1 FROM departments").fetchone():
        con.executemany("INSERT INTO departments(name) VALUES(?)",
                        [("Administration",), ("Human Resources",), ("Accounts",)])
    if not con.execute("SELECT 1 FROM users WHERE role='admin'").fetchone():
        con.execute("INSERT INTO users(username,password_hash,role,must_change) VALUES(?,?,?,1)",
                    ("admin", generate_password_hash("admin123"), "admin"))
    con.commit()
    con.close()


def db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH, timeout=15)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys=ON")
        g.db.execute("PRAGMA secure_delete=ON")        # deleted records are wiped, not just hidden
    return g.db


@app.teardown_appcontext
def _close_db(exc):
    con = g.pop("db", None)
    if con is not None:
        con.close()


def q(sql, args=(), one=False):
    rows = db().execute(sql, args).fetchall()
    return (rows[0] if rows else None) if one else rows


def x(sql, args=()):
    cur = db().execute(sql, args)
    db().commit()
    return cur.lastrowid


# --------------------------------------------------------------------------- helpers
def get_settings():
    if "settings" not in g:
        g.settings = {r["key"]: r["value"] for r in q("SELECT key,value FROM settings")}
    return g.settings


def weekly_off():
    raw = get_settings().get("weekly_off", "")
    return {int(p) for p in raw.split(",") if p.strip().isdigit()}


def holiday_map():
    """{'YYYY-MM-DD': 'Holiday name'} for all public holidays (paid days off, never counted absent)."""
    from flask import has_app_context
    if has_app_context() and "holidays" in g:
        return g.holidays
    try:
        if has_app_context():
            m = {r["date"]: r["name"] for r in q("SELECT date, name FROM holidays")}
        else:
            con = sqlite3.connect(DB_PATH, timeout=15)
            m = dict(con.execute("SELECT date, name FROM holidays").fetchall())
            con.close()
    except Exception:                                  # noqa: BLE001
        m = {}
    if has_app_context():
        g.holidays = m
    return m


def holiday_name(d):
    return holiday_map().get(d.isoformat() if hasattr(d, "isoformat") else str(d))


def is_workday(d):
    return d.weekday() not in weekly_off() and not holiday_name(d)


def parse_date(s, default=None):
    try:
        return datetime.strptime((s or "").strip(), "%Y-%m-%d").date()
    except ValueError:
        return default


def parse_month(s):
    try:
        dt = datetime.strptime((s or "").strip(), "%Y-%m")
        return dt.year, dt.month
    except ValueError:
        t = date.today()
        return t.year, t.month


def month_dates(y, m):
    return [date(y, m, d) for d in range(1, calendar.monthrange(y, m)[1] + 1)]


def daterange(a, b):
    d = a
    while d <= b:
        yield d
        d += timedelta(days=1)


def to_float(v, default=0.0):
    try:
        return float(str(v).replace(",", "").strip() or default)
    except ValueError:
        return default


def parse_time(s):
    try:
        return datetime.strptime((s or "").strip(), "%H:%M").time()
    except ValueError:
        return None


def shift_month(y, m, delta):
    m += delta
    while m < 1:
        m += 12
        y -= 1
    while m > 12:
        m -= 12
        y += 1
    return f"{y:04d}-{m:02d}"


def next_emp_code():
    rows = q("SELECT emp_code FROM employees WHERE emp_code LIKE 'EMP-%'")
    nums = [int(r["emp_code"][4:]) for r in rows if r["emp_code"][4:].isdigit()]
    return f"EMP-{(max(nums) + 1 if nums else 1):04d}"


def leave_balances(emp_id, year=None):
    year = str(year or date.today().year)
    out = []
    for t in q("SELECT * FROM leave_types ORDER BY paid DESC, name"):
        used = q("""SELECT COALESCE(SUM(days),0) s FROM leaves WHERE employee_id=? AND leave_type_id=?
                    AND status='approved' AND substr(start_date,1,4)=?""", (emp_id, t["id"], year), one=True)["s"]
        pend = q("""SELECT COALESCE(SUM(days),0) s FROM leaves WHERE employee_id=? AND leave_type_id=?
                    AND status='pending' AND substr(start_date,1,4)=?""", (emp_id, t["id"], year), one=True)["s"]
        allowed = t["days_per_year"] or 0
        remaining = (allowed - used) if allowed > 0 else None
        out.append({"id": t["id"], "name": t["name"], "allowed": allowed, "used": used, "pending": pend,
                    "paid": t["paid"], "remaining": remaining, "encash": bool(t["encash"]),
                    "payable_days": max(remaining or 0, 0) if t["encash"] and remaining is not None else 0})
    return out


ENCASH_BASIS = {"gross26": "Full salary (Basic + Allowances) ÷ 26", "basic26": "Basic salary ÷ 26",
                "gross30": "Full salary (Basic + Allowances) ÷ 30", "basic30": "Basic salary ÷ 30"}


def leave_day_rate(emp):
    """Salary for one leave day, used to pay out unused (encashable) leave."""
    basis = get_settings().get("encash_basis") or "gross26"
    basic, allow = emp["basic_salary"] or 0, emp["allowances"] or 0
    if basis == "basic30":
        return basic / 30
    if basis == "basic26":
        return basic / 26
    if basis == "gross30":
        return (basic + allow) / 30
    return (basic + allow) / 26


def check_leave(emp_id, type_id, start_s, end_s, exclude_id=None):
    """Validate leave dates/type/balance. Returns (error, leave_type, start, end, days)."""
    start, end = parse_date(start_s), parse_date(end_s)
    if not start or not end:
        return "Please choose valid start and end dates.", None, None, None, 0
    if end < start:
        return "End date cannot be before start date.", None, None, None, 0
    if start.year != end.year:
        return "A leave cannot span two calendar years. Please split it into two requests.", None, None, None, 0
    lt = q("SELECT * FROM leave_types WHERE id=?", (type_id,), one=True)
    if not lt:
        return "Please choose a leave type.", None, None, None, 0
    days = sum(1 for d in daterange(start, end) if is_workday(d))
    if days == 0:
        return "The selected dates are all weekly off days.", None, None, None, 0
    overlap = q("""SELECT 1 FROM leaves WHERE employee_id=? AND status IN ('pending','approved')
                   AND start_date<=? AND end_date>=? AND id<>?""",
                (emp_id, end.isoformat(), start.isoformat(), exclude_id or 0), one=True)
    if overlap:
        return "These dates overlap with another pending or approved leave.", None, None, None, 0
    if lt["days_per_year"] and lt["days_per_year"] > 0:
        bal = next(b for b in leave_balances(emp_id, start.year) if b["id"] == lt["id"])
        taken = bal["used"] + bal["pending"]
        if exclude_id:
            old = q("SELECT * FROM leaves WHERE id=?", (exclude_id,), one=True)
            if old and old["leave_type_id"] == lt["id"] and old["start_date"][:4] == str(start.year) \
                    and old["status"] in ("approved", "pending"):
                taken -= old["days"]
        if taken + days > lt["days_per_year"]:
            left = lt["days_per_year"] - taken
            return (f"Not enough {lt['name']} leave balance. Available: {left:g} day(s), requested: {days}.",
                    None, None, None, 0)
    return None, lt, start, end, days


def create_leave(emp_id, type_id, start_s, end_s, reason, status="pending", decided_by=None):
    """Validate and insert a leave. Returns error message or None."""
    err, lt, start, end, days = check_leave(emp_id, type_id, start_s, end_s)
    if err:
        return err
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    lid = x("""INSERT INTO leaves(employee_id,leave_type_id,start_date,end_date,days,reason,status,applied_on,
               decided_by,decided_on) VALUES(?,?,?,?,?,?,?,?,?,?)""",
            (emp_id, lt["id"], start.isoformat(), end.isoformat(), days, reason.strip(), status, now,
             decided_by, now if decided_by else None))
    if status == "approved":
        apply_leave_attendance(q("SELECT * FROM leaves WHERE id=?", (lid,), one=True))
    return None


def apply_leave_attendance(lv):
    note = "Approved leave #%d" % lv["id"]
    for d in daterange(parse_date(lv["start_date"]), parse_date(lv["end_date"])):
        if is_workday(d):
            db().execute("""INSERT INTO attendance(employee_id,date,status,note) VALUES(?,?,'leave',?)
                            ON CONFLICT(employee_id,date) DO UPDATE SET status='leave', note=excluded.note,
                            check_in=NULL, check_out=NULL""", (lv["employee_id"], d.isoformat(), note))
    db().commit()


def remove_leave_attendance(lv):
    x("DELETE FROM attendance WHERE employee_id=? AND status='leave' AND date BETWEEN ? AND ?",
      (lv["employee_id"], lv["start_date"], lv["end_date"]))


def compute_payroll(emp, y, m):
    dates = month_dates(y, m)
    first, last = dates[0], dates[-1]
    s = get_settings()
    workdays = [d for d in dates if is_workday(d)]
    jd = parse_date(emp["join_date"])
    not_joined = sum(1 for d in workdays if jd and d < jd)
    rows = q("SELECT status FROM attendance WHERE employee_id=? AND date BETWEEN ? AND ?",
             (emp["id"], first.isoformat(), last.isoformat()))
    c = Counter(r["status"] for r in rows)
    unpaid = 0
    for lv in q("""SELECT l.start_date,l.end_date FROM leaves l JOIN leave_types t ON t.id=l.leave_type_id
                   WHERE l.employee_id=? AND l.status='approved' AND t.paid=0
                   AND l.start_date<=? AND l.end_date>=?""", (emp["id"], last.isoformat(), first.isoformat())):
        a, b = max(parse_date(lv["start_date"]), first), min(parse_date(lv["end_date"]), last)
        unpaid += sum(1 for d in daterange(a, b) if is_workday(d))
    lates_per = int(to_float(s.get("lates_per_day"), 0))
    late_penalty = (c["late"] // lates_per) if lates_per > 0 else 0
    absent = c["absent"] + 0.5 * c["half_day"]
    present = c["present"] + c["late"] + 0.5 * c["half_day"]
    deduction_days = absent + unpaid + not_joined + late_penalty
    basic, allow = emp["basic_salary"] or 0, emp["allowances"] or 0
    per_day = (basic + allow) / len(workdays) if workdays else 0
    return {
        "basic": basic, "allowances": allow, "working_days": len(workdays),
        "present_days": present, "absent_days": absent, "late_days": c["late"],
        "unpaid_leave_days": unpaid, "deduction_days": deduction_days,
        "absent_deduction": round(per_day * deduction_days),
    }


def calc_net(p):
    return round((p["basic"] or 0) + (p["allowances"] or 0) + (p["overtime"] or 0) + (p["bonus"] or 0)
                 - (p["absent_deduction"] or 0) - (p["tax"] or 0) - (p["other_deductions"] or 0)
                 - ((p["loan_deduction"] if "loan_deduction" in p.keys() else 0) or 0))


# --------------------------------------------------------------------------- auth & security
@app.before_request
def _before():
    g.user = None
    if request.method == "POST":
        token = session.get("csrf")
        if not token or token != request.form.get("csrf"):
            abort(400, "Your form expired. Please go back, refresh the page and try again.")
    uid = session.get("uid")
    if uid:
        g.user = q("""SELECT u.*, e.name AS emp_name, COALESCE(r.name, u.role) AS role_name FROM users u
                      LEFT JOIN employees e ON e.id=u.employee_id LEFT JOIN roles r ON r.key=u.role
                      WHERE u.id=? AND u.active=1""", (uid,), one=True)
        if g.user is None:
            session.pop("uid", None)
        else:
            now = datetime.now().timestamp()
            idle = int(to_float(get_settings().get("idle_minutes"), 30) or 0)
            last = session.get("seen", now)
            if request.endpoint == "api_notifications":
                g.idle_expired = idle > 0 and now - last > idle * 60
                return None
            if idle > 0 and now - last > idle * 60 and request.endpoint not in ("static",):
                audit("Logged out (idle)", f"no activity for {idle} min", user=g.user["username"])
                token = session.get("csrf")
                session.clear()
                session["csrf"] = token or secrets.token_hex(16)
                g.user = None
                flash(f"You were signed out after {idle} minutes without activity. Please sign in again.", "warn")
                return redirect(url_for("login", next=request.path if request.method == "GET" else None))
            session["seen"] = now
    auto_backup()
    start_zkt_thread()


_last_auto = {"day": None}
BACKUP_KINDS = {"auto": "Auto", "manual": "Manual", "upload": "Uploaded", "before-restore": "Safety copy"}
_BK_NEW = re.compile(r"^backup-(\d{4})-(\d\d)-(\d\d)-(\d\d)-(\d\d)-(\d\d)-([a-z-]+\d*)\.db$")
_BK_OLD = re.compile(r"^(?:old-)?hr-backup-(\d{4})(\d\d)(\d\d)-(\d\d)(\d\d)(\d\d)\.db$")
_BK_SAFE = re.compile(r"^[A-Za-z0-9._-]{1,120}\.db$")


def make_backup(kind="manual"):
    """Copy the whole live database into backups/backup-YYYY-MM-DD-HH-MM-SS-<kind>.db and return the file name."""
    os.makedirs(BACKUP_DIR, exist_ok=True)
    stamp = datetime.now().strftime("%Y-%m-%d-%H-%M-%S")
    name = f"backup-{stamp}-{kind}.db"
    n = 1
    while os.path.exists(os.path.join(BACKUP_DIR, name)):
        n += 1
        name = f"backup-{stamp}-{kind}{n}.db"
    src = sqlite3.connect(DB_PATH, timeout=30)
    dest = sqlite3.connect(os.path.join(BACKUP_DIR, name))
    src.backup(dest)
    dest.close()
    src.close()
    return name


def backup_info(name):
    """Type, date/time and size of one backup file."""
    path = os.path.join(BACKUP_DIR, name)
    st = os.stat(path)
    when, kind = None, "manual"
    m = _BK_NEW.match(name)
    if m:
        when = datetime(*map(int, m.groups()[:6]))
        kind = re.sub(r"\d+$", "", m.group(7))
    else:
        m = _BK_OLD.match(name)
        if m:
            when = datetime(*map(int, m.groups()))
        elif name.startswith("auto-"):
            kind = "auto"
    when = when or datetime.fromtimestamp(st.st_mtime)
    return {"name": name, "kind": kind if kind in BACKUP_KINDS else "manual", "size": st.st_size, "when": when}


def list_backups():
    if not os.path.isdir(BACKUP_DIR):
        return []
    out = []
    for f in os.listdir(BACKUP_DIR):
        if f.endswith(".db") and _BK_SAFE.match(f) and os.path.isfile(os.path.join(BACKUP_DIR, f)):
            try:
                out.append(backup_info(f))
            except OSError:
                pass
    return sorted(out, key=lambda b: b["when"], reverse=True)


def auto_backup():
    """Once a day, save a full copy of the database (type Auto) and keep the last 30 automatic copies."""
    day = date.today().isoformat()
    if _last_auto["day"] == day:
        return
    _last_auto["day"] = day
    try:
        backups = list_backups()
        autos = [b for b in backups if b["kind"] == "auto"]
        if not any(b["when"].date() == date.today() for b in autos):
            make_backup("auto")
            autos = [b for b in list_backups() if b["kind"] == "auto"]
        for old in autos[AUTO_BACKUPS_KEEP:]:
            os.remove(os.path.join(BACKUP_DIR, old["name"]))
        _last_auto["error"] = None
    except Exception as e:                                       # noqa: BLE001 - never block the app
        _last_auto["error"] = str(e)[:200]
        print("Auto-backup failed:", e)


def client_ip():
    return request.headers.get("X-Forwarded-For", request.remote_addr or "")[:45] if request else ""


def audit(action, detail="", user=None):
    try:
        con = sqlite3.connect(DB_PATH, timeout=15)
        con.execute("INSERT INTO audit_log(at,username,ip,action,detail) VALUES(?,?,?,?,?)",
                    (datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                     user if user is not None else (g.user["username"] if g.get("user") else ""),
                     client_ip(), action, (detail or "")[:500]))
        con.commit()
        con.close()
    except Exception as e:                                       # noqa: BLE001
        print("Audit log failed:", e)


AUDIT_NAMES = {
    "employee_new": "Added employee", "employee_edit": "Edited employee", "employee_bulk_edit": "Bulk-edited employees",
    "employee_delete": "Deleted employee", "employee_account": "Changed employee login",
    "employee_import_confirm": "Imported employees", "department_new": "Saved department",
    "departments_manage": "Changed departments", "department_delete": "Deleted department",
    "attendance": "Marked attendance", "attendance_mark": "Quick-marked attendance", "attendance_manage": "Edited attendance",
    "attendance_import_confirm": "Imported attendance", "leaves": "Recorded leave", "leave_decide": "Leave decision", "leave_edit": "Edited leave",
    "payroll_generate": "Generated payroll", "payroll_edit": "Edited salary", "payroll_status": "Changed salary status",
    "payroll_pay_all": "Marked all salaries paid", "payroll_deduction": "Salary deduction decision",
    "payroll_deduct_all": "Applied all salary deductions", "payroll_delete": "Deleted salary row",
    "loans": "Gave a loan", "loan_view": "Changed a loan", "settings_page": "Changed settings", "recycle_bin": "Recycle bin",
    "users": "Changed users", "user_perms_page": "Changed a user's permissions", "roles_page": "Changed roles & permissions", "change_password": "Changed own password",
    "asset_new": "Added asset", "asset_edit": "Edited asset", "asset_assign": "Assigned asset",
    "asset_return": "Received asset back", "holidays": "Changed public holidays",
    "increments": "Gave increment / promotion", "increment_undo": "Undid salary change", "settlement_edit": "Final settlement",
    "settlement_delete": "Deleted / reopened settlement", "doc_upload": "Uploaded employee document",
    "doc_delete": "Deleted employee document", "doc_edit": "Edited employee document", "warnings_page": "Issued warning",
    "warning_update": "Changed warning", "asset_management": "Asset management", "asset_delete": "Deleted asset", "assets_assign": "Assigned asset", "assets_import": "Imported assets (page)",
}
_SKIP_FIELDS = {"csrf", "password", "new", "confirm", "current", "token", "changed"}


@app.after_request
def _after(resp):
    resp.headers.setdefault("X-Frame-Options", "DENY")
    resp.headers.setdefault("X-Content-Type-Options", "nosniff")
    resp.headers.setdefault("Referrer-Policy", "same-origin")
    if resp.mimetype == "text/html" or resp.headers.get("Content-Disposition"):
        resp.headers["Cache-Control"] = "no-store"             # don't keep HR data in the browser cache
    try:
        if request.method == "POST" and g.get("user") and resp.status_code < 400 and \
                request.endpoint not in ("login", "logout", "checkin", "checkout", "backups", "backup_restore", "backup_delete", "assets_import", "recycle_bin"):
            name = AUDIT_NAMES.get(request.endpoint, request.endpoint or "?")
            bits = [f"{k}={v}" for k, v in request.form.items() if k not in _SKIP_FIELDS and v and len(v) < 60][:8]
            ids = ", ".join(f"{k} {v}" for k, v in (request.view_args or {}).items())
            audit(name, " · ".join(x_ for x_ in [ids, ", ".join(bits)] if x_))
        elif request.endpoint == "backup" and g.get("user") and resp.status_code < 400:
            audit("Downloaded backup", "")
    except Exception as e:                                       # noqa: BLE001
        print("Audit failed:", e)
    return resp


def csrf_token():
    if "csrf" not in session:
        session["csrf"] = secrets.token_hex(16)
    return session["csrf"]


app.jinja_env.globals["csrf_token"] = csrf_token


def _split(v):
    return {p for p in (v or "").split(",") if p}


def role_perm_set(role):
    if role == "admin":
        return set(ALL_PERMS)
    r = q("SELECT perms FROM roles WHERE key=?", (role,), one=True)
    return _split(r["perms"] if r else "")


def effective_perms(u):
    """Role permissions, plus anything switched on for this user, minus anything switched off."""
    if u["role"] == "admin":
        return set(ALL_PERMS)
    keys = u.keys() if hasattr(u, "keys") else []
    add = _split(u["perm_add"]) if "perm_add" in keys else set()
    rem = _split(u["perm_remove"]) if "perm_remove" in keys else set()
    return ((role_perm_set(u["role"]) | add) - rem) & set(ALL_PERMS)


def has_access(u):
    """Only HR / managers may sign in: admin, or someone with at least one permission."""
    return bool(effective_perms(u))


def login_required(f):
    @wraps(f)
    def wrapper(*a, **kw):
        if g.user is None:
            return redirect(url_for("login", next=request.path))
        if not has_access(g.user):
            session.clear()
            flash("This system is for HR and managers only. Ask the administrator for access.", "error")
            return redirect(url_for("login"))
        if g.user["must_change"] and request.endpoint not in ("change_password", "logout"):
            flash("Please set a new password before continuing.", "warn")
            return redirect(url_for("change_password"))
        return f(*a, **kw)
    return wrapper


def user_perms():
    if "perms" not in g:
        u = g.get("user")
        if not u:
            g.perms = set()
        else:
            g.perms = effective_perms(u)
    return g.perms


def can(*perms):
    have = user_perms()
    return any(p in have for p in perms)


def perm_required(*perms):
    def deco(f):
        @wraps(f)
        @login_required
        def wrapper(*a, **kw):
            if not can(*perms):
                abort(403)
            return f(*a, **kw)
        return wrapper
    return deco


def is_hr():
    """True for any staff role that can see the management side of the system."""
    return bool(user_perms() - {"backup"})


app.jinja_env.globals["can"] = can


@app.context_processor
def _ctx():
    pending = 0
    emp_total = 0
    if can("leaves.view", "leaves.approve"):
        pending = q("SELECT COUNT(*) c FROM leaves WHERE status='pending'", one=True)["c"]
    if can("employees.view"):
        emp_total = q("SELECT COUNT(*) c FROM employees", one=True)["c"]
    loans_pending = 0
    if can("payroll.view"):
        loans_pending = q("""SELECT COUNT(*) c FROM loans l WHERE l.status<>'closed' AND l.amount -
                             COALESCE((SELECT SUM(amount) FROM loan_payments p WHERE p.loan_id=l.id),0) > 0.5""",
                          one=True)["c"]
    hour = datetime.now().hour
    greeting = "Good Morning" if hour < 12 else ("Good Afternoon" if hour < 17 else "Good Evening")
    return {"user": g.get("user"), "settings": get_settings(), "today": date.today(), "is_hr": is_hr(),
            "greeting": greeting, "now": datetime.now(), "emp_total": emp_total,
            "pending_leaves": pending, "loans_pending": loans_pending, "ATT_STATUSES": ATT_STATUSES, "ATT_CODES": ATT_CODES}


@app.template_filter("money")
def money(v):
    try:
        v = float(v or 0)
    except (TypeError, ValueError):
        return v
    return f"{v:,.0f}"


@app.template_filter("num")
def num(v):
    try:
        return f"{float(v or 0):g}"
    except (TypeError, ValueError):
        return v


@app.template_filter("nicedate")
def nicedate(s):
    d = parse_date(s) if isinstance(s, str) else s
    return d.strftime("%d %b %Y") if d else (s or "")


@app.template_filter("dayname")
def dayname(s):
    d = parse_date(s) if isinstance(s, str) else s
    return d.strftime("%A") if d else ""


@app.template_filter("monthname")
def monthname(s):
    y, m = parse_month(s)
    return f"{calendar.month_name[m]} {y}"


@app.errorhandler(403)
def _e403(e):
    return render_template("error.html", code=403, msg="You don't have permission to open this page."), 403


@app.errorhandler(404)
def _e404(e):
    return render_template("error.html", code=404, msg="Page not found."), 404


@app.errorhandler(400)
def _e400(e):
    return render_template("error.html", code=400, msg=getattr(e, "description", "Bad request")), 400


@app.errorhandler(500)
def _e500(e):
    """Show what went wrong and keep it in error.log (in the data folder) so it can be fixed."""
    import traceback
    orig = getattr(e, "original_exception", None) or e
    try:
        with open(os.path.join(DATA_DIR, "error.log"), "a", encoding="utf-8") as f:
            f.write(f"\n[{datetime.now():%Y-%m-%d %H:%M:%S}] {request.method} {request.path}\n")
            f.write("".join(traceback.format_exception(type(orig), orig, orig.__traceback__)))
    except Exception:                                                # noqa: BLE001
        pass
    return render_template("error.html", code=500,
                           msg=f"Something went wrong: {type(orig).__name__}: {orig}. Details were saved in error.log."), 500


# --------------------------------------------------------------------------- login
@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        uname_in = request.form.get("username", "").strip()
        u = q("SELECT * FROM users WHERE username=?", (uname_in,), one=True)
        now_s = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        if u and u["locked_until"] and u["locked_until"] > now_s:
            mins = max(1, int((datetime.strptime(u["locked_until"], "%Y-%m-%d %H:%M:%S") - datetime.now()).total_seconds() // 60) + 1)
            audit("Login blocked (locked)", "", user=uname_in)
            flash(f"Too many wrong passwords. This account is locked for {mins} more minute(s).", "error")
            return render_template("login.html")
        ok_pw = bool(u) and check_password_hash(u["password_hash"], request.form.get("password", ""))
        if not ok_pw:
            if u:
                fails = (u["failed_count"] or 0) + 1
                if fails >= MAX_FAILS:
                    until = (datetime.now() + timedelta(minutes=LOCK_MINUTES)).strftime("%Y-%m-%d %H:%M:%S")
                    x("UPDATE users SET failed_count=0, locked_until=? WHERE id=?", (until, u["id"]))
                    audit("Account locked", f"{MAX_FAILS} wrong passwords", user=uname_in)
                    flash(f"Too many wrong passwords. The account is locked for {LOCK_MINUTES} minutes.", "error")
                    return render_template("login.html")
                x("UPDATE users SET failed_count=? WHERE id=?", (fails, u["id"]))
            audit("Failed login", "wrong password" if u else "unknown username", user=uname_in)
        if u and u["active"] and ok_pw and not has_access(u):
            flash("This system is for HR and managers only. Ask the administrator for access.", "error")
            return render_template("login.html")
        if u and u["active"] and check_password_hash(u["password_hash"], request.form.get("password", "")):
            token = session.get("csrf")
            session.clear()
            session["csrf"] = token or secrets.token_hex(16)
            session.permanent = True
            session["uid"] = u["id"]
            session["seen"] = datetime.now().timestamp()
            x("UPDATE users SET last_login=?, failed_count=0, locked_until=NULL WHERE id=?",
              (datetime.now().strftime("%Y-%m-%d %H:%M"), u["id"]))
            audit("Signed in", "", user=u["username"])
            nxt = request.args.get("next", "")
            return redirect(nxt if nxt.startswith("/") and not nxt.startswith("//") else url_for("dashboard"))
        flash("Wrong username or password, or the account is disabled.", "error")
    return render_template("login.html")


@app.post("/logout")
def logout():
    if session.get("uid"):
        r_ = q("SELECT username FROM users WHERE id=?", (session["uid"],), one=True)
        audit("Signed out", "", user=r_["username"] if r_ else "")
    session.clear()
    return redirect(url_for("login"))


@app.route("/change-password", methods=["GET", "POST"])
@login_required
def change_password():
    if request.method == "POST":
        cur, new, conf = (request.form.get(k, "") for k in ("current", "new", "confirm"))
        if not check_password_hash(g.user["password_hash"], cur):
            flash("Current password is wrong.", "error")
        elif len(new) < 6:
            flash("New password must be at least 6 characters.", "error")
        elif new != conf:
            flash("New passwords do not match.", "error")
        else:
            x("UPDATE users SET password_hash=?, must_change=0 WHERE id=?", (generate_password_hash(new), g.user["id"]))
            flash("Password changed.", "ok")
            return redirect(url_for("dashboard"))
    return render_template("change_password.html")


# --------------------------------------------------------------------------- dashboard
@app.route("/")
@login_required
def dashboard():
    today = date.today().isoformat()
    if is_hr():
        active = q("SELECT COUNT(*) c FROM employees WHERE status='active'", one=True)["c"]
        by = {r["status"]: r["c"] for r in q(
            """SELECT a.status, COUNT(*) c FROM attendance a JOIN employees e ON e.id=a.employee_id
               WHERE a.date=? AND e.status='active' GROUP BY a.status""", (today,))}
        marked = sum(by.values())
        depts = q("""SELECT d.name, COUNT(e.id) c FROM departments d LEFT JOIN employees e
                     ON e.department_id=d.id AND e.status='active' GROUP BY d.id ORDER BY c DESC, d.name""")
        pend = q("""SELECT l.*, e.name, t.name tname FROM leaves l JOIN employees e ON e.id=l.employee_id
                    JOIN leave_types t ON t.id=l.leave_type_id WHERE l.status='pending' ORDER BY l.start_date LIMIT 6""")
        on_leave = q("""SELECT e.name, t.name tname, l.end_date FROM leaves l JOIN employees e ON e.id=l.employee_id
                        JOIN leave_types t ON t.id=l.leave_type_id WHERE l.status='approved'
                        AND l.start_date<=? AND l.end_date>=? ORDER BY e.name""", (today, today))
        tomorrow = (date.today() + timedelta(days=1)).isoformat()
        upcoming = q("""SELECT e.name, t.name tname, l.start_date, l.end_date, l.days FROM leaves l
                        JOIN employees e ON e.id=l.employee_id JOIN leave_types t ON t.id=l.leave_type_id
                        WHERE l.status='approved' AND l.start_date BETWEEN ? AND ? ORDER BY l.start_date LIMIT 5""",
                     (tomorrow, (date.today() + timedelta(days=14)).isoformat()))
        upcoming_n = q("""SELECT COUNT(*) c FROM leaves WHERE status='approved' AND start_date BETWEEN ? AND ?""",
                       (tomorrow, (date.today() + timedelta(days=14)).isoformat()), one=True)["c"]
        m0 = date.today().replace(day=1).isoformat()
        m1 = (date.today().replace(day=28) + timedelta(days=4)).replace(day=1).isoformat()
        leave_month = q("""SELECT COUNT(*) c, COALESCE(SUM(days),0) d FROM leaves WHERE status='approved'
                           AND start_date<? AND end_date>=?""", (m1, m0), one=True)
        late_today = q("""SELECT e.id, e.name, a.check_in, d.name dept FROM attendance a JOIN employees e ON e.id=a.employee_id
                          LEFT JOIN departments d ON d.id=e.department_id
                          WHERE a.date=? AND a.status='late' ORDER BY a.check_in DESC LIMIT 6""", (today,))
        absent_today = q("""SELECT e.id, e.name, d.name dept FROM attendance a JOIN employees e ON e.id=a.employee_id
                            LEFT JOIN departments d ON d.id=e.department_id
                            WHERE a.date=? AND a.status='absent' ORDER BY e.name LIMIT 6""", (today,))
        last_sync = q("SELECT MAX(date) d FROM attendance", one=True)["d"]
        month = date.today().strftime("%Y-%m")
        pay = q("""SELECT COUNT(*) c, COALESCE(SUM(net),0) s, COALESCE(SUM(status='paid'),0) p,
                   COALESCE(SUM(CASE WHEN status='paid' THEN net END),0) paid_amt,
                   COALESCE(SUM(absent_deduction+tax+other_deductions),0) ded FROM payroll WHERE month=?""",
                (month,), one=True)
        monthly_bill = q("SELECT COALESCE(SUM(basic_salary+allowances),0) s FROM employees WHERE status='active'",
                         one=True)["s"]
        # last 7 days attendance trend
        trend = []
        d = date.today()
        while len(trend) < 7:
            if is_workday(d):
                c = {r["status"]: r["c"] for r in q(
                    "SELECT status, COUNT(*) c FROM attendance WHERE date=? GROUP BY status", (d.isoformat(),))}
                trend.append({"d": d, "present": c.get("present", 0), "late": c.get("late", 0) + c.get("half_day", 0),
                              "leave": c.get("leave", 0), "absent": c.get("absent", 0)})
            d -= timedelta(days=1)
            if (date.today() - d).days > 30:
                break
        trend.reverse()
        tmax = max([active] + [t["present"] + t["late"] + t["leave"] + t["absent"] for t in trend] or [1]) or 1
        new_joiners = q("""SELECT e.name, e.designation, e.join_date, d.name dept FROM employees e
                           LEFT JOIN departments d ON d.id=e.department_id WHERE e.status='active'
                           ORDER BY e.join_date DESC, e.id DESC LIMIT 5""")
        mmdd = date.today().strftime("%m-%d")
        birthdays = q("""SELECT name, dob FROM employees WHERE status='active' AND dob IS NOT NULL AND dob<>''
                         AND substr(dob,6,5) BETWEEN ? AND ? ORDER BY substr(dob,6,5) LIMIT 6""",
                      (mmdd, (date.today() + timedelta(days=30)).strftime("%m-%d")))
        return render_template("dashboard_hr.html", active=active, by=by, unmarked=max(active - marked, 0),
                               depts=depts, pend=pend, on_leave=on_leave, pay=pay, month=month,
                               workday=is_workday(date.today()), trend=trend, tmax=tmax,
                               monthly_bill=monthly_bill, new_joiners=new_joiners, birthdays=birthdays,
                               dept_max=max([r["c"] for r in depts] or [1]) or 1, upcoming=upcoming,
                               upcoming_n=upcoming_n, leave_month=leave_month, late_today=late_today,
                               absent_today=absent_today, last_sync=last_sync)
    return my_dashboard()


def my_dashboard():
    emp_id = g.user["employee_id"]
    if not emp_id:
        return render_template("error.html", code="", msg="Your account is not linked to an employee record.")
    today = date.today()
    att = q("SELECT * FROM attendance WHERE employee_id=? AND date=?", (emp_id, today.isoformat()), one=True)
    month = today.strftime("%Y-%m")
    c = {r["status"]: r["c"] for r in q(
        "SELECT status, COUNT(*) c FROM attendance WHERE employee_id=? AND substr(date,1,7)=? GROUP BY status",
        (emp_id, month))}
    leaves = q("""SELECT l.*, t.name tname FROM leaves l JOIN leave_types t ON t.id=l.leave_type_id
                  WHERE l.employee_id=? ORDER BY l.id DESC LIMIT 5""", (emp_id,))
    emp = q("""SELECT e.*, d.name dept FROM employees e LEFT JOIN departments d ON d.id=e.department_id
               WHERE e.id=?""", (emp_id,), one=True)
    return render_template("dashboard_emp.html", att=att, counts=c, leaves=leaves, emp=emp,
                           balances=leave_balances(emp_id), workday=is_workday(today), now=datetime.now())


@app.post("/checkin")
@login_required
def checkin():
    emp_id = g.user["employee_id"] or abort(403)
    now = datetime.now()
    today = now.date().isoformat()
    row = q("SELECT * FROM attendance WHERE employee_id=? AND date=?", (emp_id, today), one=True)
    if row and row["check_in"]:
        flash(f"You already checked in at {row['check_in']}.", "warn")
    elif row and row["status"] == "leave":
        flash("You are on approved leave today.", "warn")
    else:
        s = get_settings()
        start = parse_time(s.get("office_start")) or datetime.strptime("09:00", "%H:%M").time()
        limit = datetime.combine(now.date(), start) + timedelta(minutes=int(to_float(s.get("grace_minutes"), 0)))
        status = "late" if now > limit else "present"
        x("""INSERT INTO attendance(employee_id,date,status,check_in) VALUES(?,?,?,?)
             ON CONFLICT(employee_id,date) DO UPDATE SET status=excluded.status, check_in=excluded.check_in""",
          (emp_id, today, status, now.strftime("%H:%M")))
        flash(f"Checked in at {now:%H:%M}" + (" (marked late)." if status == "late" else "."), "ok")
    return redirect(url_for("dashboard"))


@app.post("/checkout")
@login_required
def checkout():
    emp_id = g.user["employee_id"] or abort(403)
    now = datetime.now()
    row = q("SELECT * FROM attendance WHERE employee_id=? AND date=?", (emp_id, now.date().isoformat()), one=True)
    if not row or not row["check_in"]:
        flash("Please check in first.", "warn")
    else:
        x("UPDATE attendance SET check_out=? WHERE id=?", (now.strftime("%H:%M"), row["id"]))
        flash(f"Checked out at {now:%H:%M}. Have a good evening!", "ok")
    return redirect(url_for("dashboard"))


# --------------------------------------------------------------------------- employees
EMP_FIELDS = ["name", "father_name", "cnic", "gender", "dob", "email", "phone", "address", "department_id",
              "designation", "join_date", "employment_type", "basic_salary", "allowances", "bank_account",
              "emergency_contact", "status", "manager_id"]


def manager_choices(exclude=None):
    return q("""SELECT id, name, emp_code, designation FROM employees WHERE status='active' AND id<>?
                ORDER BY name""", (exclude or 0,))


def manager_loop(eid, mid):
    """True if making `mid` the manager of `eid` would create a circle (A reports to B reports to A)."""
    seen = set()
    while mid and mid not in seen:
        if mid == eid:
            return True
        seen.add(mid)
        r = q("SELECT manager_id FROM employees WHERE id=?", (mid,), one=True)
        mid = r["manager_id"] if r else None
    return False


def emp_form(**kw):
    kw.setdefault("depts", all_depts())
    kw.setdefault("managers", manager_choices(kw.get("eid")))
    return render_template("employee_form.html", **kw)


@app.route("/employees")
@perm_required("employees.view")
def employees():
    sql = """SELECT e.*, d.name dept, m.name mgr_name, m.emp_code mgr_code,
                    (SELECT COUNT(*) FROM employees t WHERE t.manager_id=e.id AND t.status='active') team_size
             FROM employees e LEFT JOIN departments d ON d.id=e.department_id
             LEFT JOIN employees m ON m.id=e.manager_id WHERE 1=1"""
    args = []
    term = request.args.get("q", "").strip()
    if term:
        sql += " AND (e.name LIKE ? OR e.emp_code LIKE ? OR e.cnic LIKE ? OR e.phone LIKE ? OR e.designation LIKE ?)"
        args += [f"%{term}%"] * 5
    if request.args.get("dept"):
        sql += " AND e.department_id=?"
        args.append(request.args["dept"])
    if request.args.get("manager", "").isdigit():
        sql += " AND e.manager_id=?"
        args.append(int(request.args["manager"]))
    status = request.args.get("status", "active")
    if status in ("active", "inactive"):
        sql += " AND e.status=?"
        args.append(status)
    rows = q(sql + " ORDER BY e.name", args)
    if request.args.get("export") == "csv":
        cols = ["emp_code", "name", "father_name", "cnic", "gender", "dob", "email", "phone", "address", "dept",
                "designation", "join_date", "employment_type", "basic_salary", "allowances", "bank_account",
                "emergency_contact", "status", "mgr_code", "mgr_name"]
        return csv_response("employees.csv", cols, [[r[c] for c in cols] for r in rows])
    managers = q("""SELECT DISTINCT m.id, m.name FROM employees e JOIN employees m ON m.id=e.manager_id ORDER BY m.name""")
    return render_template("employees.html", rows=rows, depts=q("SELECT * FROM departments ORDER BY name"),
                           status=status, managers=managers)


def emp_form_data():
    d = {f: request.form.get(f, "").strip() for f in EMP_FIELDS}
    d["basic_salary"] = to_float(d["basic_salary"])
    d["allowances"] = to_float(d["allowances"])
    d["department_id"] = int(d["department_id"]) if d["department_id"].isdigit() else None
    d["status"] = d["status"] if d["status"] in ("active", "inactive") else "active"
    d["manager_id"] = int(d["manager_id"]) if d["manager_id"].isdigit() else None
    return d


@app.route("/employees/new", methods=["GET", "POST"])
@perm_required("employees.manage")
def employee_new():
    if request.method == "POST":
        d = emp_form_data()
        code = request.form.get("emp_code", "").strip() or next_emp_code()
        if not d["name"]:
            flash("Name is required.", "error")
        elif q("SELECT 1 FROM employees WHERE emp_code=?", (code,), one=True):
            flash(f"Employee code {code} is already used.", "error")
        else:
            uname = request.form.get("username", "").strip()
            pw = request.form.get("password", "")
            if request.form.get("create_login"):
                if not uname or len(pw) < 6:
                    flash("For a login, enter a username and a password of at least 6 characters.", "error")
                    return emp_form(e=d, code=code, new=True)
                if q("SELECT 1 FROM users WHERE username=?", (uname,), one=True):
                    flash(f"Username '{uname}' is already taken.", "error")
                    return emp_form(e=d, code=code, new=True)
            cols = ["emp_code", "source"] + EMP_FIELDS
            eid = x(f"INSERT INTO employees({','.join(cols)}) VALUES({','.join('?' * len(cols))})",
                    [code, "manual"] + [d[f] for f in EMP_FIELDS])
            if request.form.get("create_login"):
                x("INSERT INTO users(username,password_hash,role,employee_id,must_change) VALUES(?,?,?,?,1)",
                  (uname, generate_password_hash(pw), "employee", eid))
            flash(f"Employee {d['name']} ({code}) added.", "ok")
            return redirect(url_for("employee_view", eid=eid))
        return emp_form(e=d, code=code, new=True)
    return emp_form(e={"status": "active", "join_date": date.today().isoformat()}, code=next_emp_code(), new=True)


def all_depts():
    return q("SELECT * FROM departments ORDER BY name")


def get_emp(eid):
    e = q("""SELECT e.*, d.name dept, m.name mgr_name, m.emp_code mgr_code, m.designation mgr_desig
             FROM employees e LEFT JOIN departments d ON d.id=e.department_id
             LEFT JOIN employees m ON m.id=e.manager_id WHERE e.id=?""", (eid,), one=True)
    if not e:
        abort(404)
    return e


@app.route("/employees/<int:eid>/edit", methods=["GET", "POST"])
@perm_required("employees.manage")
def employee_edit(eid):
    e = get_emp(eid)
    if request.method == "POST":
        d = emp_form_data()
        code = request.form.get("emp_code", "").strip() or e["emp_code"]
        if not d["name"]:
            flash("Name is required.", "error")
        elif q("SELECT 1 FROM employees WHERE emp_code=? AND id<>?", (code, eid), one=True):
            flash(f"Employee code {code} is already used.", "error")
        elif d["manager_id"] and manager_loop(eid, d["manager_id"]):
            flash("That manager reports to this employee (directly or through others). Choose a different manager.", "error")
        else:
            x(f"UPDATE employees SET emp_code=?, {', '.join(f + '=?' for f in EMP_FIELDS)} WHERE id=?",
              [code] + [d[f] for f in EMP_FIELDS] + [eid])
            log_salary_change(eid, e, d.get("basic_salary"), d.get("allowances"), d.get("designation"), "Changed in Edit profile")
            if d["status"] == "inactive":
                x("UPDATE users SET active=0 WHERE employee_id=?", (eid,))
            flash("Employee updated.", "ok")
            return redirect(url_for("employee_view", eid=eid))
        return emp_form(e=d, code=code, new=False, eid=eid)
    return emp_form(e=e, code=e["emp_code"], new=False, eid=eid)


QUICK_FIELDS = ["name", "department_id", "designation", "manager_id", "phone", "cnic", "join_date",
                "employment_type", "basic_salary", "allowances", "status"]


@app.route("/employees/edit", methods=["GET", "POST"])
@perm_required("employees.manage")
def employee_bulk_edit():
    if request.method == "POST":
        ids = [int(i) for i in request.form.getlist("changed") if i.isdigit()]
        saved, problems = 0, []
        valid_depts = {r["id"] for r in q("SELECT id FROM departments")}
        for eid in ids:
            e = q("SELECT * FROM employees WHERE id=?", (eid,), one=True)
            if not e:
                continue
            vals = {f: request.form.get(f"{f}_{eid}", "").strip() for f in QUICK_FIELDS}
            if not vals["name"]:
                problems.append(f"{e['emp_code']}: name can't be empty")
                continue
            vals["department_id"] = int(vals["department_id"]) if vals["department_id"].isdigit() \
                and int(vals["department_id"]) in valid_depts else None
            mid = int(vals["manager_id"]) if vals["manager_id"].isdigit() else None
            if mid and (mid == eid or manager_loop(eid, mid)):
                problems.append(f"{e['emp_code']}: manager not changed (would report to themselves)")
                mid = e["manager_id"]
            vals["manager_id"] = mid
            vals["basic_salary"] = to_float(vals["basic_salary"], e["basic_salary"] or 0)
            vals["allowances"] = to_float(vals["allowances"], e["allowances"] or 0)
            vals["status"] = vals["status"] if vals["status"] in ("active", "inactive") else e["status"]
            if vals["join_date"] and not parse_date(vals["join_date"]):
                vals["join_date"] = e["join_date"]
            db().execute(f"UPDATE employees SET {', '.join(f + '=?' for f in QUICK_FIELDS)} WHERE id=?",
                         [vals[f] for f in QUICK_FIELDS] + [eid])
            log_salary_change(eid, e, vals["basic_salary"], vals["allowances"], vals["designation"], "Changed in Edit Employees", commit=False)
            if vals["status"] == "inactive" and e["status"] != "inactive":
                db().execute("UPDATE users SET active=0 WHERE employee_id=?", (eid,))
            saved += 1
        db().commit()
        if saved:
            flash(f"Saved changes for {saved} employee(s).", "ok")
        for p_ in problems:
            flash(p_, "error")
        if not saved and not problems:
            flash("No changes to save.", "warn")
        return redirect(url_for("employee_bulk_edit", q=request.form.get("q", ""), dept=request.form.get("dept", ""),
                                status=request.form.get("status", "all")))
    sql = "SELECT e.* FROM employees e WHERE 1=1"
    args = []
    term = request.args.get("q", "").strip()
    if term:
        sql += " AND (e.name LIKE ? OR e.emp_code LIKE ? OR e.cnic LIKE ? OR e.phone LIKE ? OR e.designation LIKE ?)"
        args += [f"%{term}%"] * 5
    if request.args.get("dept"):
        sql += " AND e.department_id=?"
        args.append(request.args["dept"])
    status = request.args.get("status", "all")
    if status in ("active", "inactive"):
        sql += " AND e.status=?"
        args.append(status)
    rows = q(sql + " ORDER BY e.name", args)
    return render_template("employee_bulk_edit.html", rows=rows, depts=all_depts(), status=status, term=term,
                           managers=manager_choices())


@app.route("/employees/<int:eid>")
@perm_required("employees.view")
def employee_view(eid):
    e = get_emp(eid)
    y, m = parse_month(request.args.get("month"))
    month = f"{y:04d}-{m:02d}"
    recs = {r["date"]: r for r in q("SELECT * FROM attendance WHERE employee_id=? AND substr(date,1,7)=?", (eid, month))}
    counts = Counter(r["status"] for r in recs.values())
    first = date(y, m, 1)
    cal = [None] * first.weekday() + [
        {"d": d, "rec": recs.get(d.isoformat()), "work": is_workday(d), "future": d > date.today(),
         "hol": holiday_name(d)} for d in month_dates(y, m)]
    workdays_so_far = sum(1 for d in month_dates(y, m) if is_workday(d) and d <= date.today())
    present = counts["present"] + counts["late"] + 0.5 * counts["half_day"]
    att_pct = round(present * 100 / workdays_so_far) if workdays_so_far and recs else None
    pays = q("SELECT * FROM payroll WHERE employee_id=? ORDER BY month DESC LIMIT 12", (eid,)) if can("payroll.view") else []
    acct = q("""SELECT u.*, r.name role_name FROM users u LEFT JOIN roles r ON r.key=u.role WHERE u.employee_id=?""",
             (eid,), one=True)
    leaves = q("""SELECT l.*, t.name tname FROM leaves l JOIN leave_types t ON t.id=l.leave_type_id
                  WHERE l.employee_id=? ORDER BY l.start_date DESC LIMIT 10""", (eid,))
    assets_now = q("""SELECT a.*, s.assigned_on, s.expected_return, s.condition_out FROM asset_assignments s
                      JOIN assets a ON a.id=s.asset_id WHERE s.employee_id=? AND s.returned_on IS NULL
                      ORDER BY s.assigned_on""", (eid,))
    assets_past = q("""SELECT a.asset_code, a.name, a.id, s.assigned_on, s.returned_on, s.condition_in FROM asset_assignments s
                       JOIN assets a ON a.id=s.asset_id WHERE s.employee_id=? AND s.returned_on IS NOT NULL
                       ORDER BY s.returned_on DESC LIMIT 10""", (eid,))
    jd = parse_date(e["join_date"])
    tenure = None
    if jd and jd <= date.today():
        months = (date.today().year - jd.year) * 12 + date.today().month - jd.month - (date.today().day < jd.day)
        tenure = (f"{months // 12} yr " if months >= 12 else "") + f"{months % 12} mo"
    is_head = q("SELECT name FROM departments WHERE head_id=?", (eid,))
    team = q("""SELECT e.id, e.name, e.emp_code, e.designation, e.status FROM employees e
                WHERE e.manager_id=? ORDER BY e.status, e.name""", (eid,))
    emp_loans = [dict(r, balance=max(r["amount"] - r["paid"], 0), state=loan_state(r))
                 for r in q(LOAN_SQL + " WHERE l.employee_id=? ORDER BY l.issued_on DESC", (eid,))] \
        if can("payroll.view") else []
    bal = leave_balances(eid)
    rate = leave_day_rate(e)
    for b in bal:
        b["payable"] = round(b["payable_days"] * rate)
    encash_total = sum(b["payable"] for b in bal)
    incs = q("SELECT * FROM increments WHERE employee_id=? ORDER BY eff_date DESC, id DESC", (eid,)) \
        if can("payroll.view") else []
    docs = [dict(r, exp_state=doc_state(r["expiry"])) for r in
            q("SELECT * FROM emp_documents WHERE employee_id=? ORDER BY uploaded_at DESC", (eid,))]
    warns = q("SELECT * FROM warnings WHERE employee_id=? ORDER BY issued_on DESC, id DESC", (eid,))
    sett = q("SELECT * FROM settlements WHERE employee_id=? ORDER BY id DESC LIMIT 1", (eid,), one=True) \
        if can("payroll.view") else None
    return render_template("employee_view.html", incs=incs, docs=docs, warns=warns, sett=sett, DOC_TYPES=DOC_TYPES,
                           WARN_KINDS=WARN_KINDS, encash_total=encash_total, day_rate=rate,
                           encash_label=ENCASH_BASIS.get(get_settings().get("encash_basis") or "gross26"), team=team, emp_loans=emp_loans, e=e, counts=counts, pays=pays, acct=acct, leaves=leaves,
                           balances=bal, month=month, cal=cal, att_pct=att_pct, tenure=tenure,
                           assets_now=assets_now, assets_past=assets_past, is_head=is_head,
                           prev=shift_month(y, m, -1), next=shift_month(y, m, 1), workdays=workdays_so_far)


@app.post("/employees/<int:eid>/account")
@perm_required("employees.manage")
def employee_account(eid):
    e = get_emp(eid)
    acct = q("SELECT * FROM users WHERE employee_id=?", (eid,), one=True)
    action = request.form.get("action")
    pw = request.form.get("password", "")
    if action == "create" and not acct:
        uname = request.form.get("username", "").strip()
        if not uname or len(pw) < 6:
            flash("Enter a username and a password of at least 6 characters.", "error")
        elif q("SELECT 1 FROM users WHERE username=?", (uname,), one=True):
            flash(f"Username '{uname}' is already taken.", "error")
        else:
            x("INSERT INTO users(username,password_hash,role,employee_id,must_change) VALUES(?,?,'employee',?,1)",
              (uname, generate_password_hash(pw), eid))
            flash(f"Login created for {e['name']}. They'll be asked to change the password on first login.", "ok")
    elif acct and action == "reset":
        if len(pw) < 6:
            flash("Password must be at least 6 characters.", "error")
        else:
            x("UPDATE users SET password_hash=?, must_change=1 WHERE id=?", (generate_password_hash(pw), acct["id"]))
            flash("Password reset. The employee must change it on next login.", "ok")
    elif acct and action == "toggle":
        x("UPDATE users SET active=? WHERE id=?", (0 if acct["active"] else 1, acct["id"]))
        flash("Login " + ("disabled." if acct["active"] else "enabled."), "ok")
    return redirect(url_for("employee_view", eid=eid))


@app.post("/employees/<int:eid>/delete")
@perm_required("employees.delete")
def employee_delete(eid):
    e = get_emp(eid)
    to_bin("employee", eid)
    x("UPDATE departments SET head_id=NULL WHERE head_id=?", (eid,))
    x("UPDATE employees SET manager_id=NULL WHERE manager_id=?", (eid,))
    for r in q("SELECT asset_id FROM asset_assignments WHERE employee_id=? AND returned_on IS NULL", (eid,)):
        x("UPDATE assets SET status='available' WHERE id=?", (r["asset_id"],))
    x("""UPDATE asset_assignments SET returned_on=?, return_notes='Employee record deleted'
         WHERE employee_id=? AND returned_on IS NULL""", (date.today().isoformat(), eid))
    x("DELETE FROM employees WHERE id=?", (eid,))
    flash(f"{e['name']} and all their records were deleted.", "ok")
    return redirect(url_for("employees"))


# --------------------------------------------------------------------------- employee import
IMPORT_DIR = os.path.join(DATA_DIR, "imports")
IMPORT_COLS = ["emp_code", "name", "father_name", "cnic", "gender", "dob", "email", "phone", "address",
               "department", "designation", "manager", "join_date", "employment_type", "basic_salary", "allowances",
               "bank_account", "emergency_contact", "status"]
IMPORT_LABELS = {"emp_code": "Employee Code", "name": "Full Name", "father_name": "Father/Husband Name",
                 "cnic": "CNIC", "gender": "Gender", "dob": "Date of Birth", "email": "Email", "phone": "Phone",
                 "address": "Address", "department": "Department", "designation": "Designation",
                 "manager": "Reporting Manager",
                 "join_date": "Joining Date", "employment_type": "Employment Type", "basic_salary": "Basic Salary",
                 "allowances": "Allowances", "bank_account": "Bank Account", "emergency_contact": "Emergency Contact",
                 "status": "Status", "username": "Login Username", "password": "Login Password"}
HEADER_ALIASES = {
    "emp_code": ["emp_code", "employee_code", "code", "emp_id", "employee_id", "emp_no", "employee_no", "staff_id", "id"],
    "name": ["name", "full_name", "employee_name", "emp_name", "staff_name"],
    "father_name": ["father_name", "father_husband_name", "father", "father_s_name", "husband_name", "f_name", "s_o_d_o"],
    "cnic": ["cnic", "nic", "cnic_no", "national_id", "id_card", "cnic_number"],
    "gender": ["gender", "sex"],
    "dob": ["dob", "date_of_birth", "birth_date", "birthday"],
    "email": ["email", "email_address", "e_mail"],
    "phone": ["phone", "mobile", "contact", "cell", "phone_no", "mobile_no", "contact_no", "phone_number"],
    "address": ["address", "home_address", "residential_address"],
    "department": ["department", "dept", "department_name"],
    "designation": ["designation", "job_title", "title", "position", "post"],
    "manager": ["manager", "reporting_manager", "reports_to", "supervisor", "line_manager", "manager_code",
                "manager_name", "reporting_to", "boss", "team_lead", "manager_emp_code"],
    "join_date": ["join_date", "joining_date", "date_of_joining", "doj", "joined", "hire_date"],
    "employment_type": ["employment_type", "type", "job_type", "contract_type"],
    "basic_salary": ["basic_salary", "basic", "salary", "basic_pay", "monthly_salary"],
    "allowances": ["allowances", "allowance", "total_allowances"],
    "bank_account": ["bank_account", "account", "account_no", "iban", "bank_account_iban", "account_number"],
    "emergency_contact": ["emergency_contact", "emergency", "emergency_phone"],
    "status": ["status", "active"],
}
_ALIAS = {a: k for k, v in HEADER_ALIASES.items() for a in v}


def _norm_header(h):
    h = str(h or "").strip().lower()
    out = "".join(c if c.isalnum() else "_" for c in h)
    while "__" in out:
        out = out.replace("__", "_")
    return out.strip("_")


def _cell_text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return str(v).strip()


def _parse_any_date(v):
    if v is None or v == "":
        return "", None
    if isinstance(v, datetime):
        return v.date().isoformat(), None
    if isinstance(v, date):
        return v.isoformat(), None
    if isinstance(v, (int, float)) and 20000 < v < 80000:     # Excel serial number
        return (date(1899, 12, 30) + timedelta(days=int(v))).isoformat(), None
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d-%b-%Y", "%d %b %Y", "%d-%B-%Y", "%d %B %Y",
                "%Y/%m/%d", "%d/%m/%y", "%d-%m-%y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).date().isoformat(), None
        except ValueError:
            pass
    return "", f"date '{s}' not understood (use DD/MM/YYYY)"


def read_upload(fs):
    """Return list of raw rows (lists) from an uploaded CSV or XLSX file."""
    name = (fs.filename or "").lower()
    data = fs.read()
    if name.endswith((".xlsx", ".xlsm")):
        try:
            import openpyxl
        except ImportError:
            raise ValueError("Excel support is not installed yet. Close the server and run Start Server again, "
                             "or save the file as CSV from Excel (File → Save As → CSV).")
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if name.endswith(".xls"):
        raise ValueError("Old .xls files are not supported. In Excel use File → Save As → Excel Workbook (.xlsx) or CSV.")
    if name.endswith((".csv", ".txt")):
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        sample = text[:2000]
        delim = ";" if sample.count(";") > sample.count(",") else ("\t" if sample.count("\t") > sample.count(",") else ",")
        return [r for r in csv.reader(io.StringIO(text), delimiter=delim)]
    raise ValueError("Please upload an Excel (.xlsx) or CSV file.")


def analyse_import(raw, update_existing):
    # find header row: first row with a recognised 'name' column (within first 10 rows)
    hdr_idx, mapping = None, {}
    for i, r in enumerate(raw[:10]):
        m = {}
        for j, h in enumerate(r):
            key = _ALIAS.get(_norm_header(h))
            if key and key not in m.values():
                m[j] = key
        if "name" in m.values():
            hdr_idx, mapping = i, m
            break
    if hdr_idx is None:
        raise ValueError("Couldn't find a 'Name' column in the first rows. Use the template, or make sure the first row has headings like Name, Department, Basic Salary.")
    depts = {r["name"].lower(): r["id"] for r in q("SELECT id, name FROM departments")}
    by_code = {r["emp_code"].lower(): r["id"] for r in q("SELECT id, emp_code FROM employees")}
    by_cnic = {r["cnic"].replace("-", ""): r["id"] for r in q("SELECT id, cnic FROM employees WHERE cnic IS NOT NULL AND cnic<>''")}
    taken_users = {r["username"].lower() for r in q("SELECT username FROM users")}
    emp_names = {}
    for r_ in q("SELECT name FROM employees"):
        emp_names[r_["name"].lower()] = emp_names.get(r_["name"].lower(), 0) + 1
    seen_codes, seen_users, seen_cnics, rows = set(), set(), set(), []
    new_depts = set()
    for n, r in enumerate(raw[hdr_idx + 1:], start=hdr_idx + 2):
        rec = {}
        for j, key in mapping.items():
            rec[key] = r[j] if j < len(r) else None
        if not any(_cell_text(v) for v in rec.values()):
            continue                                            # blank line
        errors, warns = [], []
        d = {k: _cell_text(rec.get(k)) for k in IMPORT_COLS}
        for k in ("dob", "join_date"):
            iso, err = _parse_any_date(rec.get(k))
            d[k] = iso
            if err:
                warns.append(f"{IMPORT_LABELS[k]}: {err} – left empty")
        for k in ("basic_salary", "allowances"):
            raw_v = d[k]
            d[k] = to_float(raw_v, None) if raw_v else None
            if raw_v and d[k] is None:
                errors.append(f"{IMPORT_LABELS[k]} '{raw_v}' is not a number")
        if not d["name"]:
            errors.append("Name is empty")
        if d["cnic"] and d["cnic"].isdigit() and len(d["cnic"]) == 13:
            d["cnic"] = f"{d['cnic'][:5]}-{d['cnic'][5:12]}-{d['cnic'][12]}"
        if d["cnic"]:
            ck = d["cnic"].replace("-", "")
            if ck in seen_cnics:
                errors.append(f"CNIC {d['cnic']} appears twice in the file")
            seen_cnics.add(ck)
        if d["phone"].isdigit() and len(d["phone"]) == 10 and d["phone"].startswith("3"):
            d["phone"] = "0" + d["phone"]                      # Excel dropped the leading 0
        st = d["status"].lower()
        d["status"] = "" if not st else ("inactive" if st in ("inactive", "left", "no", "0", "resigned", "terminated") else "active")
        g_ = d["gender"].lower()
        d["gender"] = {"m": "Male", "male": "Male", "f": "Female", "female": "Female"}.get(g_, d["gender"].title())
        if d["department"] and d["department"].lower() not in depts:
            new_depts.add(d["department"])
            warns.append(f"New department '{d['department']}' will be created")
        # match existing
        existing = None
        if d["emp_code"] and d["emp_code"].lower() in by_code:
            existing = by_code[d["emp_code"].lower()]
        elif d["cnic"] and d["cnic"].replace("-", "") in by_cnic:
            existing = by_cnic[d["cnic"].replace("-", "")]
        if d["emp_code"]:
            if d["emp_code"].lower() in seen_codes:
                errors.append(f"Employee code {d['emp_code']} appears twice in the file")
            seen_codes.add(d["emp_code"].lower())
        if existing and not update_existing:
            action = "skip"
        else:
            action = "update" if existing else "new"
        d.setdefault("username", "")
        d.setdefault("password", "")
        if d["username"]:
            u = d["username"].lower()
            if action == "new" and (u in taken_users or u in seen_users):
                warns.append(f"Username '{d['username']}' already taken – login not created")
                d["username"] = ""
            elif action == "new" and len(d["password"]) < 6:
                warns.append("Password missing or under 6 characters – login not created")
                d["username"] = ""
            elif action == "update":
                d["username"] = ""                               # never touch logins on update
            seen_users.add(u)
        rows.append({"line": n, "data": d, "errors": errors, "warns": warns,
                     "action": "error" if errors else action, "existing": existing})
    # check every Reporting Manager points to someone (already in the system or in this file)
    file_codes = {r["data"]["emp_code"].lower() for r in rows if r["data"]["emp_code"]}
    file_names = Counter(r["data"]["name"].lower() for r in rows if r["data"]["name"])
    for r in rows:
        m = r["data"].get("manager", "").strip()
        if not m:
            continue
        ml = m.lower()
        if ml in by_code or ml in file_codes:
            pass
        elif emp_names.get(ml, 0) + file_names.get(ml, 0) == 1:
            pass
        elif emp_names.get(ml, 0) + file_names.get(ml, 0) > 1:
            r["warns"].append(f"Manager '{m}': more than one employee has this name – use the employee code instead; left empty")
            r["data"]["manager"] = ""
        else:
            r["warns"].append(f"Manager '{m}' not found – left empty")
            r["data"]["manager"] = ""
        if r["data"]["manager"] and (ml == r["data"]["emp_code"].lower() or ml == r["data"]["name"].lower()):
            r["warns"].append("An employee can't be their own manager – left empty")
            r["data"]["manager"] = ""
    return rows, sorted(new_depts), [IMPORT_LABELS[v] for v in mapping.values()]


@app.route("/employees/import", methods=["GET", "POST"])
@perm_required("employees.manage")
def employee_import():
    if request.method == "POST":
        fs = request.files.get("file")
        if not fs or not fs.filename:
            flash("Please choose a file to upload.", "error")
            return redirect(url_for("employee_import"))
        update = bool(request.form.get("update_existing"))
        try:
            rows, new_depts, found_cols = analyse_import(read_upload(fs), update)
        except Exception as e:                                   # noqa: BLE001
            flash(f"Couldn't read the file: {e}", "error")
            return redirect(url_for("employee_import"))
        if not rows:
            flash("The file has no employee rows under the heading row.", "error")
            return redirect(url_for("employee_import"))
        os.makedirs(IMPORT_DIR, exist_ok=True)
        token = secrets.token_hex(12)
        import json
        with open(os.path.join(IMPORT_DIR, token + ".json"), "w", encoding="utf-8") as f:
            json.dump({"rows": rows, "new_depts": new_depts, "update": update}, f)
        counts = Counter(r["action"] for r in rows)
        return render_template("employee_import.html", step="preview", rows=rows, counts=counts, token=token,
                               new_depts=new_depts, found_cols=found_cols, filename=fs.filename)
    src = request.args.get("source", "all")
    sql = """SELECT e.*, d.name dept, (SELECT username FROM users u WHERE u.employee_id=e.id) uname
               FROM employees e LEFT JOIN departments d ON d.id=e.department_id"""
    args = []
    if src in ("manual", "import"):
        sql += " WHERE COALESCE(e.source,'manual')=?"
        args.append(src)
    staff = q(sql + " ORDER BY e.id DESC", args)
    src_counts = {r["s"]: r["c"] for r in q("SELECT COALESCE(source,'manual') s, COUNT(*) c FROM employees GROUP BY s")}
    return render_template("employee_import.html", step="upload", labels=IMPORT_LABELS, cols=IMPORT_COLS,
                           staff=staff, src=src, src_counts=src_counts)


@app.post("/employees/import/confirm")
@perm_required("employees.manage")
def employee_import_confirm():
    import json
    token = request.form.get("token", "")
    path = os.path.join(IMPORT_DIR, token + ".json")
    if not token.isalnum() or not os.path.exists(path):
        flash("This import has expired. Please upload the file again.", "error")
        return redirect(url_for("employee_import"))
    with open(path, encoding="utf-8") as f:
        job = json.load(f)
    os.remove(path)
    con = db()
    try:
        for dn in job["new_depts"]:
            con.execute("INSERT OR IGNORE INTO departments(name) VALUES(?)", (dn,))
        depts = {r["name"].lower(): r["id"] for r in con.execute("SELECT id, name FROM departments")}
        added = updated = logins = 0
        mgr_links = []                                           # (employee id, manager text)
        nums = [int(r[0][4:]) for r in con.execute("SELECT emp_code FROM employees WHERE emp_code LIKE 'EMP-%'")
                if r[0][4:].isdigit()]
        nxt = max(nums) + 1 if nums else 1
        for r in job["rows"]:
            if r["action"] not in ("new", "update"):
                continue
            d = r["data"]
            vals = {k: d[k] for k in EMP_FIELDS if k in d}
            vals["department_id"] = depts.get(d["department"].lower()) if d["department"] else None
            if r["action"] == "update":
                # only overwrite fields that have a value in the file
                sets = {k: v for k, v in vals.items() if v not in ("", None)}
                if d["emp_code"]:
                    sets["emp_code"] = d["emp_code"]
                if sets:
                    con.execute(f"UPDATE employees SET {', '.join(k + '=?' for k in sets)} WHERE id=?",
                                list(sets.values()) + [r["existing"]])
                updated += 1
                if d.get("manager"):
                    mgr_links.append((r["existing"], d["manager"]))
            else:
                code = d["emp_code"]
                if not code:
                    while True:
                        code = f"EMP-{nxt:04d}"
                        nxt += 1
                        if not con.execute("SELECT 1 FROM employees WHERE emp_code=?", (code,)).fetchone():
                            break
                elif con.execute("SELECT 1 FROM employees WHERE emp_code=?", (code,)).fetchone():
                    continue
                vals["status"] = vals.get("status") or "active"
                vals["basic_salary"] = vals.get("basic_salary") or 0
                vals["allowances"] = vals.get("allowances") or 0
                cols = ["emp_code", "source"] + list(vals.keys())
                cur = con.execute(f"INSERT INTO employees({','.join(cols)}) VALUES({','.join('?' * len(cols))})",
                                  [code, "import"] + list(vals.values()))
                added += 1
                if d.get("manager"):
                    mgr_links.append((cur.lastrowid, d["manager"]))
                if d.get("username") and not con.execute("SELECT 1 FROM users WHERE username=?", (d["username"],)).fetchone():
                    con.execute("INSERT INTO users(username,password_hash,role,employee_id,must_change) VALUES(?,?,'employee',?,1)",
                                (d["username"], generate_password_hash(d["password"]), cur.lastrowid))
                    logins += 1
        linked = 0
        if mgr_links:
            codes = {r[1].lower(): r[0] for r in con.execute("SELECT id, emp_code FROM employees")}
            names = {}
            for r_ in con.execute("SELECT id, name FROM employees"):
                names.setdefault(r_[1].lower(), []).append(r_[0])
            for eid_, m in mgr_links:
                ml = m.strip().lower()
                mid = codes.get(ml) or (names[ml][0] if len(names.get(ml, [])) == 1 else None)
                if mid and mid != eid_:
                    # refuse circles (A -> B -> A)
                    cur_, seen_ = mid, set()
                    while cur_ and cur_ not in seen_ and cur_ != eid_:
                        seen_.add(cur_)
                        row_ = con.execute("SELECT manager_id FROM employees WHERE id=?", (cur_,)).fetchone()
                        cur_ = row_[0] if row_ else None
                    if cur_ != eid_:
                        con.execute("UPDATE employees SET manager_id=? WHERE id=?", (mid, eid_))
                        linked += 1
        con.commit()
    except Exception as e:                                       # noqa: BLE001
        con.rollback()
        flash(f"Import failed and nothing was changed: {e}", "error")
        return redirect(url_for("employee_import"))
    msg = f"Import complete: {added} employee(s) added, {updated} updated"
    if linked:
        msg += f", {linked} reporting manager(s) linked"
    if logins:
        msg += f", {logins} login(s) created"
    if job["new_depts"]:
        msg += f", {len(job['new_depts'])} new department(s)"
    flash(msg + ".", "ok")
    return redirect(url_for("employees", status="all"))


@app.route("/employees/import/template")
@perm_required("employees.manage")
def employee_import_template():
    sample = [["EMP-0101", "Ahmed Raza", "Muhammad Raza", "35202-1234567-1", "Male", "15/03/1992",
               "ahmed@company.com", "03001234567", "House 12, Street 4, Lahore", "Accounts", "Accountant", "",
               "01/02/2024", "Permanent", 60000, 10000, "PK00HABB0000001234567890", "Raza 03211234567", "Active"],
              ["", "Fatima Noor", "Noor Ahmed", "35202-7654321-2", "Female", "22/08/1995", "", "03111234567",
               "", "Human Resources", "HR Officer", "EMP-0101", "10/06/2025", "Contract", 45000, 5000, "", "", "Active"]]
    head = [IMPORT_LABELS[c] for c in IMPORT_COLS]
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Employees"
        ws.append(head)
        for r in sample:
            ws.append(r)
        fill = PatternFill("solid", fgColor="0F2A44")
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = fill
            c.alignment = Alignment(vertical="center")
        for i, h in enumerate(head, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(14, len(h) + 4)
            for row in ws.iter_rows(min_row=2, max_row=500, min_col=i, max_col=i):
                for c in row:
                    if IMPORT_COLS[i - 1] in ("cnic", "phone", "bank_account", "emergency_contact", "emp_code", "manager",
                                              "dob", "join_date"):
                        c.number_format = "@"                # keep as text (leading zeros)
        ws.freeze_panes = "A2"
        info = wb.create_sheet("Instructions")
        for line in ["How to fill this template",
                     "• Only 'Full Name' is required. Leave other cells empty if unknown.",
                     "• Employee Code: leave empty to auto-generate (EMP-0001, EMP-0002 …).",
                     "• Dates: DD/MM/YYYY, e.g. 15/03/1992.",
                     "• Department: written exactly; new departments are created automatically.",
                     "• Status: Active or Inactive.",
                     "• Reporting Manager: the manager's Employee Code (best, e.g. EMP-0101) or exact full name. The manager can be an existing employee or someone else in this same file.",
                     "• Delete the two example rows before importing.",
                     "• Existing employees are matched by Employee Code or CNIC."]:
            info.append([line])
        info.column_dimensions["A"].width = 110
        info["A1"].font = Font(bold=True, size=13)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf, as_attachment=True, download_name="employee-import-template.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except ImportError:
        return csv_response("employee-import-template.csv", head, sample)


# --------------------------------------------------------------------------- departments
DEPT_SQL = """SELECT d.*, h.name head_name, h.emp_code head_code,
                COUNT(e.id) total, COALESCE(SUM(e.status='active'),0) active,
                COALESCE(SUM(CASE WHEN e.status='active' THEN e.basic_salary+e.allowances END),0) bill
              FROM departments d LEFT JOIN employees e ON e.department_id=d.id
              LEFT JOIN employees h ON h.id=d.head_id
              GROUP BY d.id ORDER BY d.name"""


def _dept_form():
    name = request.form.get("name", "").strip()
    code = request.form.get("code", "").strip().upper() or None
    head = request.form.get("head_id", "")
    head = int(head) if head.isdigit() else None
    desc = request.form.get("description", "").strip() or None
    return name, code, head, desc


@app.route("/departments")
@perm_required("departments.view")
def departments():
    rows = q(DEPT_SQL)
    today = date.today().isoformat()
    present = {r["department_id"]: r["c"] for r in q(
        """SELECT e.department_id, COUNT(*) c FROM attendance a JOIN employees e ON e.id=a.employee_id
           WHERE a.date=? AND a.status IN ('present','late','half_day') AND e.status='active'
           GROUP BY e.department_id""", (today,))}
    unassigned = q("SELECT COUNT(*) c FROM employees WHERE status='active' AND department_id IS NULL", one=True)["c"]
    return render_template("departments.html", rows=rows, present=present, unassigned=unassigned,
                           maxc=max([r["active"] for r in rows] or [1]) or 1)


@app.route("/departments/new", methods=["GET", "POST"])
@perm_required("departments.manage")
def department_new():
    if request.method == "POST":
        name, code, head, desc = _dept_form()
        if not name:
            flash("Department name is required.", "error")
        else:
            try:
                x("INSERT INTO departments(name, code, head_id, description) VALUES(?,?,?,?)", (name, code, head, desc))
                flash(f"Department '{name}' added.", "ok")
                if request.form.get("again"):
                    return redirect(url_for("department_new"))
                return redirect(url_for("departments"))
            except sqlite3.IntegrityError:
                flash(f"A department named '{name}' already exists.", "error")
        return render_template("department_form.html", d=request.form, emps=_active_emps())
    return render_template("department_form.html", d={}, emps=_active_emps())


def _active_emps():
    return q("SELECT id, name, emp_code, designation FROM employees WHERE status='active' ORDER BY name")


@app.route("/departments/manage", methods=["GET", "POST"])
@perm_required("departments.manage")
def departments_manage():
    if request.method == "POST":
        action = request.form.get("action")
        if action == "add":
            name, code, head, desc = _dept_form()
            if not name:
                flash("Department name is required.", "error")
            else:
                try:
                    x("INSERT INTO departments(name, code, head_id, description) VALUES(?,?,?,?)", (name, code, head, desc))
                    flash(f"Department '{name}' added.", "ok")
                except sqlite3.IntegrityError:
                    flash(f"A department named '{name}' already exists.", "error")
        elif action == "save":
            ids = [int(i) for i in request.form.getlist("changed") if i.isdigit()]
            saved, errs = 0, []
            for did in ids:
                name = request.form.get(f"name_{did}", "").strip()
                if not name:
                    errs.append("A department name can't be empty.")
                    continue
                head = request.form.get(f"head_{did}", "")
                try:
                    db().execute("UPDATE departments SET name=?, code=?, head_id=?, description=? WHERE id=?",
                                 (name, request.form.get(f"code_{did}", "").strip().upper() or None,
                                  int(head) if head.isdigit() else None,
                                  request.form.get(f"desc_{did}", "").strip() or None, did))
                    saved += 1
                except sqlite3.IntegrityError:
                    errs.append(f"Another department is already named '{name}'.")
            db().commit()
            if saved:
                flash(f"Saved {saved} department(s).", "ok")
            for e_ in errs:
                flash(e_, "error")
        elif action == "delete":
            did = request.form.get("id", "")
            move_to = request.form.get("move_to", "")
            d = q("SELECT * FROM departments WHERE id=?", (did,), one=True)
            if d:
                n = q("SELECT COUNT(*) c FROM employees WHERE department_id=?", (did,), one=True)["c"]
                if n and not move_to:
                    flash(f"'{d['name']}' has {n} employee(s). Choose where to move them first.", "error")
                else:
                    if n:
                        target = int(move_to) if move_to.isdigit() and move_to != did else None
                        db().execute("UPDATE employees SET department_id=? WHERE department_id=?", (target, did))
                    to_bin("department", did)
                    db().execute("DELETE FROM departments WHERE id=?", (did,))
                    db().commit()
                    flash(f"Department '{d['name']}' deleted" + (f" and {n} employee(s) moved." if n else "."), "ok")
        elif action == "move":
            src, dst = request.form.get("src", ""), request.form.get("dst", "")
            if src.isdigit() and (dst.isdigit() or dst == "none") and src != dst:
                n = db().execute("UPDATE employees SET department_id=? WHERE department_id=?",
                                 (int(dst) if dst.isdigit() else None, int(src))).rowcount
                db().commit()
                flash(f"Moved {n} employee(s).", "ok")
            else:
                flash("Choose two different departments.", "error")
        return redirect(url_for("departments_manage"))
    return render_template("departments_manage.html", rows=q(DEPT_SQL), emps=_active_emps())


@app.post("/departments/<int:did>/delete")
@perm_required("departments.manage")
def department_delete(did):
    if q("SELECT 1 FROM employees WHERE department_id=?", (did,), one=True):
        flash("Move the employees in this department first, then delete it.", "error")
    else:
        to_bin("department", did)
        x("DELETE FROM departments WHERE id=?", (did,))
        flash("Department deleted.", "ok")
    return redirect(url_for("departments_manage"))


# --------------------------------------------------------------------------- attendance
@app.route("/attendance", methods=["GET", "POST"])
@perm_required("attendance.manage")
def attendance():
    d = parse_date(request.values.get("date"), date.today())
    ds = d.isoformat()
    emps = q("""SELECT e.id, e.emp_code, e.name, e.designation, d.name dept FROM employees e
                LEFT JOIN departments d ON d.id=e.department_id
                WHERE e.status='active' AND (e.join_date IS NULL OR e.join_date='' OR e.join_date<=?)
                ORDER BY d.name, e.name""", (ds,))
    if request.method == "POST":
        saved = 0
        for e in emps:
            st = request.form.get(f"status_{e['id']}", "")
            if st == "":
                db().execute("DELETE FROM attendance WHERE employee_id=? AND date=?", (e["id"], ds))
                continue
            if st not in ATT_STATUSES:
                continue
            cin = request.form.get(f"in_{e['id']}", "").strip() or None
            cout = request.form.get(f"out_{e['id']}", "").strip() or None
            note = request.form.get(f"note_{e['id']}", "").strip() or None
            if st in ("absent", "leave"):
                cin = cout = None
            db().execute("""INSERT INTO attendance(employee_id,date,status,check_in,check_out,note) VALUES(?,?,?,?,?,?)
                            ON CONFLICT(employee_id,date) DO UPDATE SET status=excluded.status,
                            check_in=excluded.check_in, check_out=excluded.check_out, note=excluded.note""",
                         (e["id"], ds, st, cin, cout, note))
            saved += 1
        db().commit()
        flash(f"Attendance saved for {nicedate(ds)} ({saved} marked).", "ok")
        return redirect(url_for("attendance", date=ds))
    att = {r["employee_id"]: r for r in q("SELECT * FROM attendance WHERE date=?", (ds,))}
    counts = Counter(r["status"] for r in att.values())
    return render_template("attendance_daily.html", d=d, emps=emps, att=att, counts=counts,
                           prev=(d - timedelta(days=1)).isoformat(), next=(d + timedelta(days=1)).isoformat(),
                           workday=is_workday(d), weekday=WEEKDAYS[d.weekday()], holiday=holiday_name(d))


@app.post("/attendance/mark")
@perm_required("attendance.manage")
def attendance_mark():
    """One-click marking from the Daily Attendance page (Present / Leave / Absent)."""
    eid = request.form.get("eid", type=int)
    ds = parse_date(request.form.get("date"), date.today()).isoformat()
    st = request.form.get("status", "")
    if st not in ("present", "leave", "absent") or not q("SELECT 1 FROM employees WHERE id=?", (eid,), one=True):
        return {"ok": False, "error": "Invalid request"}, 400
    if st == "present":
        db().execute("""INSERT INTO attendance(employee_id,date,status) VALUES(?,?,'present')
                        ON CONFLICT(employee_id,date) DO UPDATE SET status='present'""", (eid, ds))
    else:
        db().execute("""INSERT INTO attendance(employee_id,date,status) VALUES(?,?,?)
                        ON CONFLICT(employee_id,date) DO UPDATE SET status=excluded.status, check_in=NULL, check_out=NULL""",
                     (eid, ds, st))
    db().commit()
    counts = Counter(r["status"] for r in q("SELECT status FROM attendance WHERE date=?", (ds,)))
    return {"ok": True, "status": st, "counts": dict(counts)}


@app.route("/attendance/report")
@perm_required("attendance.view")
def attendance_report():
    y, m = parse_month(request.args.get("month"))
    month = f"{y:04d}-{m:02d}"
    dates = month_dates(y, m)
    dept = request.args.get("dept", "")
    search = request.args.get("q", "").strip()
    emp_id = request.args.get("emp", type=int)
    sql = """SELECT e.id, e.emp_code, e.name, e.designation, d.name dept FROM employees e LEFT JOIN departments d
             ON d.id=e.department_id WHERE e.status='active'"""
    args = []
    if dept:
        sql += " AND e.department_id=?"
        args.append(dept)
    if search:
        sql += " AND (e.name LIKE ? OR e.emp_code LIKE ?)"
        args += [f"%{search}%", f"%{search}%"]
    emps = q(sql + " ORDER BY d.name, e.name", args)
    if search and len(emps) == 1 and not emp_id:
        emp_id = emps[0]["id"]
    grid, times = {}, {}
    for r in q("SELECT employee_id, date, status, check_in, check_out FROM attendance WHERE substr(date,1,7)=?", (month,)):
        grid[(r["employee_id"], r["date"])] = r["status"]
        times[(r["employee_id"], r["date"])] = (r["check_in"] or "", r["check_out"] or "")
    off = weekly_off()
    offd = {d for d in dates if not is_workday(d)}
    hol = holiday_map()
    workdays_so_far = sum(1 for d in dates if d not in offd and d <= date.today())
    rows = []
    team = Counter()
    for e in emps:
        cells = [grid.get((e["id"], d.isoformat())) for d in dates]
        c = Counter(s for s in cells if s)
        team.update(c)
        attended = c["present"] + c["late"] + c["half_day"] * 0.5
        pct = round(attended * 100 / workdays_so_far) if workdays_so_far else 0
        tips = [times.get((e["id"], d.isoformat()), ("", "")) for d in dates]
        rows.append({"e": e, "cells": cells, "c": c, "pct": min(pct, 100), "tips": tips,
                     "initials": "".join(w[0] for w in e["name"].split()[:2]).upper()})
    team_attended = team["present"] + team["late"] + team["half_day"] * 0.5
    team_pct = round(team_attended * 100 / (workdays_so_far * len(emps))) if workdays_so_far and emps else 0
    if request.args.get("export") == "csv":
        head = ["Code", "Name", "Department"] + [str(d.day) for d in dates] + ["Present", "Late", "Half", "Absent", "Leave"]
        body = [[r["e"]["emp_code"], r["e"]["name"], r["e"]["dept"] or ""]
                + [ATT_CODES.get(s, "") if s else ("HOLIDAY" if hol.get(d.isoformat()) else ("OFF" if d in offd else "")) for s, d in zip(r["cells"], dates)]
                + [r["c"]["present"], r["c"]["late"], r["c"]["half_day"], r["c"]["absent"], r["c"]["leave"]]
                for r in rows]
        return csv_response(f"attendance-{month}.csv", head, body)
    detail = None
    if emp_id:
        e = q("""SELECT e.id, e.emp_code, e.name, e.designation, d.name dept FROM employees e
                 LEFT JOIN departments d ON d.id=e.department_id WHERE e.id=?""", (emp_id,), one=True)
        if e:
            recs = {r["date"]: r for r in q("SELECT * FROM attendance WHERE employee_id=? AND substr(date,1,7)=?",
                                            (emp_id, month))}
            days = [{"d": d, "r": recs.get(d.isoformat()), "off": d in offd, "hol": hol.get(d.isoformat())} for d in dates]
            c = Counter(r["status"] for r in recs.values())
            workdays = sum(1 for x_ in days if not x_["off"] and x_["d"] <= date.today())
            attended = c["present"] + c["late"] + c["half_day"] * 0.5
            detail = {"e": e, "days": days, "c": c, "workdays": workdays,
                      "pct": round(attended * 100 / workdays) if workdays else 0,
                      "unmarked": sum(1 for x_ in days if not x_["off"] and not x_["r"] and x_["d"] <= date.today())}
    return render_template("attendance_report.html", month=month, dates=dates, rows=rows, off=off, offd=offd, hol=hol,
                           prev=shift_month(y, m, -1), next=shift_month(y, m, 1), depts=all_depts(), dept=dept,
                           search=search, detail=detail, today=date.today(), team=team,
                           team_pct=min(team_pct, 100), workdays_so_far=workdays_so_far)


@app.route("/my/attendance")
@login_required
def my_attendance():
    emp_id = g.user["employee_id"] or abort(403)
    y, m = parse_month(request.args.get("month"))
    month = f"{y:04d}-{m:02d}"
    recs = {r["date"]: r for r in q("SELECT * FROM attendance WHERE employee_id=? AND substr(date,1,7)=?",
                                    (emp_id, month))}
    days = [(d, recs.get(d.isoformat()), is_workday(d)) for d in month_dates(y, m) if d <= date.today()]
    c = Counter(r["status"] for r in recs.values())
    return render_template("my_attendance.html", days=list(reversed(days)), c=c, month=month,
                           prev=shift_month(y, m, -1), next=shift_month(y, m, 1))


# --------------------------------------------------------------------------- attendance: manage
@app.route("/attendance/manage", methods=["GET", "POST"])
@perm_required("attendance.manage")
def attendance_manage():
    if request.method == "POST":
        action = request.form.get("action")
        back = {k: request.form.get(k, "") for k in ("frm", "to", "emp", "dept", "st")}
        if action == "save":
            ids = [int(i) for i in request.form.getlist("changed") if i.isdigit()]
            dels = {int(i) for i in request.form.getlist("del") if i.isdigit()}
            saved = deleted = 0
            for rid in dels:
                db().execute("DELETE FROM attendance WHERE id=?", (rid,))
                deleted += 1
            for rid in ids:
                if rid in dels:
                    continue
                st = request.form.get(f"st_{rid}", "")
                if st not in ATT_STATUSES:
                    continue
                cin = request.form.get(f"in_{rid}", "").strip() or None
                cout = request.form.get(f"out_{rid}", "").strip() or None
                if st in ("absent", "leave"):
                    cin = cout = None
                db().execute("UPDATE attendance SET status=?, check_in=?, check_out=?, note=? WHERE id=?",
                             (st, cin, cout, request.form.get(f"note_{rid}", "").strip() or None, rid))
                saved += 1
            db().commit()
            flash(f"Saved {saved} record(s)" + (f", deleted {deleted}" if deleted else "") + ".", "ok")
        elif action == "mark":
            eids = [int(i) for i in request.form.getlist("m_emp") if i.isdigit()]
            if request.form.get("m_all"):
                eids = [r["id"] for r in q("SELECT id FROM employees WHERE status='active'")]
            a, b = parse_date(request.form.get("m_from")), parse_date(request.form.get("m_to"))
            st = request.form.get("m_status")
            overwrite = bool(request.form.get("m_over"))
            if not eids or not a or not b or b < a or st not in ATT_STATUSES + ["clear"]:
                flash("Choose employee(s), a valid date range and a status.", "error")
            elif (b - a).days > 92:
                flash("Please mark at most 3 months at a time.", "error")
            else:
                n = 0
                note = request.form.get("m_note", "").strip() or None
                for eid in eids:
                    for d in daterange(a, b):
                        if not is_workday(d) and not request.form.get("m_off"):
                            continue
                        ds = d.isoformat()
                        if st == "clear":
                            n += db().execute("DELETE FROM attendance WHERE employee_id=? AND date=?", (eid, ds)).rowcount
                            continue
                        ex = db().execute("SELECT id FROM attendance WHERE employee_id=? AND date=?", (eid, ds)).fetchone()
                        if ex and not overwrite:
                            continue
                        db().execute("""INSERT INTO attendance(employee_id,date,status,note) VALUES(?,?,?,?)
                                        ON CONFLICT(employee_id,date) DO UPDATE SET status=excluded.status,
                                        note=excluded.note, check_in=NULL, check_out=NULL""", (eid, ds, st, note))
                        n += 1
                db().commit()
                flash(f"{'Cleared' if st == 'clear' else 'Marked'} {n} day-record(s).", "ok")
            back.update(frm=request.form.get("m_from", back["frm"]), to=request.form.get("m_to", back["to"]))
        return redirect(url_for("attendance_manage", **{k: v for k, v in back.items() if v}))
    today = date.today()
    frm = parse_date(request.args.get("frm"), today.replace(day=1))
    to = parse_date(request.args.get("to"), today)
    if to < frm:
        frm, to = to, frm
    sql = """SELECT a.*, e.name, e.emp_code, d.name dept FROM attendance a JOIN employees e ON e.id=a.employee_id
             LEFT JOIN departments d ON d.id=e.department_id WHERE a.date BETWEEN ? AND ?"""
    args = [frm.isoformat(), to.isoformat()]
    emp, dept, st = request.args.get("emp", ""), request.args.get("dept", ""), request.args.get("st", "")
    if emp.isdigit():
        sql += " AND a.employee_id=?"
        args.append(int(emp))
    if dept.isdigit():
        sql += " AND e.department_id=?"
        args.append(int(dept))
    if st in ATT_STATUSES:
        sql += " AND a.status=?"
        args.append(st)
    rows = q(sql + " ORDER BY a.date DESC, e.name LIMIT 2000", args)
    counts = Counter(r["status"] for r in rows)
    emps = q("SELECT id, name, emp_code, status FROM employees ORDER BY status, name")
    return render_template("attendance_manage.html", rows=rows, counts=counts, frm=frm, to=to, emps=emps,
                           depts=all_depts(), emp=emp, dept=dept, st=st, limited=len(rows) >= 2000)


# --------------------------------------------------------------------------- attendance: import
ATT_ALIASES = {
    "emp": ["emp_code", "employee_code", "code", "emp_id", "employee_id", "emp_no", "employee_no", "staff_id", "id",
            "ac_no", "acno", "no", "user_id", "userid", "enroll_no", "enrollno", "badge", "badge_no", "person_id",
            "card_no", "pin"],
    "name": ["name", "full_name", "employee_name", "emp_name", "staff_name", "employee"],
    "cnic": ["cnic", "nic", "cnic_no", "national_id"],
    "date": ["date", "att_date", "attendance_date", "day", "work_date"],
    "datetime": ["time", "date_time", "datetime", "check_time", "checktime", "punch_time", "timestamp", "log_time",
                 "punch", "date_and_time"],
    "status": ["status", "attendance", "att", "state", "remarks_status", "p_a"],
    "in": ["in", "check_in", "time_in", "in_time", "checkin", "clock_in", "first_in", "arrival"],
    "out": ["out", "check_out", "time_out", "out_time", "checkout", "clock_out", "last_out", "departure"],
    "note": ["note", "notes", "remark", "remarks", "comment", "reason"],
}
_ATT_ALIAS = {a: k for k, v in ATT_ALIASES.items() for a in v}
STATUS_WORDS = {
    "p": "present", "present": "present", "1": "present", "y": "present", "yes": "present", "✓": "present", "✔": "present",
    "a": "absent", "absent": "absent", "0": "absent", "n": "absent", "x": "absent", "ab": "absent",
    "l": "late", "late": "late", "lt": "late",
    "h": "half_day", "hd": "half_day", "half": "half_day", "half_day": "half_day", "half day": "half_day", "½": "half_day",
    "lv": "leave", "leave": "leave", "cl": "leave", "sl": "leave", "al": "leave", "el": "leave", "ml": "leave", "on leave": "leave",
}
SKIP_WORDS = {"", "-", "off", "wo", "w", "holiday", "hol", "sun", "sunday", "weekend", "ph", "gh", "--"}


def _parse_time(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.strftime("%H:%M")
    if hasattr(v, "hour") and hasattr(v, "minute") and not isinstance(v, date):
        return f"{v.hour:02d}:{v.minute:02d}"
    if isinstance(v, float) and 0 <= v < 1:               # Excel time fraction
        m = round(v * 24 * 60)
        return f"{m // 60:02d}:{m % 60:02d}"
    s = str(v).strip().upper().replace(".", ":")
    for fmt in ("%H:%M", "%H:%M:%S", "%I:%M %p", "%I:%M:%S %p", "%I:%M%p", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M",
                "%d/%m/%Y %H:%M", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %I:%M %p", "%d-%m-%Y %H:%M", "%m/%d/%Y %I:%M:%S %p"):
        try:
            return datetime.strptime(s, fmt).strftime("%H:%M")
        except ValueError:
            pass
    return None


def _split_datetime(v):
    """Return (iso_date or '', 'HH:MM' or None) from a combined date-time cell."""
    if isinstance(v, datetime):
        return v.date().isoformat(), v.strftime("%H:%M")
    if isinstance(v, (int, float)) and v > 20000:          # Excel serial with fraction
        dt = datetime(1899, 12, 30) + timedelta(days=float(v))
        return dt.date().isoformat(), dt.strftime("%H:%M")
    s = str(v or "").strip()
    if not s:
        return "", None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y %I:%M %p",
                "%d/%m/%Y %I:%M:%S %p", "%m/%d/%Y %I:%M:%S %p", "%m/%d/%Y %I:%M %p", "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M",
                "%Y/%m/%d %H:%M:%S", "%Y/%m/%d %H:%M", "%d.%m.%Y %H:%M"):
        try:
            dt = datetime.strptime(s.upper() if "M" in s.upper()[-2:] else s, fmt)
            return dt.date().isoformat(), dt.strftime("%H:%M")
        except ValueError:
            pass
    d_, err = _parse_any_date(s)
    if d_:
        return d_, None
    return "", _parse_time(s)


def _status_from_time(cin):
    if not cin:
        return "present"
    s = get_settings()
    start = parse_time(s.get("office_start")) or datetime.strptime("09:00", "%H:%M").time()
    limit = (datetime.combine(date.today(), start) + timedelta(minutes=int(to_float(s.get("grace_minutes"), 0)))).time()
    return "late" if datetime.strptime(cin, "%H:%M").time() > limit else "present"


def _employee_matcher():
    emps = q("SELECT id, emp_code, name, cnic, status, join_date FROM employees")
    by_code = {e["emp_code"].lower(): e for e in emps}
    by_num = {}
    for e in emps:                                          # "EMP-0012" also matches "12" (biometric numbers)
        digits = "".join(ch for ch in e["emp_code"] if ch.isdigit())
        if digits:
            by_num.setdefault(str(int(digits)), []).append(e)
    by_cnic = {(e["cnic"] or "").replace("-", ""): e for e in emps if e["cnic"]}
    by_name = {}
    for e in emps:
        by_name.setdefault(e["name"].strip().lower(), []).append(e)

    def match(code, name, cnic):
        code = (code or "").strip()
        if code:
            if code.lower() in by_code:
                return by_code[code.lower()], None
            if code.isdigit() and str(int(code)) in by_num and len(by_num[str(int(code))]) == 1:
                return by_num[str(int(code))][0], None
        if cnic and cnic.replace("-", "") in by_cnic:
            return by_cnic[cnic.replace("-", "")], None
        if name:
            hits = by_name.get(name.strip().lower(), [])
            if len(hits) == 1:
                return hits[0], None
            if len(hits) > 1:
                return None, f"more than one employee named '{name}', add the Employee Code"
        return None, f"employee not found ({code or name or cnic or 'blank'})"
    return match


def analyse_att_import(raw, grid_month, overwrite):
    match = _employee_matcher()
    # locate header row
    hdr_idx, mapping, day_cols = None, {}, {}
    for i, r in enumerate(raw[:12]):
        m, days = {}, {}
        for j, h in enumerate(r):
            hn = _norm_header(_cell_text(h))
            if hn.isdigit() and 1 <= int(hn) <= 31:
                days[j] = int(hn)
                continue
            key = _ATT_ALIAS.get(hn)
            if key and key not in m.values():
                m[j] = key
        has_emp = any(v in ("emp", "name", "cnic") for v in m.values())
        if has_emp and (len(days) >= 7 or any(v in ("date", "datetime") for v in m.values())):
            hdr_idx, mapping, day_cols = i, m, (days if len(days) >= 7 else {})
            break
    if hdr_idx is None:
        raise ValueError("Couldn't find the heading row. The file needs an Employee Code or Name column, plus either "
                         "a Date column (list format) or day columns 1, 2, 3 … 31 (monthly sheet format).")
    existing = {(r["employee_id"], r["date"]): r for r in q("SELECT employee_id, date, status, check_in, check_out FROM attendance")}
    recs, problems = {}, []
    fmt = "grid" if day_cols else "list"
    if fmt == "grid":
        y, mth = parse_month(grid_month)
        last_day = calendar.monthrange(y, mth)[1]
    for n, r in enumerate(raw[hdr_idx + 1:], start=hdr_idx + 2):
        cell = {k: (r[j] if j < len(r) else None) for j, k in mapping.items()}
        code, name, cnic = _cell_text(cell.get("emp")), _cell_text(cell.get("name")), _cell_text(cell.get("cnic"))
        if not (code or name or cnic):
            continue
        e, err = match(code, name, cnic)
        if not e:
            problems.append({"line": n, "who": code or name or cnic, "msg": err})
            continue
        if fmt == "grid":
            for j, day in day_cols.items():
                if day > last_day:
                    continue
                v = _cell_text(r[j] if j < len(r) else None).strip()
                key = v.lower()
                if key in SKIP_WORDS:
                    continue
                st = STATUS_WORDS.get(key)
                if not st:
                    problems.append({"line": n, "who": e["name"], "msg": f"day {day}: '{v}' is not a status code"})
                    continue
                ds = date(y, mth, day).isoformat()
                recs[(e["id"], ds)] = {"eid": e["id"], "code": e["emp_code"], "name": e["name"], "date": ds,
                                       "status": st, "in": None, "out": None, "note": None, "line": n}
            continue
        # list format
        ds, tm = "", None
        if "date" in cell and cell.get("date") not in (None, ""):
            dv = cell["date"]
            ds, tm0 = _split_datetime(dv) if isinstance(dv, datetime) else (_parse_any_date(dv)[0], None)
            tm = tm0 if tm0 and tm0 != "00:00" else None
        if "datetime" in cell and cell.get("datetime") not in (None, ""):
            d2, t2 = _split_datetime(cell["datetime"])
            ds = ds or d2
            tm = t2 or _parse_time(cell["datetime"])
        if not ds:
            problems.append({"line": n, "who": e["name"], "msg": "date missing or not understood (use DD/MM/YYYY)"})
            continue
        st_raw = _cell_text(cell.get("status")).strip().lower()
        st = STATUS_WORDS.get(st_raw) if st_raw else None
        if st_raw and not st and st_raw not in SKIP_WORDS and st_raw not in ("c/in", "c/out", "check in", "check out",
                                                                              "checkin", "checkout", "in", "out", "0", "1"):
            problems.append({"line": n, "who": e["name"], "msg": f"status '{st_raw}' not understood"})
            continue
        if st_raw in SKIP_WORDS and st_raw:
            continue
        cin, cout = _parse_time(cell.get("in")), _parse_time(cell.get("out"))
        k = (e["id"], ds)
        rec = recs.get(k) or {"eid": e["id"], "code": e["emp_code"], "name": e["name"], "date": ds, "status": None,
                              "in": None, "out": None, "note": None, "line": n, "punches": []}
        if tm and not cin and not cout:                      # machine punch log: one row per punch
            rec.setdefault("punches", []).append(tm)
        rec["in"] = cin or rec["in"]
        rec["out"] = cout or rec["out"]
        if st:
            rec["status"] = st
        note = _cell_text(cell.get("note"))
        if note:
            rec["note"] = note
        recs[k] = rec
    rows = []
    for rec in sorted(recs.values(), key=lambda x: (x["date"], x["name"])):
        p = sorted(rec.pop("punches", []) or [])
        if p:
            rec["in"] = rec["in"] or p[0]
            if len(p) > 1:
                rec["out"] = rec["out"] or p[-1]
        if not rec["status"]:
            rec["status"] = _status_from_time(rec["in"])
        if rec["status"] in ("absent", "leave"):
            rec["in"] = rec["out"] = None
        ex = existing.get((rec["eid"], rec["date"]))
        if not ex:
            rec["action"] = "new"
        elif ex["status"] == rec["status"] and (ex["check_in"] or None) == rec["in"] and (ex["check_out"] or None) == rec["out"]:
            rec["action"] = "same"
        else:
            rec["action"] = "update" if overwrite else "keep"
            rec["old"] = ex["status"]
        rows.append(rec)
    return rows, problems, fmt


@app.route("/attendance/import", methods=["GET", "POST"])
@perm_required("attendance.manage")
def attendance_import():
    if request.method == "POST":
        fs = request.files.get("file")
        if not fs or not fs.filename:
            flash("Please choose a file to upload.", "error")
            return redirect(url_for("attendance_import"))
        overwrite = bool(request.form.get("overwrite"))
        month = request.form.get("month") or date.today().strftime("%Y-%m")
        try:
            rows, problems, fmt = analyse_att_import(read_upload(fs), month, overwrite)
        except Exception as e:                                       # noqa: BLE001
            flash(f"Couldn't read the file: {e}", "error")
            return redirect(url_for("attendance_import"))
        if not rows and not problems:
            flash("No attendance marks found in the file. Fill in P / A / L / H / LV (or dates and times) for your employees, save, and upload again.", "error")
            return redirect(url_for("attendance_import"))
        import json
        os.makedirs(IMPORT_DIR, exist_ok=True)
        token = secrets.token_hex(12)
        with open(os.path.join(IMPORT_DIR, "att_" + token + ".json"), "w", encoding="utf-8") as f:
            json.dump([r for r in rows if r["action"] in ("new", "update")], f)
        counts = Counter(r["action"] for r in rows)
        dates = sorted({r["date"] for r in rows})
        return render_template("attendance_import.html", step="preview", rows=rows[:1500], total=len(rows),
                               problems=problems, counts=counts, token=token, fmt=fmt, filename=fs.filename,
                               emp_count=len({r["eid"] for r in rows}), dates=dates, month=month)
    return render_template("attendance_import.html", step="upload", month=date.today().strftime("%Y-%m"))


@app.post("/attendance/import/confirm")
@perm_required("attendance.manage")
def attendance_import_confirm():
    import json
    token = request.form.get("token", "")
    path = os.path.join(IMPORT_DIR, "att_" + token + ".json")
    if not token.isalnum() or not os.path.exists(path):
        flash("This import has expired. Please upload the file again.", "error")
        return redirect(url_for("attendance_import"))
    with open(path, encoding="utf-8") as f:
        rows = json.load(f)
    os.remove(path)
    con = db()
    try:
        for r in rows:
            con.execute("""INSERT INTO attendance(employee_id,date,status,check_in,check_out,note) VALUES(?,?,?,?,?,?)
                           ON CONFLICT(employee_id,date) DO UPDATE SET status=excluded.status, check_in=excluded.check_in,
                           check_out=excluded.check_out, note=COALESCE(excluded.note, attendance.note)""",
                        (r["eid"], r["date"], r["status"], r["in"], r["out"], r["note"]))
        con.commit()
    except Exception as e:                                           # noqa: BLE001
        con.rollback()
        flash(f"Import failed and nothing was changed: {e}", "error")
        return redirect(url_for("attendance_import"))
    dates = sorted({r["date"] for r in rows})
    flash(f"Attendance imported: {len(rows)} record(s)" + (f" from {nicedate(dates[0])} to {nicedate(dates[-1])}" if dates else "") + ".", "ok")
    if dates:
        return redirect(url_for("attendance_report", month=dates[-1][:7]))
    return redirect(url_for("attendance_import"))


@app.route("/attendance/import/template")
@perm_required("attendance.manage")
def attendance_import_template():
    kind = request.args.get("kind", "grid")
    y, m = parse_month(request.args.get("month"))
    emps = q("""SELECT e.emp_code, e.name, d.name dept FROM employees e LEFT JOIN departments d ON d.id=e.department_id
                WHERE e.status='active' ORDER BY d.name, e.name""")
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    head_fill, off_fill = PatternFill("solid", fgColor="0F2A44"), PatternFill("solid", fgColor="E5E7EB")
    if kind == "list":
        ws.title = "Attendance"
        ws.append(["Employee Code", "Name", "Date", "Status", "Check In", "Check Out", "Note"])
        d = date.today()
        for e in emps[:3] or [{"emp_code": "EMP-0001", "name": "Example"}]:
            ws.append([e["emp_code"], e["name"], d.strftime("%d/%m/%Y"), "P", "09:02", "17:05", ""])
        ws.append([emps[0]["emp_code"] if emps else "EMP-0001", emps[0]["name"] if emps else "Example",
                   (d - timedelta(days=1)).strftime("%d/%m/%Y"), "A", "", "", "sick"])
        widths = [16, 26, 14, 10, 11, 11, 24]
        for c in ws["C"][1:]:
            c.number_format = "@"
        name = "attendance-list-template.xlsx"
    else:
        ws.title = f"{calendar.month_abbr[m]} {y}"
        days = month_dates(y, m)
        ws.append(["Employee Code", "Name", "Department"] + [d.day for d in days])
        ws.append(["", "", "Day →"] + [d.strftime("%a")[:2] for d in days])
        for e in emps:
            ws.append([e["emp_code"], e["name"], e["dept"] or ""] + ["OFF" if not is_workday(d) else "" for d in days])
        for col, d in enumerate(days, start=4):
            if not is_workday(d):
                for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=col, max_col=col):
                    for c in row:
                        c.fill = off_fill
        widths = [16, 26, 18] + [4.5] * len(days)
        ws.freeze_panes = "D3"
        name = f"attendance-{y}-{m:02d}-sheet.xlsx"
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center")
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    info = wb.create_sheet("How to fill")
    lines = (["Monthly sheet: one row per employee, one column per day.",
              "Codes: P = Present, L = Late, H = Half day, A = Absent, LV = Leave (CL / SL / AL also mean leave).",
              "Leave a cell empty (or OFF) to skip that day. Grey columns are weekly off days.",
              "When uploading, choose the same month on the import page."] if kind != "list" else
             ["List format: one row per employee per day.",
              "Date as DD/MM/YYYY. Status: P, L, H, A or LV (or leave empty and give Check In; late is detected automatically).",
              "Times as HH:MM (24h) or 9:05 AM.",
              "Biometric machine exports also work: one row per punch with Employee No + Date/Time; first punch = in, last = out."])
    for l_ in lines:
        info.append([l_])
    info.column_dimensions["A"].width = 120
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=name,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# --------------------------------------------------------------------------- leaves
@app.route("/leaves", methods=["GET", "POST"])
@perm_required("leaves.view", "leaves.approve")
def leaves():
    if request.method == "POST":
        if not can("leaves.approve"):
            abort(403)
        eid = request.form.get("employee_id", "")
        if not q("SELECT 1 FROM employees WHERE id=?", (eid,), one=True):
            flash("Please choose an employee.", "error")
        else:
            err = create_leave(int(eid), request.form.get("leave_type_id"), request.form.get("start_date"),
                               request.form.get("end_date"), request.form.get("reason", ""),
                               status="approved", decided_by=g.user["username"])
            flash(err or "Leave recorded and approved.", "error" if err else "ok")
        return redirect(url_for("leaves", status="approved"))
    status = request.args.get("status") or ("pending" if q("SELECT 1 FROM leaves WHERE status='pending'", one=True) else "all")
    sql = """SELECT l.*, e.name, e.emp_code, t.name tname, t.paid FROM leaves l
             JOIN employees e ON e.id=l.employee_id JOIN leave_types t ON t.id=l.leave_type_id"""
    args = []
    if status in ("pending", "approved", "rejected", "cancelled"):
        sql += " WHERE l.status=?"
        args.append(status)
    rows = q(sql + " ORDER BY l.start_date DESC, l.id DESC LIMIT 500", args)
    emps = q("""SELECT e.id, e.name, e.emp_code, d.name dept FROM employees e LEFT JOIN departments d ON d.id=e.department_id
                WHERE e.status='active' ORDER BY e.name""")
    today = date.today()
    t = today.isoformat()
    month = today.strftime("%Y-%m")
    stats = {
        "pending": q("SELECT COUNT(*) c FROM leaves WHERE status='pending'", one=True)["c"],
        "on_leave": q("""SELECT e.name FROM leaves l JOIN employees e ON e.id=l.employee_id WHERE l.status='approved'
                         AND l.start_date<=? AND l.end_date>=? ORDER BY e.name""", (t, t)),
        "upcoming": q("""SELECT COUNT(*) c FROM leaves WHERE status='approved' AND start_date>? AND start_date<=?""",
                      (t, (today + timedelta(days=7)).isoformat()), one=True)["c"],
        "month_days": q("""SELECT COALESCE(SUM(days),0) s FROM leaves WHERE status='approved' AND substr(start_date,1,7)=?""",
                        (month,), one=True)["s"],
    }
    counts = {r["status"]: r["c"] for r in q("SELECT status, COUNT(*) c FROM leaves GROUP BY status")}
    return render_template("leaves.html", rows=rows, status=status, emps=emps, stats=stats, counts=counts, today=today,
                           types=q("SELECT * FROM leave_types ORDER BY name"))


@app.post("/leaves/<int:lid>/decide")
@perm_required("leaves.approve")
def leave_decide(lid):
    lv = q("SELECT * FROM leaves WHERE id=?", (lid,), one=True) or abort(404)
    action = request.form.get("action")
    remarks = request.form.get("remarks", "").strip()
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    if action in ("approve", "reject") and lv["status"] == "pending":
        new = "approved" if action == "approve" else "rejected"
        x("UPDATE leaves SET status=?, decided_by=?, decided_on=?, remarks=? WHERE id=?",
          (new, g.user["username"], now, remarks, lid))
        if new == "approved":
            apply_leave_attendance(lv)
        flash(f"Leave {new}.", "ok")
    elif action == "cancel" and lv["status"] == "approved":
        x("UPDATE leaves SET status='cancelled', decided_by=?, decided_on=?, remarks=? WHERE id=?",
          (g.user["username"], now, remarks or "Cancelled by HR", lid))
        remove_leave_attendance(lv)
        flash("Approved leave cancelled and attendance cleared for those days.", "ok")
    return redirect(request.referrer or url_for("leaves"))


@app.route("/leaves/<int:lid>/edit", methods=["GET", "POST"])
@perm_required("leaves.approve")
def leave_edit(lid):
    lv = q("""SELECT l.*, e.name, e.emp_code FROM leaves l JOIN employees e ON e.id=l.employee_id WHERE l.id=?""",
           (lid,), one=True) or abort(404)
    if lv["status"] not in ("pending", "approved"):
        flash("Only pending or approved leaves can be edited.", "error")
        return redirect(url_for("leaves"))
    if request.method == "POST":
        err, lt, start, end, days = check_leave(lv["employee_id"], request.form.get("leave_type_id"),
                                                request.form.get("start_date"), request.form.get("end_date"), exclude_id=lid)
        if err:
            flash(err, "error")
            return redirect(url_for("leave_edit", lid=lid))
        if lv["status"] == "approved":
            remove_leave_attendance(lv)
        x("""UPDATE leaves SET leave_type_id=?, start_date=?, end_date=?, days=?, reason=?, remarks=? WHERE id=?""",
          (lt["id"], start.isoformat(), end.isoformat(), days, request.form.get("reason", "").strip(),
           request.form.get("remarks", "").strip() or lv["remarks"], lid))
        if lv["status"] == "approved":
            apply_leave_attendance(q("SELECT * FROM leaves WHERE id=?", (lid,), one=True))
        flash(f"Leave updated for {lv['name']}: {nicedate(start.isoformat())} – {nicedate(end.isoformat())} ({days} day(s)).", "ok")
        return redirect(url_for("leaves", status=lv["status"]))
    return render_template("leave_edit.html", lv=lv, types=q("SELECT * FROM leave_types ORDER BY name"),
                           balances=leave_balances(lv["employee_id"], int(lv["start_date"][:4])))


@app.route("/api/unmarked")
@perm_required("leaves.approve")
def api_unmarked():
    """Active employees with no attendance (and no leave) on a date – to pick who to give leave."""
    d = (parse_date(request.args.get("date")) or date.today()).isoformat()
    rows = q("""SELECT id FROM employees e WHERE e.status='active'
                AND NOT EXISTS (SELECT 1 FROM attendance a WHERE a.employee_id=e.id AND a.date=?)
                AND NOT EXISTS (SELECT 1 FROM leaves l WHERE l.employee_id=e.id AND l.status IN ('pending','approved')
                                AND l.start_date<=? AND l.end_date>=?)""", (d, d, d))
    absent = q("SELECT employee_id FROM attendance WHERE date=? AND status='absent'", (d,))
    return {"date": d, "unmarked": [r["id"] for r in rows], "absent": [r["employee_id"] for r in absent]}


@app.route("/my/leaves", methods=["GET", "POST"])
@login_required
def my_leaves():
    emp_id = g.user["employee_id"] or abort(403)
    if request.method == "POST":
        err = create_leave(emp_id, request.form.get("leave_type_id"), request.form.get("start_date"),
                           request.form.get("end_date"), request.form.get("reason", ""))
        flash(err or "Leave request sent to HR.", "error" if err else "ok")
        return redirect(url_for("my_leaves"))
    rows = q("""SELECT l.*, t.name tname FROM leaves l JOIN leave_types t ON t.id=l.leave_type_id
                WHERE l.employee_id=? ORDER BY l.start_date DESC""", (emp_id,))
    return render_template("my_leaves.html", rows=rows, balances=leave_balances(emp_id),
                           types=q("SELECT * FROM leave_types ORDER BY name"))


@app.post("/my/leaves/<int:lid>/cancel")
@login_required
def my_leave_cancel(lid):
    lv = q("SELECT * FROM leaves WHERE id=? AND employee_id=?", (lid, g.user["employee_id"]), one=True) or abort(404)
    if lv["status"] == "pending":
        x("UPDATE leaves SET status='cancelled', remarks='Withdrawn by employee' WHERE id=?", (lid,))
        flash("Leave request withdrawn.", "ok")
    return redirect(url_for("my_leaves"))


# --------------------------------------------------------------------------- payroll
@app.route("/payroll")
@perm_required("payroll.view")
def payroll():
    y, m = parse_month(request.args.get("month"))
    month = f"{y:04d}-{m:02d}"
    rows = q("""SELECT p.*, e.name, e.emp_code, e.designation, d.name dept FROM payroll p
                JOIN employees e ON e.id=p.employee_id LEFT JOIN departments d ON d.id=e.department_id
                WHERE p.month=? ORDER BY d.name, e.name""", (month,))
    tot = {k: sum((r[k] or 0) for r in rows) for k in
           ("basic", "allowances", "overtime", "bonus", "absent_deduction", "tax", "other_deductions",
            "loan_deduction", "net")}
    missing = q("""SELECT COUNT(*) c FROM employees WHERE status='active' AND id NOT IN
                   (SELECT employee_id FROM payroll WHERE month=?)""", (month,), one=True)["c"]
    if request.args.get("export") == "csv":
        head = ["Code", "Name", "Department", "Designation", "Basic", "Allowances", "Overtime", "Bonus",
                "Working days", "Present", "Absent", "Late", "Unpaid leave", "Deduction days", "Absent deduction",
                "Tax", "Other deductions", "Loan installment", "Net", "Status"]
        body = [[r["emp_code"], r["name"], r["dept"] or "", r["designation"] or "", r["basic"], r["allowances"],
                 r["overtime"], r["bonus"], r["working_days"], r["present_days"], r["absent_days"], r["late_days"],
                 r["unpaid_leave_days"], r["deduction_days"], r["absent_deduction"], r["tax"],
                 r["other_deductions"], r["loan_deduction"] or 0, r["net"], r["status"]] for r in rows]
        return csv_response(f"payroll-{month}.csv", head, body)
    loan_out = q("""SELECT COUNT(*) c, COALESCE(SUM(bal),0) s FROM (SELECT l.amount - COALESCE((SELECT SUM(amount)
                    FROM loan_payments p WHERE p.loan_id=l.id),0) bal FROM loans l WHERE l.status<>'closed') WHERE bal>0.5""",
                 one=True)
    pend = [r for r in rows if r["ded_status"] == "pending" and r["status"] != "paid"]
    return render_template("payroll.html", rows=rows, month=month, tot=tot, missing=missing, loan_out=loan_out, pend=pend,
                           prev=shift_month(y, m, -1), next=shift_month(y, m, 1))


@app.post("/payroll/generate")
@perm_required("payroll.manage")
def payroll_generate():
    y, m = parse_month(request.form.get("month"))
    month = f"{y:04d}-{m:02d}"
    last = month_dates(y, m)[-1].isoformat()
    emps = q("""SELECT * FROM employees WHERE status='active'
                AND (join_date IS NULL OR join_date='' OR join_date<=?)""", (last,))
    made = updated = skipped = review = 0
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    for e in emps:
        c = compute_payroll(e, y, m)
        ex = q("SELECT * FROM payroll WHERE employee_id=? AND month=?", (e["id"], month), one=True)
        if ex and ex["status"] == "paid":
            skipped += 1
            continue
        extra = {k: (ex[k] if ex else 0) or 0 for k in ("overtime", "bonus", "tax", "other_deductions")}
        sug = c["absent_deduction"]
        exk = ex.keys() if ex else []
        if ex and "ded_status" in exk and ex["ded_status"] in ("applied", "waived") \
                and abs((ex["suggested_deduction"] or 0) - sug) < 1:
            c["absent_deduction"], dst = ex["absent_deduction"] or 0, ex["ded_status"]   # HR already decided
        else:
            c["absent_deduction"], dst = 0, ("pending" if sug > 0 else None)          # HR reviews and deducts
        if dst == "pending":
            review += 1
        row = {**c, **extra, "suggested_deduction": sug, "ded_status": dst}
        if ex:                                   # recalculating: forget this salary's previous loan cuts
            x("DELETE FROM loan_payments WHERE payroll_id=?", (ex["id"],))
        row["loan_deduction"] = 0
        room = max(calc_net(row), 0)             # never take more than the salary
        cuts = []
        for ln in q("""SELECT * FROM loans WHERE employee_id=? AND status='active' AND
                       (start_month IS NULL OR start_month='' OR start_month<=?) ORDER BY issued_on, id""",
                    (e["id"], month)):
            ded = round(min(ln["installment"] or 0, loan_balance(ln), room))
            if ded > 0:
                cuts.append((ln["id"], ded))
                room -= ded
        row["loan_deduction"] = sum(d for _, d in cuts)
        row["net"] = calc_net(row)
        cols = list(row.keys())
        if ex:
            x(f"UPDATE payroll SET {', '.join(k + '=?' for k in cols)}, generated_on=? WHERE id=?",
              [row[k] for k in cols] + [now, ex["id"]])
            pid_ = ex["id"]
            updated += 1
        else:
            pid_ = x(f"INSERT INTO payroll(employee_id,month,{','.join(cols)},generated_on) "
                     f"VALUES(?,?,{','.join('?' * len(cols))},?)", [e["id"], month] + [row[k] for k in cols] + [now])
            made += 1
        for lid_, ded in cuts:
            x("""INSERT INTO loan_payments(loan_id,payroll_id,month,amount,paid_on,method,note,created_by)
                 VALUES(?,?,?,?,?,'salary','Deducted from salary',?)""",
              (lid_, pid_, month, ded, date.today().isoformat(), g.user["username"]))
    flash(f"Payroll for {monthname(month)}: {made} created, {updated} recalculated, {skipped} already paid (unchanged).", "ok")
    if review:
        flash(f"🔔 {review} salary(ies) have absent / late / unpaid-leave deductions to review. Nothing was cut yet – "
              "check each one below and click Deduct or Don't deduct.", "warn")
    return redirect(url_for("payroll", month=month))


@app.post("/payroll/<int:pid>/deduction")
@perm_required("payroll.manage")
def payroll_deduction(pid):
    p = get_pay(pid)
    if p["status"] == "paid":
        flash("This salary is already paid. Mark it unpaid first.", "warn")
        return redirect(url_for("payroll", month=p["month"]))
    act = request.form.get("act")
    row = dict(p)
    if act == "apply":
        row["absent_deduction"] = to_float(request.form.get("amount"), p["suggested_deduction"] or 0)
        row["ded_status"] = "applied"
    elif act == "waive":
        row["absent_deduction"], row["ded_status"] = 0, "waived"
    elif act == "undo":
        row["absent_deduction"], row["ded_status"] = 0, "pending"
    else:
        abort(400)
    row["net"] = calc_net(row)
    x("UPDATE payroll SET absent_deduction=?, ded_status=?, net=? WHERE id=?",
      (row["absent_deduction"], row["ded_status"], row["net"], pid))
    flash({"apply": f"Deducted {money(row['absent_deduction'])} from {p['name']}'s salary.",
           "waive": f"No deduction for {p['name']} this month.",
           "undo": f"{p['name']}'s deduction is back to review."}[act], "ok")
    return redirect(url_for("payroll", month=p["month"]) + "#ded")


@app.post("/payroll/deductions/apply-all")
@perm_required("payroll.manage")
def payroll_deduct_all():
    y, m = parse_month(request.form.get("month"))
    month = f"{y:04d}-{m:02d}"
    n = 0
    for p in q("SELECT * FROM payroll WHERE month=? AND ded_status='pending' AND status<>'paid'", (month,)):
        row = dict(p, absent_deduction=p["suggested_deduction"] or 0)
        x("UPDATE payroll SET absent_deduction=?, ded_status='applied', net=? WHERE id=?",
          (row["absent_deduction"], calc_net(row), p["id"]))
        n += 1
    flash(f"Applied all {n} suggested deduction(s) for {monthname(month)}.", "ok")
    return redirect(url_for("payroll", month=month))


def get_pay(pid):
    p = q("""SELECT p.*, e.name, e.emp_code, e.designation, e.cnic, e.bank_account, e.join_date, d.name dept
             FROM payroll p JOIN employees e ON e.id=p.employee_id LEFT JOIN departments d ON d.id=e.department_id
             WHERE p.id=?""", (pid,), one=True)
    if not p:
        abort(404)
    return p


@app.route("/payroll/<int:pid>/edit", methods=["GET", "POST"])
@perm_required("payroll.manage")
def payroll_edit(pid):
    p = get_pay(pid)
    if p["status"] == "paid":
        flash("This salary is marked paid. Mark it unpaid first to edit.", "warn")
        return redirect(url_for("payroll", month=p["month"]))
    if request.method == "POST":
        row = dict(p)
        for k in ("basic", "allowances", "overtime", "bonus", "absent_deduction", "tax", "other_deductions"):
            row[k] = to_float(request.form.get(k), row[k] or 0)
        row["remarks"] = request.form.get("remarks", "").strip()
        row["net"] = calc_net(row)
        dst = p["ded_status"]
        if abs(row["absent_deduction"] - (p["absent_deduction"] or 0)) >= 0.5:
            dst = "applied" if row["absent_deduction"] > 0 else "waived"
        x("""UPDATE payroll SET basic=?, allowances=?, overtime=?, bonus=?, absent_deduction=?, tax=?,
             other_deductions=?, remarks=?, net=?, ded_status=? WHERE id=?""",
          (row["basic"], row["allowances"], row["overtime"], row["bonus"], row["absent_deduction"], row["tax"],
           row["other_deductions"], row["remarks"], row["net"], dst, pid))
        flash(f"Salary updated for {p['name']}. Net: {money(row['net'])}", "ok")
        return redirect(url_for("payroll", month=p["month"]))
    cuts = q("""SELECT lp.amount, l.id loan_id, l.reason FROM loan_payments lp JOIN loans l ON l.id=lp.loan_id
                WHERE lp.payroll_id=?""", (pid,))
    return render_template("payroll_edit.html", p=p, cuts=cuts)


@app.post("/payroll/<int:pid>/status")
@perm_required("payroll.manage")
def payroll_status(pid):
    p = get_pay(pid)
    if p["status"] == "paid":
        x("UPDATE payroll SET status='draft', paid_on=NULL WHERE id=?", (pid,))
    else:
        x("UPDATE payroll SET status='paid', paid_on=? WHERE id=?", (date.today().isoformat(), pid))
    return redirect(url_for("payroll", month=p["month"]))


@app.post("/payroll/pay-all")
@perm_required("payroll.manage")
def payroll_pay_all():
    y, m = parse_month(request.form.get("month"))
    month = f"{y:04d}-{m:02d}"
    x("UPDATE payroll SET status='paid', paid_on=? WHERE month=? AND status<>'paid'", (date.today().isoformat(), month))
    flash(f"All salaries for {monthname(month)} marked paid.", "ok")
    return redirect(url_for("payroll", month=month))


@app.post("/payroll/<int:pid>/delete")
@perm_required("payroll.manage")
def payroll_delete(pid):
    p = get_pay(pid)
    if p["status"] != "paid":
        to_bin("payroll", pid)
        x("DELETE FROM loan_payments WHERE payroll_id=?", (pid,))
        x("DELETE FROM payroll WHERE id=?", (pid,))
        flash("Salary row removed.", "ok")
    return redirect(url_for("payroll", month=p["month"]))


@app.route("/payslip/<int:pid>")
@login_required
def payslip(pid):
    p = get_pay(pid)
    if not can("payroll.view") and (p["employee_id"] != g.user["employee_id"] or p["status"] != "paid"):
        abort(403)
    earnings = [("Basic salary", p["basic"]), ("Allowances", p["allowances"]),
                ("Overtime", p["overtime"]), ("Bonus", p["bonus"])]
    deductions = [("Absent / unpaid days (%s)" % num(p["deduction_days"]), p["absent_deduction"]),
                  ("Income tax", p["tax"]), ("Other deductions", p["other_deductions"])]
    for lp in q("""SELECT lp.amount, l.id FROM loan_payments lp JOIN loans l ON l.id=lp.loan_id
                   WHERE lp.payroll_id=?""", (pid,)):
        left = q("""SELECT l.amount - COALESCE((SELECT SUM(amount) FROM loan_payments WHERE loan_id=l.id),0) b
                    FROM loans l WHERE l.id=?""", (lp["id"],), one=True)["b"]
        deductions.append((f"Loan installment (balance left {money(max(left, 0))})", lp["amount"]))
    if (p["loan_deduction"] or 0) and not any(l.startswith("Loan") for l, _ in deductions):
        deductions.append(("Loan installment", p["loan_deduction"]))
    return render_template("payslip.html", p=p, earnings=earnings, deductions=deductions,
                           gross=sum(v or 0 for _, v in earnings), total_ded=sum(v or 0 for _, v in deductions))


@app.route("/my/payslips")
@login_required
def my_payslips():
    emp_id = g.user["employee_id"] or abort(403)
    rows = q("SELECT * FROM payroll WHERE employee_id=? AND status='paid' ORDER BY month DESC", (emp_id,))
    return render_template("my_payslips.html", rows=rows)


@app.route("/my/profile")
@login_required
def my_profile():
    emp_id = g.user["employee_id"] or abort(403)
    return render_template("my_profile.html", e=get_emp(emp_id))


# --------------------------------------------------------------------------- loans
def loan_balance(ln):
    paid = q("SELECT COALESCE(SUM(amount),0) s FROM loan_payments WHERE loan_id=?", (ln["id"],), one=True)["s"]
    return max(round((ln["amount"] or 0) - paid, 2), 0)


LOAN_SQL = """SELECT l.*, e.name, e.emp_code, e.designation, e.status emp_status, d.name dept,
                     COALESCE((SELECT SUM(amount) FROM loan_payments p WHERE p.loan_id=l.id),0) paid,
                     (SELECT COUNT(*) FROM loan_payments p WHERE p.loan_id=l.id) n_pay,
                     (SELECT MAX(month) FROM loan_payments p WHERE p.loan_id=l.id) last_month
              FROM loans l JOIN employees e ON e.id=l.employee_id LEFT JOIN departments d ON d.id=e.department_id"""


def loan_state(r):
    bal = max((r["amount"] or 0) - (r["paid"] or 0), 0)
    if bal <= 0.5:
        return "completed"
    return r["status"]                      # active / paused / closed


@app.route("/payroll/loans", methods=["GET", "POST"])
@perm_required("payroll.view")
def loans():
    if request.method == "POST":
        if not can("payroll.manage"):
            abort(403)
        eid = request.form.get("employee_id", "")
        amount = to_float(request.form.get("amount"), 0)
        inst = to_float(request.form.get("installment"), 0)
        issued = request.form.get("issued_on") or date.today().isoformat()
        start = request.form.get("start_month") or issued[:7]
        if not q("SELECT 1 FROM employees WHERE id=?", (eid,), one=True):
            flash("Please choose an employee.", "error")
        elif amount <= 0:
            flash("Enter the loan amount.", "error")
        elif inst <= 0:
            flash("Enter the monthly installment to deduct from salary.", "error")
        else:
            lid = x("""INSERT INTO loans(employee_id,amount,installment,issued_on,start_month,reason,status,created_by)
                       VALUES(?,?,?,?,?,?,'active',?)""",
                    (int(eid), amount, min(inst, amount), issued, start[:7], request.form.get("reason", "").strip(),
                     g.user["username"]))
            flash(f"Loan of {settings_get('currency')} {money(amount)} added. {money(min(inst, amount))} will be "
                  f"deducted from salary every month from {monthname(start[:7])}.", "ok")
            return redirect(url_for("loan_view", lid=lid))
        return redirect(url_for("loans"))
    rows = [dict(r, balance=max((r["amount"] or 0) - (r["paid"] or 0), 0), state=loan_state(r))
            for r in q(LOAN_SQL + " ORDER BY l.issued_on DESC, l.id DESC")]
    view = request.args.get("view", "pending")
    counts = Counter(r["state"] for r in rows)
    pending = [r for r in rows if r["state"] in ("active", "paused")]
    shown = {"pending": pending, "completed": [r for r in rows if r["state"] == "completed"],
             "closed": [r for r in rows if r["state"] == "closed"], "all": rows}.get(view, pending)
    month = date.today().strftime("%Y-%m")
    kp = {"outstanding": sum(r["balance"] for r in pending), "pending": len(pending),
          "next": sum(min(r["installment"] or 0, r["balance"]) for r in pending
                      if r["state"] == "active" and (r["start_month"] or "") <= month),
          "given": sum(r["amount"] for r in rows),
          "recovered_month": q("SELECT COALESCE(SUM(amount),0) s FROM loan_payments WHERE month=?", (month,), one=True)["s"]}
    if request.args.get("export") == "csv":
        head = ["Code", "Employee", "Department", "Issued on", "Amount", "Monthly installment", "Paid", "Balance",
                "Status", "Reason"]
        body = [[r["emp_code"], r["name"], r["dept"] or "", r["issued_on"], r["amount"], r["installment"],
                 r["paid"], r["balance"], r["state"], r["reason"] or ""] for r in shown]
        return csv_response(f"loans-{view}.csv", head, body)
    emps = q("SELECT id, name, emp_code FROM employees WHERE status='active' ORDER BY name")
    return render_template("loans.html", rows=shown, view=view, counts=counts, kp=kp, emps=emps, month=month,
                           pre=request.args.get("emp", ""))


def settings_get(k):
    r = q("SELECT value FROM settings WHERE key=?", (k,), one=True)
    return r["value"] if r else ""


@app.route("/payroll/loans/<int:lid>", methods=["GET", "POST"])
@perm_required("payroll.view")
def loan_view(lid):
    ln = q(LOAN_SQL + " WHERE l.id=?", (lid,), one=True) or abort(404)
    if request.method == "POST":
        if not can("payroll.manage"):
            abort(403)
        action = request.form.get("action")
        bal = loan_balance(ln)
        if action == "edit":
            inst = to_float(request.form.get("installment"), ln["installment"])
            amount = to_float(request.form.get("amount"), ln["amount"])
            if amount < (ln["paid"] or 0):
                flash("Amount can't be less than what is already repaid.", "error")
            elif inst <= 0:
                flash("Installment must be more than 0.", "error")
            else:
                x("UPDATE loans SET amount=?, installment=?, start_month=?, reason=?, issued_on=? WHERE id=?",
                  (amount, inst, (request.form.get("start_month") or ln["start_month"] or "")[:7],
                   request.form.get("reason", "").strip(), request.form.get("issued_on") or ln["issued_on"], lid))
                flash("Loan updated. Recalculate payroll for months not yet paid to apply the new installment.", "ok")
        elif action == "pay":
            amt = to_float(request.form.get("amount"), 0)
            if amt <= 0:
                flash("Enter the amount received.", "error")
            elif amt > bal + 0.5:
                flash(f"That is more than the balance ({money(bal)}).", "error")
            else:
                paid_on = request.form.get("paid_on") or date.today().isoformat()
                x("""INSERT INTO loan_payments(loan_id,month,amount,paid_on,method,note,created_by)
                     VALUES(?,?,?,?,'cash',?,?)""", (lid, paid_on[:7], amt, paid_on,
                                                     request.form.get("note", "").strip() or "Cash repayment", g.user["username"]))
                flash(f"Repayment of {money(amt)} recorded." + (" Loan fully repaid 🎉" if amt >= bal - 0.5 else ""), "ok")
        elif action == "unpay":
            pay = q("SELECT * FROM loan_payments WHERE id=? AND loan_id=?", (request.form.get("pay_id"), lid), one=True)
            if pay and not pay["payroll_id"]:
                x("DELETE FROM loan_payments WHERE id=?", (pay["id"],))
                flash("Repayment removed.", "ok")
            else:
                flash("Salary deductions are removed by recalculating or deleting that month's salary.", "warn")
        elif action in ("pause", "resume"):
            x("UPDATE loans SET status=? WHERE id=?", ("paused" if action == "pause" else "active", lid))
            flash("Deductions paused – nothing will be cut from salary until you resume." if action == "pause"
                  else "Deductions resumed from the next payroll.", "ok")
        elif action == "close":
            x("UPDATE loans SET status='closed', closed_on=?, notes=? WHERE id=?",
              (date.today().isoformat(), request.form.get("notes", "").strip() or "Closed / waived", lid))
            flash(f"Loan closed. Remaining {money(bal)} will not be deducted.", "ok")
        elif action == "reopen":
            x("UPDATE loans SET status='active', closed_on=NULL WHERE id=?", (lid,))
            flash("Loan reopened.", "ok")
        elif action == "delete":
            if ln["n_pay"]:
                flash("This loan has repayments. Close it instead of deleting.", "error")
            else:
                to_bin("loan", lid)
                x("DELETE FROM loans WHERE id=?", (lid,))
                flash("Loan deleted.", "ok")
                return redirect(url_for("loans"))
        return redirect(url_for("loan_view", lid=lid))
    pays = q("""SELECT lp.*, p.status pay_status FROM loan_payments lp LEFT JOIN payroll p ON p.id=lp.payroll_id
                WHERE lp.loan_id=? ORDER BY lp.paid_on DESC, lp.id DESC""", (lid,))
    bal = max((ln["amount"] or 0) - (ln["paid"] or 0), 0)
    months_left = int(-(-bal // ln["installment"])) if ln["installment"] and bal > 0 else 0
    return render_template("loan_view.html", ln=ln, pays=pays, bal=bal, state=loan_state(ln), months_left=months_left,
                           others=q(LOAN_SQL + " WHERE l.employee_id=? AND l.id<>? ORDER BY l.issued_on DESC",
                                    (ln["employee_id"], lid)))


# --------------------------------------------------------------------------- activity log
@app.route("/activity")
@perm_required("users")
def activity():
    sql, args = "SELECT * FROM audit_log WHERE 1=1", []
    who = request.args.get("user", "").strip()
    day = request.args.get("date", "").strip()
    kind = request.args.get("kind", "")
    if who:
        sql += " AND username=?"
        args.append(who)
    if day:
        sql += " AND substr(at,1,10)=?"
        args.append(day)
    if kind == "security":
        sql += " AND (action LIKE '%login%' OR action LIKE 'Signed%' OR action LIKE '%locked%' OR action LIKE 'Logged out%' OR action LIKE '%backup%' OR action LIKE '%password%' OR action LIKE '%users%' OR action LIKE '%roles%')"
    rows = q(sql + " ORDER BY id DESC LIMIT 1000", args)
    if request.args.get("export") == "csv":
        return csv_response("activity-log.csv", ["Time", "User", "Computer (IP)", "Action", "Details"],
                            [[r["at"], r["username"], r["ip"], r["action"], r["detail"]] for r in rows])
    users_ = [r["username"] for r in q("SELECT DISTINCT username FROM audit_log WHERE username<>'' ORDER BY username")]
    fails24 = q("""SELECT COUNT(*) c FROM audit_log WHERE action IN ('Failed login','Account locked','Login blocked (locked)')
                   AND at>=?""", ((datetime.now() - timedelta(days=1)).strftime("%Y-%m-%d %H:%M:%S"),), one=True)["c"]
    autos = [b for b in list_backups() if b["kind"] == "auto"]
    return render_template("activity.html", rows=rows, users_=users_, who=who, day=day, kind=kind, fails24=fails24,
                           autos=autos, db_size=os.path.getsize(DB_PATH) if os.path.exists(DB_PATH) else 0)


# --------------------------------------------------------------------------- settings & users
@app.route("/settings", methods=["GET", "POST"])
@perm_required("settings")
def settings_page():
    if request.method == "POST":
        form = request.form.get("form")
        if form == "general":
            for k in ("company_name", "company_address", "currency", "office_start", "office_end",
                      "grace_minutes", "lates_per_day", "idle_minutes"):
                x("UPDATE settings SET value=? WHERE key=?", (request.form.get(k, "").strip(), k))
            off = ",".join(v for v in request.form.getlist("weekly_off") if v.isdigit())
            x("UPDATE settings SET value=? WHERE key='weekly_off'", (off,))
            flash("Settings saved.", "ok")
        elif form == "notify":
            off = [t for t, *_ in NOTIFY_TYPES if not request.form.get("n_" + t)]
            x("INSERT OR REPLACE INTO settings(key,value) VALUES('notify_off',?)", (",".join(off),))
            for k in ("notify_sound", "notify_loud", "click_sound"):
                x("INSERT OR REPLACE INTO settings(key,value) VALUES(?,?)", (k, "1" if request.form.get(k) else "0"))
            vol = request.form.get("click_volume")
            if vol in ("low", "medium", "high"):
                x("INSERT OR REPLACE INTO settings(key,value) VALUES('click_volume',?)", (vol,))
            flash("Notification settings saved.", "ok")
        elif form == "leave_type":
            name = request.form.get("name", "").strip()
            days = int(to_float(request.form.get("days_per_year"), 0))
            paid = 1 if request.form.get("paid") else 0
            enc = 1 if request.form.get("encash") else 0
            tid = request.form.get("id")
            try:
                if not name:
                    flash("Leave type name is required.", "error")
                elif tid:
                    x("UPDATE leave_types SET name=?, days_per_year=?, paid=?, encash=? WHERE id=?", (name, days, paid, enc, tid))
                    flash("Leave type updated.", "ok")
                else:
                    x("INSERT INTO leave_types(name,days_per_year,paid,encash) VALUES(?,?,?,?)", (name, days, paid, enc))
                    flash("Leave type added.", "ok")
            except sqlite3.IntegrityError:
                flash("A leave type with that name already exists.", "error")
        elif form == "encash":
            b = request.form.get("encash_basis")
            if b in ENCASH_BASIS:
                x("INSERT INTO settings(key,value) VALUES('encash_basis',?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (b,))
                flash(f"Unused leave will be paid at {ENCASH_BASIS[b]} per day.", "ok")
        elif form == "leave_type_delete":
            try:
                if not q("SELECT 1 FROM leaves WHERE leave_type_id=?", (request.form.get("id"),), one=True):
                    to_bin("leave_type", request.form.get("id", type=int))
                x("DELETE FROM leave_types WHERE id=?", (request.form.get("id"),))
                flash("Leave type deleted.", "ok")
            except sqlite3.IntegrityError:
                flash("This leave type is used by existing leaves and can't be deleted. Rename it instead.", "error")
        return redirect(url_for("settings_page"))
    return render_template("settings.html", encash_basis=ENCASH_BASIS, bin_count=q("SELECT COUNT(*) c FROM recycle_bin", one=True)["c"], types=q("SELECT * FROM leave_types ORDER BY name"),
                           off=weekly_off(), weekdays=WEEKDAYS, lan=lan_address(), pcname=socket.gethostname(),
                           notify_types=NOTIFY_TYPES,
                           notify_off=set((get_settings().get("notify_off") or "").split(",")))


@app.route("/users/<int:uid>/permissions", methods=["GET", "POST"])
@perm_required("users")
def user_perms_page(uid):
    target = q("SELECT * FROM users WHERE id=?", (uid,), one=True)
    if not target:
        abort(404)
    role = q("SELECT * FROM roles WHERE key=?", (target["role"],), one=True)
    base = role_perm_set(target["role"])
    if request.method == "POST":
        if target["role"] == "admin":
            flash("Administrators always have every permission.", "error")
        elif request.form.get("action") == "reset":
            x("UPDATE users SET perm_add='', perm_remove='' WHERE id=?", (uid,))
            flash(f"{target['username']} now has exactly the permissions of the role {role['name'] if role else target['role']}.", "ok")
        else:
            want = set(request.form.getlist("perm")) & set(ALL_PERMS)
            if target["id"] == g.user["id"] and "users" not in want:
                flash("You can't remove the Users permission from your own account.", "error")
                return redirect(url_for("user_perms_page", uid=uid))
            x("UPDATE users SET perm_add=?, perm_remove=? WHERE id=?",
              (",".join(sorted(want - base)), ",".join(sorted(base - want)), uid))
            flash(f"Permissions saved for {target['username']}.", "ok")
        return redirect(url_for("user_perms_page", uid=uid))
    groups = {}
    for key, module, label in PERMISSIONS:
        groups.setdefault(module, []).append((key, label))
    return render_template("user_perms.html", target=target, role=role, base=base, have=effective_perms(target),
                           groups=groups, add=_split(target["perm_add"]), rem=_split(target["perm_remove"]))


@app.route("/users", methods=["GET", "POST"])
@perm_required("users")
def users():
    roles = {r["key"]: r for r in q("SELECT * FROM roles")}
    if request.method == "POST":
        action = request.form.get("action")
        uid = request.form.get("id")
        pw = request.form.get("password", "")
        admins = q("SELECT COUNT(*) c FROM users WHERE role='admin' AND active=1", one=True)["c"]
        target = q("SELECT * FROM users WHERE id=?", (uid,), one=True) if uid else None
        if action == "create":
            uname = request.form.get("username", "").strip()
            role = request.form.get("role")
            if role not in roles:
                flash("Choose a role.", "error")
            elif not uname or len(pw) < 6:
                flash("Enter a username and a password of at least 6 characters.", "error")
            elif q("SELECT 1 FROM users WHERE username=?", (uname,), one=True):
                flash(f"Username '{uname}' is already taken.", "error")
            else:
                add = rem = ""
                if request.form.get("custom") == "1" and role != "admin":
                    want = set(request.form.getlist("perm")) & set(ALL_PERMS)
                    base = role_perm_set(role)
                    add, rem = ",".join(sorted(want - base)), ",".join(sorted(base - want))
                x("INSERT INTO users(username,password_hash,role,must_change,perm_add,perm_remove) VALUES(?,?,?,1,?,?)",
                  (uname, generate_password_hash(pw), role, add, rem))
                flash(f"User '{uname}' created with role {roles[role]['name']}" +
                      (" and their own permission switches." if add or rem else "."), "ok")
        elif not target:
            flash("User not found.", "error")
        elif target["id"] == g.user["id"] and action in ("toggle", "delete", "role"):
            flash("You can't disable, delete or change the role of your own account.", "error")
        elif target["role"] == "admin" and admins <= 1 and (action in ("toggle", "delete") or
                                                             (action == "role" and request.form.get("role") != "admin")):
            flash("This is the last active administrator. Make someone else an administrator first.", "error")
        elif action == "role":
            role = request.form.get("role")
            if role in roles:
                x("UPDATE users SET role=? WHERE id=?", (role, uid))
                flash(f"{target['username']} is now {roles[role]['name']}.", "ok")
        elif action == "reset":
            if len(pw) < 6:
                flash("Password must be at least 6 characters.", "error")
            else:
                x("UPDATE users SET password_hash=?, must_change=1 WHERE id=?", (generate_password_hash(pw), uid))
                flash("Password reset.", "ok")
        elif action == "unlock":
            x("UPDATE users SET failed_count=0, locked_until=NULL WHERE id=?", (uid,))
            flash(f"{target['username']} unlocked.", "ok")
        elif action == "toggle":
            x("UPDATE users SET active=1-active WHERE id=?", (uid,))
            flash("User updated.", "ok")
        elif action == "delete":
            to_bin("user", int(uid))
            x("DELETE FROM users WHERE id=?", (uid,))
            flash("User deleted.", "ok")
        return redirect(url_for("users"))
    rows = q("""SELECT u.*, e.name emp_name, e.emp_code FROM users u LEFT JOIN employees e ON e.id=u.employee_id
                ORDER BY CASE u.role WHEN 'admin' THEN 0 WHEN 'employee' THEN 2 ELSE 1 END, u.username""")
    groups = {}
    for k, module, label in PERMISSIONS:
        groups.setdefault(module, []).append((k, label))
    role_perms = {k: sorted(role_perm_set(k)) for k in roles}
    return render_template("users.html", rows=rows, roles=list(roles.values()), groups=groups, role_perms=role_perms)


@app.route("/roles", methods=["GET", "POST"])
@perm_required("users")
def roles_page():
    if request.method == "POST":
        action = request.form.get("action")
        if action == "add":
            name = request.form.get("name", "").strip()
            key = "".join(ch if ch.isalnum() else "_" for ch in name.lower()).strip("_")[:30]
            copy = request.form.get("copy_from", "")
            if not name or not key:
                flash("Role name is required.", "error")
            elif q("SELECT 1 FROM roles WHERE key=? OR name=?", (key, name), one=True):
                flash(f"A role called '{name}' already exists.", "error")
            else:
                perms = ""
                if copy:
                    src = q("SELECT perms FROM roles WHERE key=?", (copy,), one=True)
                    perms = ",".join(ALL_PERMS) if copy == "admin" else (src["perms"] if src else "")
                x("INSERT INTO roles(key,name,description,perms,system) VALUES(?,?,?,?,0)",
                  (key, name, request.form.get("description", "").strip(), perms))
                flash(f"Role '{name}' created. Tick its permissions below and save.", "ok")
        elif action == "save":
            for r in q("SELECT * FROM roles WHERE key<>'admin'"):
                perms = [p for p in request.form.getlist(f"p_{r['key']}") if p in ALL_PERMS]
                name = request.form.get(f"name_{r['key']}", r["name"]).strip() or r["name"]
                desc = request.form.get(f"desc_{r['key']}", r["description"] or "").strip()
                x("UPDATE roles SET perms=?, name=?, description=? WHERE key=?", (",".join(perms), name, desc, r["key"]))
            flash("Roles and permissions saved. Changes apply on each user's next page load.", "ok")
        elif action == "delete":
            key = request.form.get("key", "")
            r = q("SELECT * FROM roles WHERE key=?", (key,), one=True)
            n = q("SELECT COUNT(*) c FROM users WHERE role=?", (key,), one=True)["c"]
            if not r or r["system"]:
                flash("Built-in roles can't be deleted.", "error")
            elif n:
                flash(f"{n} user(s) still have this role. Change their role on the Users page first.", "error")
            else:
                to_bin("role", key)
                x("DELETE FROM roles WHERE key=?", (key,))
                flash(f"Role '{r['name']}' deleted.", "ok")
        return redirect(url_for("roles_page"))
    roles = q("""SELECT r.*, (SELECT COUNT(*) FROM users u WHERE u.role=r.key) n FROM roles r
                 ORDER BY CASE r.key WHEN 'admin' THEN 0 WHEN 'hr' THEN 1 WHEN 'employee' THEN 9 ELSE 5 END, r.name""")
    groups = {}
    for key, module, label in PERMISSIONS:
        groups.setdefault(module, []).append((key, label))
    role_perms = {r["key"]: (set(ALL_PERMS) if r["key"] == "admin" else set((r["perms"] or "").split(","))) for r in roles}
    return render_template("roles.html", roles=roles, groups=groups, role_perms=role_perms)


REQUIRED_TABLES = {"employees", "users", "settings", "departments"}


def check_backup_file(path):
    """Return None if the file is a usable NB Sons HR database, otherwise a reason it isn't."""
    try:
        with open(path, "rb") as f:
            if f.read(16) != b"SQLite format 3\x00":
                return "This is not a database (.db) file from the HR system."
        con = sqlite3.connect(f"file:{path}?mode=ro", uri=True)
        try:
            if con.execute("PRAGMA quick_check").fetchone()[0] != "ok":
                return "This database file is damaged."
            tables = {r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        finally:
            con.close()
        if not REQUIRED_TABLES <= tables:
            return "This database is not from the NB Sons HR system."
    except sqlite3.Error as e:
        return f"The file could not be read ({e})."
    return None


def backup_path(name):
    """Full path of an existing backup, or 404 (never allows paths outside the backups folder)."""
    if not _BK_SAFE.match(name or "") or not os.path.isfile(os.path.join(BACKUP_DIR, name)):
        abort(404)
    return os.path.join(BACKUP_DIR, name)


@app.route("/backups", methods=["GET", "POST"])
@perm_required("backup")
def backups():
    if request.method == "POST":
        act = request.form.get("act")
        if act == "create":
            try:
                name = make_backup("manual")
                audit("Created backup", name)
                flash(f"Backup saved: {name}", "ok")
            except Exception as e:                               # noqa: BLE001
                flash(f"Backup failed: {e}", "error")
        elif act == "upload":
            f = request.files.get("file")
            if not f or not f.filename:
                flash("Choose a .db backup file to upload.", "error")
                return redirect(url_for("backups"))
            if not f.filename.lower().endswith(".db"):
                flash("Only .db backup files can be uploaded.", "error")
                return redirect(url_for("backups"))
            os.makedirs(BACKUP_DIR, exist_ok=True)
            tmp = os.path.join(BACKUP_DIR, f".upload-{secrets.token_hex(6)}.tmp")
            f.save(tmp)
            problem = check_backup_file(tmp)
            if problem:
                os.remove(tmp)
                flash(problem, "error")
                return redirect(url_for("backups"))
            name = f"backup-{datetime.now():%Y-%m-%d-%H-%M-%S}-upload.db"
            os.replace(tmp, os.path.join(BACKUP_DIR, name))
            audit("Uploaded backup", f"{f.filename} → {name}")
            flash(f"Backup uploaded: {name}. Click Restore on it to load its data.", "ok")
        return redirect(url_for("backups"))
    items = list_backups()
    for b in items:
        b["label"] = BACKUP_KINDS[b["kind"]]
        secs = (datetime.now() - b["when"]).total_seconds()
        b["age"] = ("just now" if secs < 60 else f"{int(secs // 60)} min ago" if secs < 3600 else
                    f"{int(secs // 3600)} hr ago" if secs < 86400 else
                    "1 day ago" if secs < 172800 else f"{int(secs // 86400)} days ago")
    stats = {
        "total": len(items),
        "auto": sum(b["kind"] == "auto" for b in items),
        "manual": sum(b["kind"] != "auto" for b in items),
        "size": sum(b["size"] for b in items),
        "db_size": os.path.getsize(DB_PATH) if os.path.exists(DB_PATH) else 0,
    }
    return render_template("backups.html", items=items, stats=stats, now=datetime.now(),
                           error=_last_auto.get("error"), keep=AUTO_BACKUPS_KEEP)


@app.route("/backups/<name>/download")
@perm_required("backup")
def backup_download(name):
    path = backup_path(name)
    audit("Downloaded backup", name)
    return send_file(path, as_attachment=True, download_name=name)


@app.route("/backups/<name>/restore", methods=["POST"])
@perm_required("backup")
def backup_restore(name):
    path = backup_path(name)
    problem = check_backup_file(path)
    if problem:
        flash(f"Can't restore {name}: {problem}", "error")
        return redirect(url_for("backups"))
    who = g.user["username"]
    try:
        safety = make_backup("before-restore")          # so the restore itself can be undone
        con = g.pop("db", None)
        if con is not None:
            con.close()
        src = sqlite3.connect(f"file:{path}?mode=ro", uri=True)
        dest = sqlite3.connect(DB_PATH, timeout=30)
        src.backup(dest)
        dest.close()
        src.close()
        init_db()                                        # bring older backups up to date
    except Exception as e:                               # noqa: BLE001
        flash(f"Restore failed – nothing was changed: {e}", "error")
        return redirect(url_for("backups"))
    audit("Restored backup", f"{name} (previous data saved as {safety})", user=who)
    token = session.get("csrf")
    session.clear()
    session["csrf"] = token or secrets.token_hex(16)
    flash(f"Data restored from {name}. Please sign in again. Your data from before the restore was saved as {safety}.", "ok")
    return redirect(url_for("login"))


@app.route("/backups/<name>/delete", methods=["POST"])
@perm_required("backup")
def backup_delete(name):
    path = backup_path(name)
    os.remove(path)
    audit("Deleted backup", name)
    flash(f"Backup deleted: {name}", "ok")
    return redirect(url_for("backups"))


@app.route("/backup")
@perm_required("backup")
def backup():
    """Old link: take a backup now and download it."""
    name = make_backup("manual")
    return send_file(os.path.join(BACKUP_DIR, name), as_attachment=True, download_name=name)


# --------------------------------------------------------------------------- recycle bin
BIN_KINDS = {"employee": ("👤", "Employee"), "department": ("🏢", "Department"), "user": ("🔐", "User login"),
             "role": ("🛡️", "Role"), "asset": ("💻", "Asset"), "payroll": ("💰", "Salary row"),
             "loan": ("🏦", "Loan"), "leave_type": ("🌴", "Leave type")}


def _rows(sql, args=()):
    return [dict(r) for r in db().execute(sql, args).fetchall()]


def _ids(rows):
    return [r["id"] for r in rows]


def _inq(ids):
    return ",".join("?" * len(ids)) or "NULL"


def to_bin(kind, oid):
    """Save a full copy of an item (and the records that belong to it) before it is deleted."""
    tables, fix = [], []
    label = detail = ""
    if kind == "employee":
        e = _rows("SELECT * FROM employees WHERE id=?", (oid,))
        if not e:
            return
        label, detail = e[0]["name"], e[0]["emp_code"]
        pay = _rows("SELECT * FROM payroll WHERE employee_id=?", (oid,))
        loans = _rows("SELECT * FROM loans WHERE employee_id=?", (oid,))
        lp = _rows(f"SELECT * FROM loan_payments WHERE loan_id IN ({_inq(_ids(loans))}) OR payroll_id IN ({_inq(_ids(pay))})",
                   _ids(loans) + _ids(pay))
        tables = [("employees", e), ("users", _rows("SELECT * FROM users WHERE employee_id=?", (oid,))),
                  ("attendance", _rows("SELECT * FROM attendance WHERE employee_id=?", (oid,))),
                  ("leaves", _rows("SELECT * FROM leaves WHERE employee_id=?", (oid,))),
                  ("payroll", pay), ("loans", loans), ("loan_payments", lp),
                  ("increments", _rows("SELECT * FROM increments WHERE employee_id=?", (oid,))),
                  ("settlements", _rows("SELECT * FROM settlements WHERE employee_id=?", (oid,))),
                  ("emp_documents", _rows("SELECT * FROM emp_documents WHERE employee_id=?", (oid,))),
                  ("warnings", _rows("SELECT * FROM warnings WHERE employee_id=?", (oid,)))]
        fix = [("departments", "head_id", oid, _ids(_rows("SELECT id FROM departments WHERE head_id=?", (oid,)))),
               ("employees", "manager_id", oid, _ids(_rows("SELECT id FROM employees WHERE manager_id=?", (oid,)))),
               ("asset_assignments", "employee_id", oid, _ids(_rows("SELECT id FROM asset_assignments WHERE employee_id=?", (oid,))))]
        n = sum(len(r) for _, r in tables[2:])
        detail += f" · with {n} attendance/leave/salary/loan record(s)" if n else ""
    elif kind == "department":
        d = _rows("SELECT * FROM departments WHERE id=?", (oid,))
        if not d:
            return
        label, detail = d[0]["name"], d[0]["code"] or ""
        tables = [("departments", d)]
        fix = [("employees", "department_id", oid, _ids(_rows("SELECT id FROM employees WHERE department_id=?", (oid,))))]
    elif kind == "user":
        u = _rows("SELECT * FROM users WHERE id=?", (oid,))
        if not u:
            return
        label, detail = u[0]["username"], f"role: {u[0]['role']}"
        tables = [("users", u)]
    elif kind == "role":
        r = _rows("SELECT * FROM roles WHERE key=?", (oid,))
        if not r:
            return
        label, detail = r[0]["name"], r[0]["description"] or ""
        tables = [("roles", r)]
    elif kind == "asset":
        a = _rows("SELECT * FROM assets WHERE id=?", (oid,))
        if not a:
            return
        label, detail = f"{a[0]['asset_code']} · {a[0]['name']}", a[0]["category"] or ""
        tables = [("assets", a), ("asset_assignments", _rows("SELECT * FROM asset_assignments WHERE asset_id=?", (oid,)))]
    elif kind == "payroll":
        p_ = _rows("SELECT p.*, e.name ename FROM payroll p LEFT JOIN employees e ON e.id=p.employee_id WHERE p.id=?", (oid,))
        if not p_:
            return
        label, detail = f"{p_[0].pop('ename') or 'Employee'} – {p_[0]['month']}", f"net {p_[0]['net']:,.0f}"
        tables = [("payroll", p_), ("loan_payments", _rows("SELECT * FROM loan_payments WHERE payroll_id=?", (oid,)))]
    elif kind == "loan":
        ln = _rows("SELECT l.*, e.name ename FROM loans l LEFT JOIN employees e ON e.id=l.employee_id WHERE l.id=?", (oid,))
        if not ln:
            return
        label, detail = f"{ln[0].pop('ename') or 'Employee'} – loan", f"amount {ln[0]['amount']:,.0f}"
        tables = [("loans", ln), ("loan_payments", _rows("SELECT * FROM loan_payments WHERE loan_id=?", (oid,)))]
    elif kind == "leave_type":
        t = _rows("SELECT * FROM leave_types WHERE id=?", (oid,))
        if not t:
            return
        label, detail = t[0]["name"], f"{t[0]['days_per_year']} days/year"
        tables = [("leave_types", t)]
    import json
    db().execute("INSERT INTO recycle_bin(kind,label,detail,deleted_at,deleted_by,data) VALUES(?,?,?,?,?,?)",
                 (kind, label, detail, datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                  g.user["username"] if g.get("user") else "", json.dumps({"tables": tables, "fix": fix})))


def restore_from_bin(bid):
    """Put a deleted item back. Returns an error text, or None when it worked."""
    import json
    b = q("SELECT * FROM recycle_bin WHERE id=?", (bid,), one=True)
    if not b:
        return "This item is no longer in the bin."
    data = json.loads(b["data"])
    con = db()
    try:
        for table, rows in data["tables"]:
            have = {r[1] for r in con.execute(f"PRAGMA table_info({table})")}
            for r in rows:
                cols = [c for c in r if c in have]
                con.execute(f"INSERT INTO {table}({','.join(cols)}) VALUES({','.join('?' * len(cols))})",
                            [r[c] for c in cols])
        for table, col, val, ids in data.get("fix", []):
            if ids:
                con.execute(f"UPDATE {table} SET {col}=? WHERE {col} IS NULL AND id IN ({_inq(ids)})", [val] + ids)
        con.execute("DELETE FROM recycle_bin WHERE id=?", (bid,))
        con.commit()
    except sqlite3.IntegrityError as e:
        con.rollback()
        msg = str(e)
        if "UNIQUE" in msg:
            return ("Can't restore: something with the same code / username / name was added after this was deleted. "
                    "Rename or remove that one first, then restore.")
        return f"Can't restore: {msg}"
    return None


@app.route("/settings/bin", methods=["GET", "POST"])
@perm_required("settings")
def recycle_bin():
    if request.method == "POST":
        act, bid = request.form.get("act"), request.form.get("id", type=int)
        b = q("SELECT * FROM recycle_bin WHERE id=?", (bid,), one=True) if bid else None
        if act == "restore" and b:
            err = restore_from_bin(bid)
            if err:
                flash(err, "error")
            else:
                audit("Restored from bin", f"{BIN_KINDS.get(b['kind'], ('', b['kind']))[1]}: {b['label']}")
                flash(f"Restored: {b['label']}.", "ok")
        elif act == "purge" and b:
            x("DELETE FROM recycle_bin WHERE id=?", (bid,))
            audit("Deleted forever from bin", f"{b['kind']}: {b['label']}")
            flash(f"Deleted forever: {b['label']}.", "ok")
        elif act == "empty":
            x("DELETE FROM recycle_bin")
            audit("Emptied recycle bin", "")
            flash("Recycle bin emptied.", "ok")
        return redirect(url_for("recycle_bin", kind=request.args.get("kind", "")))
    kind = request.args.get("kind", "")
    sql, args = "SELECT id, kind, label, detail, deleted_at, deleted_by FROM recycle_bin", []
    if kind in BIN_KINDS:
        sql += " WHERE kind=?"
        args.append(kind)
    rows = q(sql + " ORDER BY id DESC", args)
    counts = {r["kind"]: r["c"] for r in q("SELECT kind, COUNT(*) c FROM recycle_bin GROUP BY kind")}
    return render_template("recycle_bin.html", rows=rows, counts=counts, kind=kind, kinds=BIN_KINDS)


# --------------------------------------------------------------------------- asset import
ASSET_IMPORT_COLS = [("asset_code", "Asset Code"), ("name", "Asset Name *"), ("category", "Category"), ("brand", "Brand"),
                     ("model", "Model"), ("serial_no", "Serial No"), ("purchase_date", "Purchase Date"),
                     ("cost", "Cost"), ("condition", "Condition"), ("location", "Location"), ("notes", "Notes"),
                     ("assign_to", "Assign To (Employee Code or Name)"), ("assigned_on", "Assigned On")]
_ASSET_ALIASES = {"code": "asset_code", "assetcode": "asset_code", "assetno": "asset_code", "tag": "asset_code",
                  "name": "name", "assetname": "name", "item": "name", "description": "name",
                  "category": "category", "type": "category", "brand": "brand", "make": "brand", "model": "model",
                  "serial": "serial_no", "serialno": "serial_no", "serialnumber": "serial_no", "imei": "serial_no",
                  "purchasedate": "purchase_date", "purchased": "purchase_date", "date": "purchase_date",
                  "cost": "cost", "price": "cost", "value": "cost", "condition": "condition", "location": "location",
                  "notes": "notes", "remarks": "notes", "assignto": "assign_to", "assignedto": "assign_to",
                  "employee": "assign_to", "employeecode": "assign_to", "user": "assign_to", "holder": "assign_to",
                  "assignedon": "assigned_on", "assigneddate": "assigned_on", "issuedon": "assigned_on"}


def _akey(h):
    k = "".join(ch for ch in str(h or "").lower() if ch.isalnum())
    k = k.replace("employeecodeorname", "")
    return _ASSET_ALIASES.get(k) or next((c for c, label in ASSET_IMPORT_COLS
                                          if "".join(ch for ch in label.lower() if ch.isalnum()) == k), None)


@app.route("/assets/import", methods=["GET", "POST"])
@perm_required("assets.manage")
def assets_import():
    result = None
    if request.method == "POST":
        fs = request.files.get("file")
        if not fs or not fs.filename:
            flash("Please choose an Excel or CSV file.", "error")
            return redirect(url_for("assets_import"))
        try:
            raw = read_upload(fs)
        except Exception as e:                                       # noqa: BLE001
            flash(f"Couldn't read the file: {e}", "error")
            return redirect(url_for("assets_import"))
        hdr_i, cols = None, {}
        for i, r in enumerate(raw[:10]):
            m = {j: _akey(v) for j, v in enumerate(r) if _akey(v)}
            if "name" in m.values():
                hdr_i, cols = i, m
                break
        if hdr_i is None:
            flash("Couldn't find the header row. The file needs at least an 'Asset Name' column – download the template.", "error")
            return redirect(url_for("assets_import"))
        match = _employee_matcher()
        added = updated = assigned = 0
        problems = []
        for n, r in enumerate(raw[hdr_i + 1:], start=hdr_i + 2):
            cell = {cols[j]: (v if not isinstance(v, str) else v.strip()) for j, v in enumerate(r) if j in cols}
            if not any(v not in (None, "") for v in cell.values()):
                continue
            name = _cell_text(cell.get("name"))
            if not name:
                problems.append((n, "Asset name is empty – row skipped"))
                continue
            code = _cell_text(cell.get("asset_code")).upper()
            cond = _cell_text(cell.get("condition")).title()
            cond = cond if cond in ASSET_CONDITIONS else "Good"
            pdate = _parse_any_date(cell.get("purchase_date"))[0] if cell.get("purchase_date") not in (None, "") else ""
            vals = {"name": name, "category": _cell_text(cell.get("category")) or "Other",
                    "brand": _cell_text(cell.get("brand")) or None, "model": _cell_text(cell.get("model")) or None,
                    "serial_no": _cell_text(cell.get("serial_no")) or None, "purchase_date": pdate or None,
                    "cost": to_float(cell.get("cost"), 0), "condition": cond,
                    "location": _cell_text(cell.get("location")) or None, "notes": _cell_text(cell.get("notes")) or None}
            ex = q("SELECT * FROM assets WHERE asset_code=?", (code,), one=True) if code else None
            if not ex and not code and vals["serial_no"]:          # same serial number = same asset, don't add twice
                ex = q("SELECT * FROM assets WHERE serial_no=?", (vals["serial_no"],), one=True)
                code = ex["asset_code"] if ex else code
            if ex:
                db().execute("""UPDATE assets SET name=?, category=?, brand=?, model=?, serial_no=?, purchase_date=?,
                                cost=?, condition=?, location=?, notes=? WHERE id=?""",
                             (*vals.values(), ex["id"]))
                aid = ex["id"]
                updated += 1
            else:
                code = code or next_asset_code()
                cur = db().execute("""INSERT INTO assets(asset_code,name,category,brand,model,serial_no,purchase_date,
                                      cost,condition,location,notes,status) VALUES(?,?,?,?,?,?,?,?,?,?,?,'available')""",
                                   (code, *vals.values()))
                aid = cur.lastrowid
                added += 1
            db().commit()
            who = _cell_text(cell.get("assign_to"))
            if who:
                emp, err = match(who, who, None)
                if not emp:
                    problems.append((n, f"{code}: saved, but not assigned – {err}"))
                    continue
                cur_holder = q("""SELECT employee_id FROM asset_assignments WHERE asset_id=? AND returned_on IS NULL""",
                               (aid,), one=True)
                if cur_holder and cur_holder["employee_id"] == emp["id"]:
                    continue
                if cur_holder:
                    problems.append((n, f"{code}: already assigned to someone else – return it first, then assign"))
                    continue
                on = _parse_any_date(cell.get("assigned_on"))[0] if cell.get("assigned_on") not in (None, "") else ""
                e2 = assign_asset(aid, emp["id"], on or None, None, cond, "Imported from Excel")
                if e2:
                    problems.append((n, f"{code}: {e2}"))
                else:
                    assigned += 1
        audit("Imported assets", f"{added} new, {updated} updated, {assigned} assigned from {fs.filename}")
        result = {"added": added, "updated": updated, "assigned": assigned, "problems": problems, "file": fs.filename}
    return render_template("assets_import.html", result=result, cols=ASSET_IMPORT_COLS)


@app.route("/assets/import/template")
@perm_required("assets.manage")
def assets_import_template():
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Assets"
    ws.append([label for _, label in ASSET_IMPORT_COLS])
    ws.append(["", "Dell Latitude 5440", "Laptop", "Dell", "Latitude 5440", "DL5440-001", "15/01/2026", 185000, "New",
               "Head office", "", "EMP-0001", "20/01/2026"])
    ws.append(["", "Samsung A15", "Mobile Phone", "Samsung", "A15", "IMEI 3567…", "01/02/2026", 45000, "Good",
               "", "Company phone", "", ""])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0F2A44")
    for i, (_, label) in enumerate(ASSET_IMPORT_COLS, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(14, len(label) + 3)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="asset-import-template.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# --------------------------------------------------------------------------- department detail
@app.route("/departments/<int:did>")
@perm_required("departments.view")
def department_view(did):
    d = q("""SELECT d.*, h.name head_name, h.emp_code head_code, h.id head_eid FROM departments d
             LEFT JOIN employees h ON h.id=d.head_id WHERE d.id=?""", (did,), one=True)
    if not d:
        abort(404)
    show = request.args.get("show", "active")
    term = request.args.get("q", "").strip()
    sql = """SELECT e.*, m.name mgr_name,
               (SELECT status FROM attendance a WHERE a.employee_id=e.id AND a.date=?) today
             FROM employees e LEFT JOIN employees m ON m.id=e.manager_id WHERE e.department_id=?"""
    args = [date.today().isoformat(), did]
    if show != "all":
        sql += " AND e.status='active'"
    if term:
        sql += " AND (e.name LIKE ? OR e.emp_code LIKE ? OR e.designation LIKE ?)"
        args += [f"%{term}%"] * 3
    emps = q(sql + " ORDER BY e.name", args)
    if request.args.get("export") == "csv":
        return csv_response(f"department-{d['name']}.csv",
                            ["Code", "Name", "Designation", "Phone", "Joined", "Reports to", "Status", "Today"],
                            [[e["emp_code"], e["name"], e["designation"] or "", e["phone"] or "", e["join_date"] or "",
                              e["mgr_name"] or "", e["status"], e["today"] or ""] for e in emps])
    active = [e for e in emps if e["status"] == "active"]
    stats = {"staff": q("SELECT COUNT(*) c FROM employees WHERE department_id=? AND status='active'", (did,), one=True)["c"],
             "in_today": sum(1 for e in active if e["today"] in ("present", "late", "half_day")),
             "absent": sum(1 for e in active if e["today"] == "absent"),
             "leave": sum(1 for e in active if e["today"] == "leave"),
             "bill": sum((e["basic_salary"] or 0) + (e["allowances"] or 0) for e in active)}
    others = q("SELECT id, name FROM departments ORDER BY name")
    return render_template("department_view.html", d=d, emps=emps, stats=stats, show=show, term=term, others=others)


# --------------------------------------------------------------------------- ZKT biometric link
_zkt_lock = threading.Lock()
_zkt_thread = {"on": False}


def _set(k, v):
    con = sqlite3.connect(DB_PATH, timeout=30)
    con.execute("INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (k, str(v)))
    con.commit()
    con.close()


def zkt_sync(force=False):
    """Read new punches from ZKT and update attendance. Returns a short result text."""
    import zkt_reader
    with _zkt_lock:
        con0 = sqlite3.connect(DB_PATH, timeout=30)
        s = dict(con0.execute("SELECT key, value FROM settings WHERE key LIKE 'zkt_%'").fetchall())
        con0.close()
        path = s.get("zkt_path") or ""
        try:
            mtime = str(os.path.getmtime(path))
            if not force and mtime == s.get("zkt_mtime"):
                _set("zkt_last", datetime.now().strftime("%Y-%m-%d %H:%M"))
                return "No new punches."
            people, punches = zkt_reader.read_mdb(path)
        except Exception as e:                                   # noqa: BLE001
            _set("zkt_err", f"Can't read ZKT data: {e}")
            return f"Can't read ZKT data: {e}"
        since = s.get("zkt_from") or date.today().replace(day=1).isoformat()
        match = _employee_matcher()
        emp_of = {}
        for b, n in people.items():
            m = match(b, n, None)[0]
            emp_of[b] = m["id"] if m else None
        con = db()
        new_days = set()
        for badge, t in punches:
            if t[:10] < since:
                continue
            if con.execute("INSERT OR IGNORE INTO zkt_punches(badge,punch_time) VALUES(?,?)", (badge, t)).rowcount:
                new_days.add((badge, t[:10]))
        updated = 0
        for badge, day in new_days:
            eid = emp_of.get(badge)
            if not eid:
                continue
            times = sorted(r[0][11:16] for r in con.execute(
                "SELECT punch_time FROM zkt_punches WHERE badge=? AND substr(punch_time,1,10)=?", (badge, day)))
            cin = times[0]
            last = times[-1]
            gap = (datetime.strptime(last, "%H:%M") - datetime.strptime(cin, "%H:%M")).seconds
            cout = last if gap >= 30 * 60 else None                   # ignore double presses at the door
            ex = con.execute("SELECT status FROM attendance WHERE employee_id=? AND date=?", (eid, day)).fetchone()
            if ex and ex[0] == "leave":
                continue                                              # don't overwrite approved leave
            con.execute("""INSERT INTO attendance(employee_id,date,status,check_in,check_out,note) VALUES(?,?,?,?,?,'ZKT')
                           ON CONFLICT(employee_id,date) DO UPDATE SET status=excluded.status,
                           check_in=excluded.check_in, check_out=excluded.check_out""",
                        (eid, day, _status_from_time(cin), cin, cout))
            updated += 1
        con.commit()
        _set("zkt_mtime", mtime)
        _set("zkt_err", "")
        _set("zkt_last", datetime.now().strftime("%Y-%m-%d %H:%M"))
        return f"{len(new_days)} new punch-day(s), {updated} attendance record(s) updated."


def _zkt_loop():
    while True:
        time.sleep(300)                                              # every 5 minutes
        try:
            with app.app_context():
                if (q("SELECT value FROM settings WHERE key='zkt_on'", one=True) or {"value": ""})["value"] == "1":
                    zkt_sync()
        except Exception as e:                                       # noqa: BLE001
            print("ZKT sync failed:", e)


def start_zkt_thread():
    if not _zkt_thread["on"]:
        _zkt_thread["on"] = True
        threading.Thread(target=_zkt_loop, daemon=True).start()


@app.route("/attendance/zkt", methods=["GET", "POST"])
@perm_required("attendance.manage")
def zkt_page():
    try:
        import zkt_reader
    except Exception as e:                                           # noqa: BLE001
        import traceback
        traceback.print_exc()
        return render_template("error.html", code="ZKT", msg=f"The ZKT reader could not start: {type(e).__name__}: {e}"), 500
    found = None
    if request.method == "POST":
        act = request.form.get("act")
        if act == "find":
            found = zkt_reader.find_mdb([d for d in os.environ.get("HR_ZKT_DIRS", "").split(os.pathsep) if d])
            if not found:
                flash("ZKT database not found on this PC. Install the HR server on the PC that has the ZKT software, "
                      "or type the path of att2000.mdb below.", "error")
        elif act == "connect":
            path = request.form.get("path", "").strip().strip('"')
            try:
                zkt_reader.read_mdb(path)
            except Exception as e:                                   # noqa: BLE001
                flash(f"Can't open this ZKT database: {e}", "error")
                return redirect(url_for("zkt_page"))
            for k, v in (("zkt_path", path), ("zkt_on", "1"), ("zkt_mtime", ""),
                         ("zkt_from", request.form.get("since") or date.today().replace(day=1).isoformat())):
                _set(k, v)
            audit("Connected ZKT", path)
            flash("ZKT connected. " + zkt_sync(force=True), "ok")
        elif act == "sync":
            flash(zkt_sync(force=True), "ok")
        elif act == "off":
            _set("zkt_on", "0")
            audit("Disconnected ZKT", "")
            flash("ZKT link turned off. Attendance already imported stays.", "ok")
        if act != "find" or not found:
            return redirect(url_for("zkt_page"))
    s = dict(get_settings())
    s.update({r["key"]: r["value"] for r in q("SELECT key, value FROM settings WHERE key LIKE 'zkt_%'")})
    people, unmatched = {}, []
    if s.get("zkt_on") == "1" and s.get("zkt_path"):
        try:
            people, _ = zkt_reader.read_mdb(s["zkt_path"])
        except Exception:                                            # noqa: BLE001
            people = {}
        match = _employee_matcher()
        unmatched = [(b, n) for b, n in sorted(people.items()) if not match(b, n, None)[0]]
    recent = q("""SELECT p.badge, p.punch_time FROM zkt_punches p ORDER BY p.punch_time DESC LIMIT 15""")
    return render_template("zkt.html", s=s, found=found, people=people, unmatched=unmatched, recent=recent,
                           punches=q("SELECT COUNT(*) c FROM zkt_punches", one=True)["c"],
                           month1=date.today().replace(day=1).isoformat())


# --------------------------------------------------------------------------- assets
ASSET_STATUSES = ["available", "assigned", "repair", "retired", "lost"]
ASSET_CONDITIONS = ["New", "Good", "Fair", "Damaged"]
ASSET_CATEGORIES = ["Laptop", "Desktop PC", "Monitor", "Mobile Phone", "SIM Card", "Tablet", "Printer", "Furniture",
                    "Vehicle", "Tools", "ID Card / Access Card", "Uniform", "Other"]


def next_asset_code():
    nums = [int(r["asset_code"][4:]) for r in q("SELECT asset_code FROM assets WHERE asset_code LIKE 'AST-%'")
            if r["asset_code"][4:].isdigit()]
    return f"AST-{(max(nums) + 1 if nums else 1):04d}"


def asset_categories():
    used = [r["category"] for r in q("SELECT DISTINCT category FROM assets WHERE category IS NOT NULL AND category<>''")]
    return sorted(set(ASSET_CATEGORIES) | set(used))


def get_asset(aid):
    a = q("""SELECT a.*, s.id asg_id, s.employee_id, s.assigned_on, s.expected_return, e.name emp_name, e.emp_code
             FROM assets a LEFT JOIN asset_assignments s ON s.asset_id=a.id AND s.returned_on IS NULL
             LEFT JOIN employees e ON e.id=s.employee_id WHERE a.id=?""", (aid,), one=True)
    if not a:
        abort(404)
    return a


ASSET_FIELDS = ["name", "category", "brand", "model", "serial_no", "purchase_date", "cost", "condition", "location", "notes"]


def _asset_form():
    d = {f: request.form.get(f, "").strip() for f in ASSET_FIELDS}
    d["cost"] = to_float(d["cost"])
    d["purchase_date"] = d["purchase_date"] if parse_date(d["purchase_date"]) else None
    d["condition"] = d["condition"] if d["condition"] in ASSET_CONDITIONS else "Good"
    return d


def assign_asset(aid, eid, on, expected, cond, note):
    emp = q("SELECT * FROM employees WHERE id=?", (eid,), one=True)
    a = q("SELECT * FROM assets WHERE id=?", (aid,), one=True)
    if not emp or not a:
        return "Choose an asset and an employee."
    if a["status"] != "available":
        return f"{a['asset_code']} is not available (status: {a['status']})."
    db().execute("""INSERT INTO asset_assignments(asset_id,employee_id,employee_name,assigned_on,expected_return,
                    condition_out,notes,assigned_by) VALUES(?,?,?,?,?,?,?,?)""",
                 (aid, eid, emp["name"], (parse_date(on) or date.today()).isoformat(),
                  expected if parse_date(expected) else None, cond or a["condition"], note or None, g.user["username"]))
    db().execute("UPDATE assets SET status='assigned' WHERE id=?", (aid,))
    db().commit()
    return None


def return_asset(aid, on, cond, note, new_status="available"):
    asg = q("SELECT * FROM asset_assignments WHERE asset_id=? AND returned_on IS NULL", (aid,), one=True)
    if not asg:
        return "This asset is not currently assigned."
    cond = cond if cond in ASSET_CONDITIONS else None
    db().execute("""UPDATE asset_assignments SET returned_on=?, condition_in=?, return_notes=?, received_by=?
                    WHERE id=?""", ((parse_date(on) or date.today()).isoformat(), cond, note or None,
                                     g.user["username"], asg["id"]))
    db().execute("UPDATE assets SET status=?, condition=COALESCE(?, condition) WHERE id=?",
                 (new_status if new_status in ASSET_STATUSES and new_status != "assigned" else "available", cond, aid))
    db().commit()
    return None


@app.route("/assets")
@perm_required("assets.view")
def assets():
    sql = """SELECT a.*, e.name emp_name, e.emp_code, e.id emp_id, s.assigned_on FROM assets a
             LEFT JOIN asset_assignments s ON s.asset_id=a.id AND s.returned_on IS NULL
             LEFT JOIN employees e ON e.id=s.employee_id WHERE 1=1"""
    args = []
    term = request.args.get("q", "").strip()
    if term:
        sql += """ AND (a.asset_code LIKE ? OR a.name LIKE ? OR a.serial_no LIKE ? OR a.brand LIKE ? OR a.model LIKE ?
                   OR a.category LIKE ? OR e.name LIKE ?)"""
        args += [f"%{term}%"] * 7
    cat, st = request.args.get("cat", ""), request.args.get("st", "")
    if cat:
        sql += " AND a.category=?"
        args.append(cat)
    if st in ASSET_STATUSES:
        sql += " AND a.status=?"
        args.append(st)
    rows = q(sql + " ORDER BY a.asset_code", args)
    if request.args.get("export") == "csv":
        head = ["Code", "Name", "Category", "Brand", "Model", "Serial", "Purchase date", "Cost", "Condition",
                "Status", "Assigned to", "Assigned on", "Location", "Notes"]
        body = [[r["asset_code"], r["name"], r["category"], r["brand"], r["model"], r["serial_no"], r["purchase_date"],
                 r["cost"], r["condition"], r["status"], r["emp_name"] or "", r["assigned_on"] or "", r["location"],
                 r["notes"]] for r in rows]
        return csv_response("assets.csv", head, body)
    stats = {r["status"]: r["c"] for r in q("SELECT status, COUNT(*) c FROM assets GROUP BY status")}
    value = q("SELECT COALESCE(SUM(cost),0) v FROM assets WHERE status<>'retired'", one=True)["v"]
    overdue = q("""SELECT COUNT(*) c FROM asset_assignments WHERE returned_on IS NULL AND expected_return IS NOT NULL
                   AND expected_return < ?""", (date.today().isoformat(),), one=True)["c"]
    return render_template("assets.html", rows=rows, stats=stats, value=value, overdue=overdue,
                           cats=asset_categories(), term=term, cat=cat, st=st, statuses=ASSET_STATUSES)


@app.route("/assets/new", methods=["GET", "POST"])
@perm_required("assets.manage")
def asset_new():
    if request.method == "POST":
        d = _asset_form()
        code = request.form.get("asset_code", "").strip().upper() or next_asset_code()
        if not d["name"]:
            flash("Asset name is required.", "error")
        elif q("SELECT 1 FROM assets WHERE asset_code=?", (code,), one=True):
            flash(f"Asset code {code} is already used.", "error")
        else:
            aid = x(f"INSERT INTO assets(asset_code,{','.join(ASSET_FIELDS)}) VALUES(?{',?' * len(ASSET_FIELDS)})",
                    [code] + [d[f] for f in ASSET_FIELDS])
            eid = request.form.get("assign_to", "")
            if eid.isdigit():
                err = assign_asset(aid, int(eid), date.today().isoformat(), request.form.get("expected_return"),
                                   d["condition"], None)
                flash(err or f"Asset {code} added and assigned.", "error" if err else "ok")
            else:
                flash(f"Asset {code} – {d['name']} added.", "ok")
            if request.form.get("again"):
                return redirect(url_for("asset_new"))
            return redirect(url_for("asset_view", aid=aid))
        return render_template("asset_form.html", a=request.form, code=code, new=True, cats=asset_categories(),
                               conds=ASSET_CONDITIONS, emps=_active_emps())
    return render_template("asset_form.html", a={"condition": "New", "purchase_date": date.today().isoformat()},
                           code=next_asset_code(), new=True, cats=asset_categories(), conds=ASSET_CONDITIONS,
                           emps=_active_emps())


@app.route("/assets/<int:aid>/edit", methods=["GET", "POST"])
@perm_required("assets.manage")
def asset_edit(aid):
    a = get_asset(aid)
    if request.method == "POST":
        d = _asset_form()
        code = request.form.get("asset_code", "").strip().upper() or a["asset_code"]
        st = request.form.get("status", a["status"])
        if a["status"] == "assigned":
            st = "assigned"                                   # change via Return
        elif st not in ASSET_STATUSES or st == "assigned":
            st = a["status"]
        if not d["name"]:
            flash("Asset name is required.", "error")
        elif q("SELECT 1 FROM assets WHERE asset_code=? AND id<>?", (code, aid), one=True):
            flash(f"Asset code {code} is already used.", "error")
        else:
            x(f"UPDATE assets SET asset_code=?, status=?, {', '.join(f + '=?' for f in ASSET_FIELDS)} WHERE id=?",
              [code, st] + [d[f] for f in ASSET_FIELDS] + [aid])
            flash("Asset updated.", "ok")
            return redirect(url_for("asset_view", aid=aid))
    return render_template("asset_form.html", a=a, code=a["asset_code"], new=False, cats=asset_categories(),
                           conds=ASSET_CONDITIONS, statuses=ASSET_STATUSES, emps=[])


@app.route("/assets/<int:aid>")
@perm_required("assets.view")
def asset_view(aid):
    a = get_asset(aid)
    hist = q("""SELECT s.*, COALESCE(e.name, s.employee_name) ename, e.emp_code FROM asset_assignments s
                LEFT JOIN employees e ON e.id=s.employee_id WHERE s.asset_id=? ORDER BY s.id DESC""", (aid,))
    return render_template("asset_view.html", a=a, hist=hist, emps=_active_emps(), conds=ASSET_CONDITIONS)


@app.post("/assets/<int:aid>/assign")
@perm_required("assets.manage")
def asset_assign(aid):
    eid = request.form.get("employee_id", "")
    err = assign_asset(aid, int(eid) if eid.isdigit() else 0, request.form.get("assigned_on"),
                       request.form.get("expected_return"), request.form.get("condition"), request.form.get("notes"))
    flash(err or "Asset assigned.", "error" if err else "ok")
    return redirect(request.form.get("next") or url_for("asset_view", aid=aid))


@app.post("/assets/<int:aid>/return")
@perm_required("assets.manage")
def asset_return(aid):
    err = return_asset(aid, request.form.get("returned_on"), request.form.get("condition"),
                       request.form.get("notes"), request.form.get("new_status", "available"))
    flash(err or "Asset returned.", "error" if err else "ok")
    return redirect(request.form.get("next") or url_for("asset_view", aid=aid))


@app.post("/assets/<int:aid>/delete")
@perm_required("assets.manage")
def asset_delete(aid):
    a = get_asset(aid)
    if a["status"] == "assigned":
        flash("Return this asset before deleting it.", "error")
        return redirect(url_for("asset_view", aid=aid))
    to_bin("asset", aid)
    x("DELETE FROM assets WHERE id=?", (aid,))
    flash(f"Asset {a['asset_code']} deleted.", "ok")
    return redirect(url_for("assets"))


@app.route("/assets/assign", methods=["GET", "POST"])
@perm_required("assets.manage")
def assets_assign():
    if request.method == "POST":
        ids = [int(i) for i in request.form.getlist("asset_id") if i.isdigit()]
        eid = request.form.get("employee_id", "")
        if not ids or not eid.isdigit():
            flash("Choose at least one asset and an employee.", "error")
        else:
            done, errs = 0, []
            for aid in ids:
                err = assign_asset(aid, int(eid), request.form.get("assigned_on"), request.form.get("expected_return"),
                                   request.form.get("condition"), request.form.get("notes"))
                if err:
                    errs.append(err)
                else:
                    done += 1
            if done:
                emp = q("SELECT name FROM employees WHERE id=?", (eid,), one=True)
                flash(f"{done} asset(s) assigned to {emp['name']}.", "ok")
            for e_ in errs:
                flash(e_, "error")
        return redirect(url_for("assets_assign"))
    available = q("SELECT * FROM assets WHERE status='available' ORDER BY category, name")
    out = q("""SELECT a.*, s.assigned_on, s.expected_return, s.condition_out, e.name emp_name, e.emp_code, e.id emp_id
               FROM asset_assignments s JOIN assets a ON a.id=s.asset_id LEFT JOIN employees e ON e.id=s.employee_id
               WHERE s.returned_on IS NULL ORDER BY e.name, a.asset_code""")
    return render_template("assets_assign.html", available=available, out=out, emps=_active_emps(),
                           conds=ASSET_CONDITIONS, pre_emp=request.args.get("emp", ""), pre_asset=request.args.get("asset", ""))


@app.route("/assets/history")
@perm_required("assets.view")
def assets_history():
    rows = q("""SELECT s.*, a.asset_code, a.name aname, a.category, COALESCE(e.name, s.employee_name) ename, e.emp_code,
                e.id emp_id FROM asset_assignments s JOIN assets a ON a.id=s.asset_id
                LEFT JOIN employees e ON e.id=s.employee_id ORDER BY s.id DESC LIMIT 1000""")
    return render_template("assets_history.html", rows=rows)


@app.route("/my/assets")
@login_required
def my_assets():
    emp_id = g.user["employee_id"] or abort(403)
    rows = q("""SELECT a.*, s.assigned_on, s.expected_return, s.returned_on, s.condition_out, s.condition_in
                FROM asset_assignments s JOIN assets a ON a.id=s.asset_id WHERE s.employee_id=?
                ORDER BY s.returned_on IS NOT NULL, s.id DESC""", (emp_id,))
    return render_template("my_assets.html", rows=rows)


# --------------------------------------------------------------------------- employee profiles
@app.route("/employees/profiles")
@perm_required("employees.view")
def employee_profiles():
    sql = """SELECT e.*, d.name dept,
               (SELECT COUNT(*) FROM asset_assignments s WHERE s.employee_id=e.id AND s.returned_on IS NULL) assets_n,
               (SELECT status FROM attendance a WHERE a.employee_id=e.id AND a.date=?) today_st
             FROM employees e LEFT JOIN departments d ON d.id=e.department_id WHERE 1=1"""
    args = [date.today().isoformat()]
    term = request.args.get("q", "").strip()
    if term:
        sql += " AND (e.name LIKE ? OR e.emp_code LIKE ? OR e.designation LIKE ? OR e.phone LIKE ? OR e.cnic LIKE ?)"
        args += [f"%{term}%"] * 5
    if request.args.get("dept"):
        sql += " AND e.department_id=?"
        args.append(request.args["dept"])
    status = request.args.get("status", "active")
    if status in ("active", "inactive"):
        sql += " AND e.status=?"
        args.append(status)
    rows = q(sql + " ORDER BY e.name", args)
    return render_template("employee_profiles.html", rows=rows, depts=all_depts(), status=status, term=term)


# --------------------------------------------------------------------------- global search
def global_search(term, limit=8):
    term = (term or "").strip()
    out = []
    if len(term) < 1:
        return out
    like = f"%{term}%"
    if can("employees.view"):
        rows = q("""SELECT e.id, e.name, e.emp_code, e.designation, e.status, d.name dept FROM employees e
                    LEFT JOIN departments d ON d.id=e.department_id
                    WHERE e.name LIKE ? OR e.emp_code LIKE ? OR e.cnic LIKE ? OR e.phone LIKE ? OR e.designation LIKE ?
                       OR e.email LIKE ? ORDER BY e.status, e.name LIMIT ?""", [like] * 6 + [limit])
        if rows:
            out.append({"group": "Employees", "icon": "👥", "more": url_for("employees", q=term, status="all"), "items": [
                {"title": r["name"], "sub": " · ".join(x_ for x_ in (r["emp_code"], r["designation"], r["dept"]) if x_),
                 "badge": r["status"], "url": url_for("employee_view", eid=r["id"])} for r in rows]})
    if can("assets.view"):
        rows = q("""SELECT a.id, a.asset_code, a.name, a.category, a.status, a.serial_no, e.name emp FROM assets a
                    LEFT JOIN asset_assignments s ON s.asset_id=a.id AND s.returned_on IS NULL
                    LEFT JOIN employees e ON e.id=s.employee_id
                    WHERE a.asset_code LIKE ? OR a.name LIKE ? OR a.serial_no LIKE ? OR a.category LIKE ?
                       OR a.brand LIKE ? OR a.model LIKE ? OR e.name LIKE ? ORDER BY a.asset_code LIMIT ?""",
                 [like] * 7 + [limit])
        if rows:
            out.append({"group": "Assets", "icon": "💻", "more": url_for("assets", q=term), "items": [
                {"title": f"{r['asset_code']} · {r['name']}",
                 "sub": " · ".join(x_ for x_ in (r["category"], ("with " + r["emp"]) if r["emp"] else None,
                                                 ("S/N " + r["serial_no"]) if r["serial_no"] else None) if x_),
                 "badge": r["status"], "url": url_for("asset_view", aid=r["id"])} for r in rows]})
    if can("attendance.view", "attendance.manage"):
        items = []
        d_, _ = _parse_any_date(term)
        if d_:
            c = q("SELECT status, COUNT(*) n FROM attendance WHERE date=? GROUP BY status", (d_,))
            summary = ", ".join(f"{r['n']} {r['status'].replace('_', ' ')}" for r in c) or "nothing marked yet"
            items.append({"title": f"Attendance on {nicedate(d_)}", "sub": summary, "badge": "",
                          "url": url_for("attendance", date=d_) if can("attendance.manage")
                          else url_for("attendance_report", month=d_[:7])})
        rows = q("""SELECT e.id, e.name, e.emp_code,
                      SUM(a.status IN ('present','late','half_day')) p, SUM(a.status='absent') ab, SUM(a.status='late') lt,
                      (SELECT status FROM attendance x WHERE x.employee_id=e.id AND x.date=?) today
                    FROM employees e LEFT JOIN attendance a ON a.employee_id=e.id AND substr(a.date,1,7)=?
                    WHERE e.name LIKE ? OR e.emp_code LIKE ? GROUP BY e.id ORDER BY e.name LIMIT ?""",
                 (date.today().isoformat(), date.today().strftime("%Y-%m"), like, like, limit))
        for r in rows:
            frm = date.today().replace(day=1).isoformat()
            items.append({"title": f"{r['name']} – attendance this month",
                          "sub": f"{r['p'] or 0} present · {r['lt'] or 0} late · {r['ab'] or 0} absent · today: "
                                 f"{(r['today'] or 'not marked').replace('_', ' ')}",
                          "badge": r["today"] or "",
                          "url": url_for("attendance_report", emp=r["id"])})
        if items:
            out.append({"group": "Attendance", "icon": "✅", "more": None, "items": items[:limit]})
    if can("leaves.view", "leaves.approve"):
        rows = q("""SELECT l.*, e.name, e.emp_code, t.name tname FROM leaves l JOIN employees e ON e.id=l.employee_id
                    JOIN leave_types t ON t.id=l.leave_type_id
                    WHERE e.name LIKE ? OR e.emp_code LIKE ? OR t.name LIKE ? OR l.reason LIKE ? OR l.status LIKE ?
                    ORDER BY l.start_date DESC LIMIT ?""", [like] * 5 + [limit])
        if rows:
            out.append({"group": "Leaves", "icon": "🌴", "more": url_for("leaves", status="all"), "items": [
                {"title": f"{r['name']} – {r['tname']} leave",
                 "sub": f"{nicedate(r['start_date'])}" + (f" to {nicedate(r['end_date'])}" if r["end_date"] != r["start_date"] else "")
                        + f" · {num(r['days'])} day(s)" + (f" · {r['reason']}" if r["reason"] else ""),
                 "badge": r["status"], "url": url_for("leaves", status=r["status"])} for r in rows]})
    if can("departments.view"):
        rows = q("SELECT id, name, code FROM departments WHERE name LIKE ? OR code LIKE ? LIMIT ?", (like, like, limit))
        if rows:
            out.append({"group": "Departments", "icon": "🏢", "more": url_for("departments"), "items": [
                {"title": r["name"], "sub": r["code"] or "", "badge": "", "url": url_for("employees", dept=r["id"], status="all")}
                for r in rows]})
    return out


@app.route("/api/search")
@login_required
def api_search():
    return {"q": request.args.get("q", ""), "groups": global_search(request.args.get("q", ""), 6)}


@app.route("/search")
@login_required
def search_page():
    term = request.args.get("q", "").strip()
    return render_template("search.html", term=term, groups=global_search(term, 50) if term else [])


# --------------------------------------------------------------------------- utils
def csv_response(name, head, body):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(head)
    w.writerows(body)
    return Response("﻿" + buf.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition": f"attachment; filename={name}"})


def lan_address():
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("10.255.255.255", 1))
        ip = s.getsockname()[0]
        s.close()
        return ip
    except OSError:
        try:
            return socket.gethostbyname(socket.gethostname())
        except OSError:
            return "127.0.0.1"


# --------------------------------------------------------------------------- notifications
NOTIFY_TYPES = [  # (type, label, permission needed to receive it, important?)
    ("summary", "Daily attendance summary (present / late / absent)", ("attendance.view", "attendance.manage"), False),
    ("leave_pending", "Leave waiting for approval", ("leaves.view", "leaves.approve"), False),
    ("leave_upcoming", "Someone goes on leave tomorrow", ("leaves.view", "leaves.approve"), False),
    ("payroll_due", "Payroll not generated yet (month end)", ("payroll.manage",), False),
    ("deduct_review", "Salary deductions waiting for HR review", ("payroll.manage",), True),
    ("loans_due", "Loan installments due this month", ("payroll.view",), False),
    ("asset_overdue", "Asset not returned on time", ("assets.view", "assets.manage"), False),
    ("birthday", "Birthdays & work anniversaries", ("employees.view",), False),
    ("doc_expiry", "Employee document expired / expiring within 30 days", ("employees.view",), False),
    ("holiday", "Public holiday tomorrow", ("attendance.view", "attendance.manage", "leaves.view"), False),
    ("security", "Security: account locked after wrong passwords", ("users",), True),
    ("backup", "Daily backup failed", ("users", "backup"), True),
]


def build_alerts():
    """Current alerts for the signed-in user. Each has a stable id so it pops up only once."""
    today = date.today()
    t = today.isoformat()
    off = set((get_settings().get("notify_off") or "").split(","))
    allowed = {k for k, _, perms, _ in NOTIFY_TYPES if k not in off and can(*perms)}
    out = []

    def add(kind, aid, title, body, url, important=False):
        if kind in allowed:
            out.append({"id": aid, "type": kind, "title": title, "body": body, "url": url, "important": important})

    if "summary" in allowed:
        by = {r["status"]: r["c"] for r in q("""SELECT a.status, COUNT(*) c FROM attendance a JOIN employees e
                                                 ON e.id=a.employee_id WHERE a.date=? AND e.status='active'
                                                 GROUP BY a.status""", (t,))}
        if by:
            absent = [r["name"] for r in q("""SELECT e.name FROM attendance a JOIN employees e ON e.id=a.employee_id
                                               WHERE a.date=? AND a.status='absent' ORDER BY e.name LIMIT 4""", (t,))]
            body = (f"{(by.get('present') or 0) + (by.get('half_day') or 0)} present · {by.get('late') or 0} late · "
                    f"{by.get('absent') or 0} absent · {by.get('leave') or 0} on leave")
            if absent:
                body += "\nAbsent: " + ", ".join(absent) + ("…" if (by.get("absent") or 0) > 4 else "")
            add("summary", f"sum-{t}", "Today's attendance", body, url_for("attendance"))
    for l in q("""SELECT l.id, e.name, t.name tn, l.start_date, l.days FROM leaves l JOIN employees e ON e.id=l.employee_id
                  JOIN leave_types t ON t.id=l.leave_type_id WHERE l.status='pending' ORDER BY l.id DESC LIMIT 10"""):
        add("leave_pending", f"lp-{l['id']}", "Leave waiting for approval",
            f"{l['name']} · {l['tn']} · {num(l['days'])} day(s) from {nicedate(l['start_date'])}", url_for("leaves"))
    tomorrow = (today + timedelta(days=1)).isoformat()
    for l in q("""SELECT l.id, e.name, t.name tn, l.end_date, l.days FROM leaves l JOIN employees e ON e.id=l.employee_id
                  JOIN leave_types t ON t.id=l.leave_type_id WHERE l.status='approved' AND l.start_date=?""", (tomorrow,)):
        add("leave_upcoming", f"lu-{l['id']}", "On leave from tomorrow",
            f"{l['name']} · {l['tn']} · {num(l['days'])} day(s), back after {nicedate(l['end_date'])}", url_for("leaves"))
    if "payroll_due" in allowed:
        active = q("SELECT COUNT(*) c FROM employees WHERE status='active'", one=True)["c"]
        if active:
            m = None
            if today.day >= 26:
                m = today.strftime("%Y-%m")
            elif today.day <= 5:
                m = (today.replace(day=1) - timedelta(days=1)).strftime("%Y-%m")
            if m and not q("SELECT 1 FROM payroll WHERE month=?", (m,), one=True):
                add("payroll_due", f"pd-{m}", "Payroll reminder",
                    f"Salaries for {monthname(m)} have not been generated yet.", url_for("payroll", month=m))
    if "deduct_review" in allowed:
        for r in q("""SELECT p.month, COUNT(*) c, SUM(p.suggested_deduction) s FROM payroll p WHERE p.ded_status='pending'
                      AND p.status<>'paid' GROUP BY p.month ORDER BY p.month DESC LIMIT 3"""):
            names = [x_["name"] for x_ in q("""SELECT e.name FROM payroll p JOIN employees e ON e.id=p.employee_id
                     WHERE p.month=? AND p.ded_status='pending' AND p.status<>'paid' ORDER BY e.name LIMIT 3""", (r["month"],))]
            add("deduct_review", f"dr-{r['month']}-{r['c']}-{int(r['s'] or 0)}", "💸 Salary deductions to review",
                f"{r['c']} salary(ies) for {monthname(r['month'])} have absences / lates – suggested {money(r['s'])} total. "
                f"Not deducted until you approve.\n{', '.join(names)}{'…' if r['c'] > 3 else ''}",
                url_for("payroll", month=r["month"]) + "#ded", True)
    if "loans_due" in allowed and today.day <= 5:
        n = q("""SELECT COUNT(*) c FROM loans l WHERE l.status='active' AND l.amount -
                 COALESCE((SELECT SUM(amount) FROM loan_payments p WHERE p.loan_id=l.id),0) > 0.5""", one=True)["c"]
        if n:
            add("loans_due", f"ln-{today:%Y-%m}", "Loan installments this month",
                f"{n} loan installment(s) will be deducted from {today:%B} salaries.", url_for("loans"))
    for a in q("""SELECT s.id, s.expected_return, a.name an, a.asset_code, e.name en FROM asset_assignments s
                  JOIN assets a ON a.id=s.asset_id LEFT JOIN employees e ON e.id=s.employee_id
                  WHERE s.returned_on IS NULL AND s.expected_return IS NOT NULL AND s.expected_return<>''
                  AND s.expected_return<? ORDER BY s.expected_return LIMIT 10""", (t,)):
        add("asset_overdue", f"ao-{a['id']}", "Asset not returned",
            f"{a['an']} ({a['asset_code']}) with {a['en'] or 'employee'} – was due {nicedate(a['expected_return'])}",
            url_for("assets_assign"))
    if "birthday" in allowed:
        md = today.strftime("%m-%d")
        for e in q("""SELECT id, name, dob, join_date FROM employees WHERE status='active'
                      AND (substr(dob,6,5)=? OR substr(join_date,6,5)=?)""", (md, md)):
            if (e["dob"] or "")[5:10] == md:
                add("birthday", f"bd-{e['id']}-{today.year}", "🎂 Birthday today", f"{e['name']} has a birthday today.",
                    url_for("employee_view", eid=e["id"]))
            jd = parse_date(e["join_date"])
            if jd and jd.strftime("%m-%d") == md and jd.year < today.year:
                yrs = today.year - jd.year
                add("birthday", f"wa-{e['id']}-{today.year}", "🎉 Work anniversary",
                    f"{e['name']} completes {yrs} year{'s' if yrs > 1 else ''} with the company today.",
                    url_for("employee_view", eid=e["id"]))
    if "doc_expiry" in allowed:
        for dc in q("""SELECT dc.id, dc.title, dc.expiry, dc.employee_id, e.name FROM emp_documents dc
                       JOIN employees e ON e.id=dc.employee_id WHERE e.status='active' AND dc.expiry IS NOT NULL
                       AND dc.expiry<>'' AND dc.expiry<=? ORDER BY dc.expiry LIMIT 10""",
                    ((today + timedelta(days=30)).isoformat(),)):
            gone = dc["expiry"] < t
            add("doc_expiry", f"dx-{dc['id']}-{dc['expiry']}", "📄 Document expired" if gone else "📄 Document expiring soon",
                f"{dc['name']} · {dc['title']} – {'expired' if gone else 'expires'} {nicedate(dc['expiry'])}",
                url_for("employee_view", eid=dc["employee_id"]) + "#docs")
    if "holiday" in allowed:
        hn = holiday_name(today + timedelta(days=1))
        if hn:
            add("holiday", f"hol-{tomorrow}", "🎉 Holiday tomorrow", f"{hn} – {nicedate(tomorrow)}. Office closed.",
                url_for("holidays"))
    if "security" in allowed:
        since = (datetime.now() - timedelta(hours=24)).strftime("%Y-%m-%d %H:%M:%S")
        for a in q("""SELECT id, at, username FROM audit_log WHERE action='Account locked' AND at>=?
                      ORDER BY id DESC LIMIT 5""", (since,)):
            add("security", f"sec-{a['id']}", "⚠️ Account locked",
                f"'{a['username']}' was locked after 5 wrong passwords ({a['at'][11:16]}).", url_for("users"), True)
    if "backup" in allowed and _last_auto.get("error"):
        add("backup", f"bk-{t}", "⚠️ Daily backup failed",
            "Today's automatic backup could not be saved: " + _last_auto["error"], url_for("activity"), True)
    return out


@app.route("/api/notifications")
def api_notifications():
    if not g.get("user") or g.get("idle_expired"):
        return app.response_class('{"error": "signed out"}', status=401, mimetype="application/json")
    s = get_settings()
    import json as _j
    resp = app.response_class(_j.dumps({"sound": s.get("notify_sound", "1") == "1",
                                        "loud": s.get("notify_loud", "1") == "1",
                                        "items": build_alerts()}), mimetype="application/json")
    resp.headers["Cache-Control"] = "no-store"
    return resp


@app.route("/api/ping")
def api_ping():
    """Lets the NB Sons HR client app find and check this server. Reveals nothing private."""
    resp = app.response_class(
        '{"app": "nbsons-hr", "version": "%s", "company": %s}' % (
            APP_VERSION, __import__("json").dumps(get_settings().get("company_name", ""))),
        mimetype="application/json")
    resp.headers["Access-Control-Allow-Origin"] = "*"
    resp.headers["Cache-Control"] = "no-store"
    return resp


# --------------------------------------------------------------------------- HR tools (holidays, increments,
# final settlement, documents, letters, warnings)
app.jinja_env.globals["holiday_today"] = lambda: holiday_name(date.today())

PK_FIXED_HOLIDAYS = [("02-05", "Kashmir Solidarity Day"), ("03-23", "Pakistan Day"), ("05-01", "Labour Day"),
                     ("05-28", "Youm-e-Takbeer"), ("08-14", "Independence Day"), ("11-09", "Iqbal Day"),
                     ("12-25", "Quaid-e-Azam Day")]


def _holiday_resync(dates):
    """After holidays change: clear 'absent' marks on new holidays and recount leave days that cover them."""
    g.pop("holidays", None)
    hol = holiday_map()
    cleared = 0
    for ds in dates:
        if ds in hol:
            cleared += db().execute("DELETE FROM attendance WHERE date=? AND status IN ('absent','leave')", (ds,)).rowcount
        for lv in q("""SELECT * FROM leaves WHERE status IN ('pending','approved') AND start_date<=? AND end_date>=?""",
                    (ds, ds)):
            days = sum(1 for d in daterange(parse_date(lv["start_date"]), parse_date(lv["end_date"])) if is_workday(d))
            db().execute("UPDATE leaves SET days=? WHERE id=?", (days, lv["id"]))
            if lv["status"] == "approved":
                apply_leave_attendance(q("SELECT * FROM leaves WHERE id=?", (lv["id"],), one=True))
    db().commit()
    return cleared


@app.route("/hr/holidays", methods=["GET", "POST"])
@perm_required("attendance.view", "attendance.manage", "leaves.view", "settings")
def holidays():
    year = request.args.get("year", type=int) or date.today().year
    if request.method == "POST":
        if not can("attendance.manage", "settings"):
            abort(403)
        act = request.form.get("act")
        changed = []
        if act == "add":
            a = parse_date(request.form.get("start"))
            b = parse_date(request.form.get("end")) or a
            name = request.form.get("name", "").strip()
            kind = request.form.get("kind") if request.form.get("kind") in ("public", "religious", "company") else "public"
            if not a or not name:
                flash("Please enter the holiday name and date.", "error")
            elif b < a or (b - a).days > 14:
                flash("Check the dates: the end date must be on/after the start, max 15 days.", "error")
            else:
                for d in daterange(a, b):
                    db().execute("""INSERT INTO holidays(date,name,kind) VALUES(?,?,?)
                                    ON CONFLICT(date) DO UPDATE SET name=excluded.name, kind=excluded.kind""",
                                 (d.isoformat(), name, kind))
                    changed.append(d.isoformat())
                n = _holiday_resync(changed)
                flash(f"Holiday saved: {name} ({len(changed)} day{'s' if len(changed) > 1 else ''})."
                      + (f" {n} absent/leave mark(s) on those days were cleared." if n else ""), "ok")
            year = a.year if a else year
        elif act == "pk":
            for md, name in PK_FIXED_HOLIDAYS:
                ds = f"{year}-{md}"
                if not q("SELECT 1 FROM holidays WHERE date=?", (ds,), one=True):
                    db().execute("INSERT INTO holidays(date,name,kind) VALUES(?,?,'public')", (ds, name))
                    changed.append(ds)
            n = _holiday_resync(changed)
            flash(f"Added {len(changed)} Pakistan national holiday(s) for {year}. Add Eid / Ashura / Eid Milad dates "
                  "yourself (they depend on the moon).", "ok")
        elif act == "delete":
            h = q("SELECT * FROM holidays WHERE id=?", (request.form.get("id", type=int),), one=True)
            if h:
                x("DELETE FROM holidays WHERE id=?", (h["id"],))
                _holiday_resync([h["date"]])
                flash(f"Removed holiday {h['name']} ({nicedate(h['date'])}). That day is a normal working day again.", "ok")
        return redirect(url_for("holidays", year=year))
    rows = q("SELECT * FROM holidays WHERE substr(date,1,4)=? ORDER BY date", (str(year),))
    t = date.today().isoformat()
    upcoming = q("SELECT * FROM holidays WHERE date>=? ORDER BY date LIMIT 1", (t,), one=True)
    months = {}
    for r in rows:
        months.setdefault(int(r["date"][5:7]), []).append(r)
    workdays = sum(1 for d in daterange(date(year, 1, 1), date(year, 12, 31)) if is_workday(d))
    return render_template("holidays.html", rows=rows, year=year, months=months, upcoming=upcoming,
                           workdays=workdays, month_names=calendar.month_name, today_s=t,
                           days_to=(parse_date(upcoming["date"]) - date.today()).days if upcoming else None)


# ---------- increments & salary history
INC_KINDS = {"increment": "Increment", "promotion": "Promotion", "decrement": "Salary decrease",
             "adjustment": "Salary adjustment"}


def log_salary_change(eid, old, basic, allow, desig, reason, commit=True):
    """Record any change of salary or designation in the employee's salary history."""
    ob, oa = old["basic_salary"] or 0, old["allowances"] or 0
    nb, na = to_float(basic, ob), to_float(allow, oa)
    od, nd = (old["designation"] or "").strip(), (desig or "").strip()
    if abs(nb - ob) < 0.5 and abs(na - oa) < 0.5 and od == nd:
        return
    kind = "promotion" if od != nd and nd else ("increment" if nb + na > ob + oa else
                                               ("decrement" if nb + na < ob + oa else "adjustment"))
    db().execute("""INSERT INTO increments(employee_id,eff_date,kind,old_basic,old_allow,new_basic,new_allow,old_desig,
                    new_desig,reason,created_by,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",
                 (eid, date.today().isoformat(), kind, ob, oa, nb, na, od, nd, reason,
                  g.user["username"] if g.get("user") else "", datetime.now().strftime("%Y-%m-%d %H:%M")))
    if commit:
        db().commit()


@app.route("/hr/increments", methods=["GET", "POST"])
@perm_required("payroll.view")
def increments():
    if request.method == "POST":
        if not can("payroll.manage"):
            abort(403)
        eid = request.form.get("employee_id", type=int)
        e = q("SELECT * FROM employees WHERE id=?", (eid,), one=True)
        eff = parse_date(request.form.get("eff_date"))
        kind = request.form.get("kind") if request.form.get("kind") in INC_KINDS else "increment"
        nb, na = to_float(request.form.get("new_basic"), -1), to_float(request.form.get("new_allow"), -1)
        nd = request.form.get("new_desig", "").strip()
        if not e or not eff or nb < 0 or na < 0:
            flash("Please choose the employee, the date and the new salary.", "error")
            return redirect(url_for("increments", emp=eid or ""))
        nd = nd or (e["designation"] or "")
        if abs(nb - (e["basic_salary"] or 0)) < 0.5 and abs(na - (e["allowances"] or 0)) < 0.5 and nd == (e["designation"] or ""):
            flash("Nothing changed – the new salary and designation are the same as now.", "warn")
            return redirect(url_for("increments", emp=eid))
        x("""INSERT INTO increments(employee_id,eff_date,kind,old_basic,old_allow,new_basic,new_allow,old_desig,new_desig,
             reason,created_by,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",
          (eid, eff.isoformat(), kind, e["basic_salary"] or 0, e["allowances"] or 0, nb, na, e["designation"] or "",
           nd, request.form.get("reason", "").strip(), g.user["username"], datetime.now().strftime("%Y-%m-%d %H:%M")))
        x("UPDATE employees SET basic_salary=?, allowances=?, designation=? WHERE id=?", (nb, na, nd or None, eid))
        old_g, new_g = (e["basic_salary"] or 0) + (e["allowances"] or 0), nb + na
        flash(f"{INC_KINDS[kind]} saved for {e['name']}: gross {money(old_g)} → {money(new_g)}"
              + (f" ({(new_g - old_g) * 100 / old_g:+.1f}%)" if old_g else "")
              + ". New salary is used from the next payroll you generate.", "ok")
        return redirect(url_for("increments", emp=eid) if request.form.get("back") != "profile"
                        else url_for("employee_view", eid=eid))
    emp = request.args.get("emp", type=int)
    sql = """SELECT i.*, e.name, e.emp_code, d.name dept FROM increments i JOIN employees e ON e.id=i.employee_id
             LEFT JOIN departments d ON d.id=e.department_id"""
    args = []
    if emp:
        sql += " WHERE i.employee_id=?"
        args.append(emp)
    rows = q(sql + " ORDER BY i.eff_date DESC, i.id DESC LIMIT 500", args)
    latest = {r["employee_id"]: r["id"] for r in q("SELECT employee_id, MAX(id) id FROM increments GROUP BY employee_id")}
    yr = str(date.today().year)
    year_rows = [r for r in q("SELECT * FROM increments WHERE substr(eff_date,1,4)=?", (yr,))]
    kp = {"count": len(year_rows), "promotions": sum(1 for r in year_rows if r["kind"] == "promotion"),
          "added": sum((r["new_basic"] + r["new_allow"]) - (r["old_basic"] + r["old_allow"]) for r in year_rows)}
    emps = q("SELECT id, name, emp_code, designation, basic_salary, allowances FROM employees WHERE status='active' ORDER BY name")
    return render_template("increments.html", rows=rows, emps=emps, emp=emp, kinds=INC_KINDS, latest=latest, kp=kp, yr=yr)


@app.post("/hr/increments/<int:iid>/undo")
@perm_required("payroll.manage")
def increment_undo(iid):
    r = q("SELECT * FROM increments WHERE id=?", (iid,), one=True) or abort(404)
    last = q("SELECT MAX(id) m FROM increments WHERE employee_id=?", (r["employee_id"],), one=True)["m"]
    if last != iid:
        flash("Only the latest change of an employee can be undone.", "error")
    else:
        x("UPDATE employees SET basic_salary=?, allowances=?, designation=? WHERE id=?",
          (r["old_basic"], r["old_allow"], r["old_desig"] or None, r["employee_id"]))
        x("DELETE FROM increments WHERE id=?", (iid,))
        flash("Undone – the employee is back on the previous salary and designation.", "ok")
    return redirect(request.form.get("next") or url_for("increments"))


# ---------- final settlement
EXIT_TYPES = ["Resignation", "Termination", "Contract ended", "Retirement", "Absconded", "Death", "Other"]


def settlement_calc(e, exit_d):
    """Work out the default final-settlement figures for an employee leaving on exit_d."""
    y, m = exit_d.year, exit_d.month
    dates = month_dates(y, m)
    workdays = [d for d in dates if is_workday(d)]
    gross = (e["basic_salary"] or 0) + (e["allowances"] or 0)
    per_day = gross / len(workdays) if workdays else 0
    jd = parse_date(e["join_date"])
    start = max(dates[0], jd) if jd and jd <= exit_d else dates[0]
    wd = [d for d in workdays if start <= d <= exit_d]
    c = Counter(r["status"] for r in q("SELECT status FROM attendance WHERE employee_id=? AND date BETWEEN ? AND ?",
                                       (e["id"], start.isoformat(), exit_d.isoformat())))
    unpaid = 0
    for lv in q("""SELECT l.start_date,l.end_date FROM leaves l JOIN leave_types t ON t.id=l.leave_type_id
                   WHERE l.employee_id=? AND l.status='approved' AND t.paid=0 AND l.start_date<=? AND l.end_date>=?""",
                (e["id"], exit_d.isoformat(), start.isoformat())):
        a, b = max(parse_date(lv["start_date"]), start), min(parse_date(lv["end_date"]), exit_d)
        unpaid += sum(1 for d in daterange(a, b) if is_workday(d))
    sal_days = max(len(wd) - c["absent"] - 0.5 * c["half_day"] - unpaid, 0)
    pay = q("SELECT id, net, status FROM payroll WHERE employee_id=? AND month=?", (e["id"], f"{y:04d}-{m:02d}"), one=True)
    rate = leave_day_rate(e)
    bal = [b for b in leave_balances(e["id"], exit_d.year) if b["encash"] and b["payable_days"] > 0]
    enc_days = sum(b["payable_days"] for b in bal)
    loans_ = [dict(r, balance=max((r["amount"] or 0) - (r["paid"] or 0), 0))
              for r in q(LOAN_SQL + " WHERE l.employee_id=? AND l.status<>'closed'", (e["id"],))]
    loans_ = [l for l in loans_ if l["balance"] > 0.5]
    assets_ = q("""SELECT s.id sid, a.id, a.asset_code, a.name, a.category, a.cost, s.assigned_on FROM asset_assignments s
                   JOIN assets a ON a.id=s.asset_id WHERE s.employee_id=? AND s.returned_on IS NULL ORDER BY a.name""", (e["id"],))
    return {"gross": gross, "per_day": per_day, "workdays_month": len(workdays), "worked_workdays": len(wd),
            "absent": c["absent"] + 0.5 * c["half_day"], "unpaid": unpaid, "salary_days": sal_days,
            "salary_amount": 0 if pay else round(per_day * sal_days), "payroll": pay, "start": start,
            "rate": rate, "enc_bal": bal, "encash_days": enc_days, "encash_amount": round(enc_days * rate),
            "loans": loans_, "loan_recovery": round(sum(l["balance"] for l in loans_)), "assets": assets_}


def _sett_net(f):
    return round(f["salary_amount"] + f["encash_amount"] + f["other_add"]
                 - f["loan_recovery"] - f["asset_deduction"] - f["other_ded"])


@app.route("/hr/settlements")
@perm_required("payroll.view")
def settlements():
    rows = q("""SELECT s.*, e.name, e.emp_code, e.designation, d.name dept FROM settlements s JOIN employees e
                ON e.id=s.employee_id LEFT JOIN departments d ON d.id=e.department_id ORDER BY s.id DESC""")
    emps = q("""SELECT id, name, emp_code, status FROM employees e WHERE NOT EXISTS
                (SELECT 1 FROM settlements s WHERE s.employee_id=e.id AND s.status='final') ORDER BY status, name""")
    kp = {"final": sum(1 for r in rows if r["status"] == "final"), "draft": sum(1 for r in rows if r["status"] == "draft"),
          "paid": sum(r["net"] or 0 for r in rows if r["status"] == "final")}
    return render_template("settlements.html", rows=rows, emps=emps, kp=kp)


@app.route("/hr/settlement/<int:eid>", methods=["GET", "POST"])
@perm_required("payroll.view")
def settlement_edit(eid):
    import json
    e = get_emp(eid)
    s = q("SELECT * FROM settlements WHERE employee_id=? ORDER BY id DESC LIMIT 1", (eid,), one=True)
    exit_d = parse_date(request.values.get("exit_date")) or (parse_date(s["exit_date"]) if s else None) or date.today()
    calc = settlement_calc(e, exit_d)
    if request.method == "POST":
        if not can("payroll.manage"):
            abort(403)
        if s and s["status"] == "final":
            flash("This settlement is already final.", "warn")
            return redirect(url_for("settlement_print", sid=s["id"]))
        f = {k: to_float(request.form.get(k), 0) for k in ("salary_days", "salary_amount", "encash_days", "encash_amount",
                                                          "loan_recovery", "other_add", "other_ded")}
        assets_ = []
        ded = 0
        for a in calc["assets"]:
            d_ = to_float(request.form.get(f"ad_{a['sid']}"), 0)
            ret = bool(request.form.get(f"ar_{a['sid']}"))
            ded += 0 if ret else d_
            assets_.append({"sid": a["sid"], "id": a["id"], "code": a["asset_code"], "name": a["name"],
                            "returned": ret, "deduction": 0 if ret else d_})
        f["asset_deduction"] = ded
        f["net"] = _sett_net(f)
        vals = (exit_d.isoformat(), request.form.get("exit_type", "Resignation"), request.form.get("reason", "").strip(),
                f["salary_days"], f["salary_amount"], f["encash_days"], f["encash_amount"], f["loan_recovery"],
                f["asset_deduction"], f["other_add"], request.form.get("other_add_note", "").strip(), f["other_ded"],
                request.form.get("other_ded_note", "").strip(), f["net"], json.dumps(assets_),
                request.form.get("remarks", "").strip())
        cols = ("exit_date,exit_type,reason,salary_days,salary_amount,encash_days,encash_amount,loan_recovery,"
                "asset_deduction,other_add,other_add_note,other_ded,other_ded_note,net,assets_json,remarks")
        if s:
            x(f"UPDATE settlements SET {', '.join(c + '=?' for c in cols.split(','))} WHERE id=?", vals + (s["id"],))
            sid = s["id"]
        else:
            sid = x(f"INSERT INTO settlements(employee_id,{cols},status,created_by,created_at) "
                    f"VALUES(?,{','.join('?' * 16)},'draft',?,?)",
                    (eid,) + vals + (g.user["username"], datetime.now().strftime("%Y-%m-%d %H:%M")))
        if request.form.get("act") == "final":
            today_s = date.today().isoformat()
            left = f["loan_recovery"]
            for ln in calc["loans"]:
                amt = min(left, ln["balance"])
                if amt > 0:
                    db().execute("""INSERT INTO loan_payments(loan_id,month,amount,paid_on,method,note,created_by)
                                    VALUES(?,?,?,?,'settlement','Recovered in final settlement',?)""",
                                 (ln["id"], exit_d.strftime("%Y-%m"), amt, today_s, g.user["username"]))
                    left -= amt
                if amt < ln["balance"] - 0.5:
                    db().execute("UPDATE loans SET status='closed', closed_on=?, notes=COALESCE(notes,'') || ? WHERE id=?",
                                 (today_s, " Closed in final settlement (balance waived).", ln["id"]))
            db().commit()
            for a in assets_:
                if a["returned"]:
                    return_asset(a["id"], exit_d.isoformat(), None, "Returned at final settlement")
            x("UPDATE employees SET status='inactive' WHERE id=?", (eid,))
            x("UPDATE users SET active=0 WHERE employee_id=?", (eid,))
            x("UPDATE settlements SET status='final', finalized_on=? WHERE id=?", (today_s, sid))
            flash(f"Final settlement done for {e['name']}. Net payable {money(f['net'])}. "
                  "The employee is now Inactive and their login is disabled.", "ok")
            return redirect(url_for("settlement_print", sid=sid))
        flash("Settlement saved as draft. Click Finalize when it's paid.", "ok")
        return redirect(url_for("settlement_edit", eid=eid))
    saved_assets = {}
    if s and s["assets_json"]:
        saved_assets = {a["sid"]: a for a in json.loads(s["assets_json"])}
    use_saved = bool(s) and not request.args.get("exit_date")
    return render_template("settlement_form.html", e=e, s=s, c=calc, exit_d=exit_d, use_saved=use_saved,
                           saved_assets=saved_assets, exit_types=EXIT_TYPES)


@app.route("/hr/settlement/<int:sid>/print")
@perm_required("payroll.view")
def settlement_print(sid):
    import json
    s = q("SELECT * FROM settlements WHERE id=?", (sid,), one=True) or abort(404)
    e = get_emp(s["employee_id"])
    assets_ = json.loads(s["assets_json"] or "[]")
    return render_template("settlement_print.html", s=s, e=e, assets=assets_)


@app.post("/hr/settlement/<int:sid>/delete")
@perm_required("payroll.manage")
def settlement_delete(sid):
    s = q("SELECT * FROM settlements WHERE id=?", (sid,), one=True) or abort(404)
    if s["status"] == "final" and request.form.get("reopen"):
        x("UPDATE settlements SET status='draft', finalized_on=NULL WHERE id=?", (sid,))
        flash("Settlement reopened as draft. (Employee status, loans and assets were not changed back.)", "warn")
        return redirect(url_for("settlement_edit", eid=s["employee_id"]))
    if s["status"] == "final":
        flash("A final settlement can't be deleted. Reopen it first.", "error")
    else:
        x("DELETE FROM settlements WHERE id=?", (sid,))
        flash("Draft settlement deleted.", "ok")
    return redirect(url_for("settlements"))


# ---------- employee documents
DOC_DIR = os.path.join(DATA_DIR, "documents")
DOC_TYPES = ["CNIC", "Contract", "CV / Resume", "Educational certificate", "Experience letter",
             "Offer / Appointment letter", "Medical certificate", "Police verification", "Passport / Visa",
             "Driving licence", "Photo", "Other"]
DOC_EXT = {".pdf", ".jpg", ".jpeg", ".png", ".gif", ".webp", ".doc", ".docx", ".xls", ".xlsx", ".txt", ".csv"}
DOC_INLINE = {".pdf", ".jpg", ".jpeg", ".png", ".gif", ".webp", ".txt"}


def doc_state(expiry):
    d = parse_date(expiry)
    if not d:
        return ""
    if d < date.today():
        return "expired"
    return "soon" if (d - date.today()).days <= 30 else "ok"


@app.route("/hr/documents")
@perm_required("employees.view")
def documents():
    show = request.args.get("show", "all")
    term = request.args.get("q", "").strip()
    sql = """SELECT dc.*, e.name, e.emp_code, e.status emp_status FROM emp_documents dc JOIN employees e ON e.id=dc.employee_id
             WHERE 1=1"""
    args = []
    if term:
        sql += " AND (e.name LIKE ? OR e.emp_code LIKE ? OR dc.title LIKE ? OR dc.doc_type LIKE ?)"
        args += [f"%{term}%"] * 4
    soon = (date.today() + timedelta(days=30)).isoformat()
    if show == "expiring":
        sql += " AND dc.expiry IS NOT NULL AND dc.expiry<>'' AND dc.expiry<=?"
        args.append(soon)
    rows = [dict(r, exp_state=doc_state(r["expiry"])) for r in q(sql + " ORDER BY dc.uploaded_at DESC LIMIT 1000", args)]
    allr = q("SELECT expiry, size FROM emp_documents")
    kp = {"total": len(allr), "size": sum(r["size"] or 0 for r in allr),
          "expired": sum(1 for r in allr if doc_state(r["expiry"]) == "expired"),
          "soon": sum(1 for r in allr if doc_state(r["expiry"]) == "soon")}
    emps = q("SELECT id, name, emp_code FROM employees ORDER BY status, name")
    return render_template("documents.html", rows=rows, kp=kp, show=show, term=term, emps=emps, types=DOC_TYPES,
                           pre=request.args.get("emp", type=int))


@app.post("/hr/documents/upload")
@perm_required("employees.manage")
def doc_upload():
    eid = request.form.get("employee_id", type=int)
    e = q("SELECT id, name FROM employees WHERE id=?", (eid,), one=True)
    files = [f for f in request.files.getlist("file") if f and f.filename]
    back = url_for("employee_view", eid=eid) + "#docs" if request.form.get("back") == "profile" else url_for("documents")
    if not e or not files:
        flash("Choose the employee and at least one file.", "error")
        return redirect(back)
    exp = parse_date(request.form.get("expiry"))
    dtype = request.form.get("doc_type") if request.form.get("doc_type") in DOC_TYPES else "Other"
    folder = os.path.join(DOC_DIR, str(eid))
    os.makedirs(folder, exist_ok=True)
    saved, bad = 0, []
    for f in files:
        ext = os.path.splitext(f.filename)[1].lower()
        if ext not in DOC_EXT:
            bad.append(f.filename)
            continue
        fname = f"{eid}/{secrets.token_hex(8)}{ext}"
        path = os.path.join(DOC_DIR, fname)
        f.save(path)
        size = os.path.getsize(path)
        title = request.form.get("title", "").strip() or os.path.splitext(f.filename)[0][:80]
        if len(files) > 1 and request.form.get("title", "").strip():
            title = f"{title} – {os.path.splitext(f.filename)[0][:40]}"
        x("""INSERT INTO emp_documents(employee_id,title,doc_type,filename,orig_name,size,expiry,notes,uploaded_by,uploaded_at)
             VALUES(?,?,?,?,?,?,?,?,?,?)""",
          (eid, title, dtype, fname, f.filename[:150], size, exp.isoformat() if exp else None,
           request.form.get("notes", "").strip(), g.user["username"], datetime.now().strftime("%Y-%m-%d %H:%M")))
        saved += 1
    if saved:
        flash(f"Uploaded {saved} document(s) for {e['name']}.", "ok")
    if bad:
        flash("Not uploaded (file type not allowed – use PDF, picture, Word or Excel): " + ", ".join(bad), "error")
    return redirect(back)


@app.route("/hr/documents/<int:did>/file")
@perm_required("employees.view")
def doc_file(did):
    dc = q("SELECT * FROM emp_documents WHERE id=?", (did,), one=True) or abort(404)
    path = os.path.realpath(os.path.join(DOC_DIR, dc["filename"]))
    if not path.startswith(os.path.realpath(DOC_DIR)) or not os.path.isfile(path):
        flash("The file for this document is missing from the server PC.", "error")
        return redirect(url_for("documents"))
    ext = os.path.splitext(path)[1].lower()
    inline = ext in DOC_INLINE and not request.args.get("download")
    return send_file(path, as_attachment=not inline, download_name=dc["orig_name"] or os.path.basename(path))


@app.post("/hr/documents/<int:did>/delete")
@perm_required("employees.manage")
def doc_delete(did):
    dc = q("SELECT * FROM emp_documents WHERE id=?", (did,), one=True) or abort(404)
    x("DELETE FROM emp_documents WHERE id=?", (did,))
    try:
        os.remove(os.path.join(DOC_DIR, dc["filename"]))
    except OSError:
        pass
    flash(f"Deleted document: {dc['title']}.", "ok")
    return redirect(request.form.get("next") or url_for("documents"))


@app.post("/hr/documents/<int:did>/edit")
@perm_required("employees.manage")
def doc_edit(did):
    dc = q("SELECT * FROM emp_documents WHERE id=?", (did,), one=True) or abort(404)
    exp = parse_date(request.form.get("expiry"))
    x("UPDATE emp_documents SET title=?, doc_type=?, expiry=? WHERE id=?",
      (request.form.get("title", "").strip() or dc["title"],
       request.form.get("doc_type") if request.form.get("doc_type") in DOC_TYPES else dc["doc_type"],
       exp.isoformat() if exp else None, did))
    flash("Document updated.", "ok")
    return redirect(request.form.get("next") or url_for("documents"))


# ---------- warnings & discipline
WARN_KINDS = {"verbal": "Verbal warning", "written": "Written warning", "final": "Final warning",
              "show_cause": "Show-cause notice"}


@app.route("/hr/warnings", methods=["GET", "POST"])
@perm_required("employees.view")
def warnings_page():
    if request.method == "POST":
        if not can("employees.manage"):
            abort(403)
        eid = request.form.get("employee_id", type=int)
        e = q("SELECT id, name FROM employees WHERE id=?", (eid,), one=True)
        on = parse_date(request.form.get("issued_on")) or date.today()
        subj = request.form.get("subject", "").strip()
        kind = request.form.get("kind") if request.form.get("kind") in WARN_KINDS else "written"
        if not e or not subj:
            flash("Choose the employee and write the subject.", "error")
        else:
            rb = parse_date(request.form.get("reply_by"))
            wid = x("""INSERT INTO warnings(employee_id,issued_on,kind,subject,details,action,reply_by,status,issued_by,created_at)
                       VALUES(?,?,?,?,?,?,?,'open',?,?)""",
                    (eid, on.isoformat(), kind, subj, request.form.get("details", "").strip(),
                     request.form.get("action", "").strip(), rb.isoformat() if rb else None, g.user["username"],
                     datetime.now().strftime("%Y-%m-%d %H:%M")))
            n = q("SELECT COUNT(*) c FROM warnings WHERE employee_id=?", (eid,), one=True)["c"]
            flash(f"{WARN_KINDS[kind]} recorded for {e['name']} (warning no. {n}). Print the letter from the list.", "ok")
            if request.form.get("print"):
                return redirect(url_for("warning_letter", wid=wid))
        return redirect(url_for("employee_view", eid=eid) + "#warns" if request.form.get("back") == "profile"
                        else url_for("warnings_page"))
    status = request.args.get("status", "all")
    sql = """SELECT w.*, e.name, e.emp_code, e.designation, d.name dept,
                    (SELECT COUNT(*) FROM warnings w2 WHERE w2.employee_id=w.employee_id) total
             FROM warnings w JOIN employees e ON e.id=w.employee_id LEFT JOIN departments d ON d.id=e.department_id"""
    args = []
    if status in ("open", "closed"):
        sql += " WHERE w.status=?"
        args.append(status)
    rows = q(sql + " ORDER BY w.issued_on DESC, w.id DESC LIMIT 1000", args)
    allw = q("SELECT kind, status, issued_on FROM warnings")
    yr = str(date.today().year)
    kp = {"open": sum(1 for r in allw if r["status"] == "open"), "year": sum(1 for r in allw if r["issued_on"][:4] == yr),
          "final": sum(1 for r in allw if r["kind"] == "final"),
          "show": sum(1 for r in allw if r["kind"] == "show_cause" and r["status"] == "open")}
    emps = q("SELECT id, name, emp_code FROM employees WHERE status='active' ORDER BY name")
    return render_template("warnings.html", rows=rows, kp=kp, status=status, emps=emps, kinds=WARN_KINDS,
                           pre=request.args.get("emp", type=int))


@app.post("/hr/warnings/<int:wid>")
@perm_required("employees.manage")
def warning_update(wid):
    w = q("SELECT * FROM warnings WHERE id=?", (wid,), one=True) or abort(404)
    act = request.form.get("act")
    if act == "delete":
        x("DELETE FROM warnings WHERE id=?", (wid,))
        flash("Warning deleted.", "ok")
    elif act == "close":
        x("UPDATE warnings SET status='closed', response=? WHERE id=?",
          (request.form.get("response", "").strip() or w["response"], wid))
        flash("Warning closed.", "ok")
    elif act == "reopen":
        x("UPDATE warnings SET status='open' WHERE id=?", (wid,))
        flash("Warning reopened.", "ok")
    return redirect(request.form.get("next") or url_for("warnings_page"))


# ---------- HR letters
LETTER_KINDS = {"offer": ("📨", "Offer letter", "Job offer with salary, before joining"),
                "appointment": ("📝", "Appointment letter", "Terms & conditions of employment"),
                "experience": ("🏅", "Experience letter", "Service certificate for a person who worked here"),
                "salary": ("💵", "Salary certificate", "For bank, loan, visa or rent"),
                "relieving": ("👋", "Relieving letter", "Confirms the person has been relieved of duties"),
                "employment": ("🪪", "Employment certificate", "Confirms the person currently works here")}
SALARY_LETTERS = {"offer", "appointment", "salary"}


@app.route("/hr/letters")
@perm_required("employees.manage")
def letters():
    emps = q("""SELECT e.id, e.name, e.emp_code, e.designation, e.status, d.name dept FROM employees e
                LEFT JOIN departments d ON d.id=e.department_id ORDER BY e.status, e.name""")
    return render_template("letters.html", emps=emps, kinds=LETTER_KINDS, pre=request.args.get("emp", type=int),
                           salary_kinds=SALARY_LETTERS)


def _letter_ref(kind, eid, on):
    return f"{(re.sub(r'[^A-Za-z]', '', get_settings().get('company_name') or '') or 'HR')[:3].upper()}/HR/{kind[:3].upper()}/{on.year}/{eid:04d}"


@app.route("/hr/letters/<int:eid>/<kind>")
@perm_required("employees.manage")
def letter(eid, kind):
    if kind not in LETTER_KINDS:
        abort(404)
    if kind in SALARY_LETTERS and not can("payroll.view"):
        abort(403)
    e = get_emp(eid)
    on = parse_date(request.args.get("date")) or date.today()
    sett = q("SELECT * FROM settlements WHERE employee_id=? ORDER BY id DESC LIMIT 1", (eid,), one=True)
    exit_d = parse_date(request.args.get("exit")) or (parse_date(sett["exit_date"]) if sett else None)
    lts = q("SELECT * FROM leave_types WHERE days_per_year>0 ORDER BY paid DESC, name")
    s = get_settings()
    off = ", ".join(WEEKDAYS[i] for i in sorted(weekly_off())) or "None"
    audit("Printed HR letter", f"{LETTER_KINDS[kind][1]} · {e['name']}")
    return render_template("letter.html", e=e, kind=kind, kinds=LETTER_KINDS, on=on, ref=_letter_ref(kind, eid, on),
                           exit_d=exit_d, lts=lts, off=off, s=s, salary_kinds=SALARY_LETTERS,
                           gross=(e["basic_salary"] or 0) + (e["allowances"] or 0), w=None)


@app.route("/hr/warnings/<int:wid>/letter")
@perm_required("employees.view")
def warning_letter(wid):
    w = q("SELECT * FROM warnings WHERE id=?", (wid,), one=True) or abort(404)
    e = get_emp(w["employee_id"])
    on = parse_date(w["issued_on"]) or date.today()
    n = q("SELECT COUNT(*) c FROM warnings WHERE employee_id=? AND (issued_on<? OR (issued_on=? AND id<=?))",
          (e["id"], w["issued_on"], w["issued_on"], wid), one=True)["c"]
    return render_template("letter.html", e=e, kind="warning", kinds=LETTER_KINDS, on=on,
                           ref=_letter_ref("warning", e["id"], on) + f"/{n}", exit_d=None, lts=[], off="",
                           s=get_settings(), salary_kinds=SALARY_LETTERS, gross=0, w=w, wkind=WARN_KINDS[w["kind"]], wno=n)


# ---------- asset management: who has what, give more assets to an existing employee
ASSET_GROUPS = [("Laptop", "💻", ("Laptop", "Desktop PC", "Monitor", "Tablet")), ("Mobile", "📱", ("Mobile Phone", "SIM Card")),
                ("Vehicle", "🚗", ("Vehicle", "Car", "Motorbike", "Bike")), ("Other", "📦", ())]


def _asset_group(cat):
    for name, _ic, cats in ASSET_GROUPS[:-1]:
        if (cat or "") in cats:
            return name
    return "Other"


@app.route("/assets/manage", methods=["GET", "POST"])
@perm_required("assets.view", "assets.manage")
def asset_management():
    if request.method == "POST":
        if not can("assets.manage"):
            abort(403)
        eid = request.form.get("employee_id", type=int)
        e = q("SELECT id, name FROM employees WHERE id=? AND status='active'", (eid,), one=True)
        back = url_for("asset_management", q=request.form.get("q", ""), grp=request.form.get("grp", "")) + f"#emp-{eid}"
        if not e:
            flash("Choose an existing active employee. New employees are added only from Add Employee.", "error")
            return redirect(back)
        act = request.form.get("act")
        if act == "return":
            aid = request.form.get("asset_id", type=int)
            err = return_asset(aid, date.today().isoformat(), request.form.get("condition"), "Returned (Asset Management)")
            flash(err or "Asset received back into store.", "error" if err else "ok")
            return redirect(back)
        ids = [int(i) for i in request.form.getlist("asset_id") if i.isdigit()]
        name = request.form.get("new_name", "").strip()
        if name:                                       # a brand-new asset given straight to the employee
            cat = request.form.get("new_cat", "").strip() or "Other"
            code = next_asset_code()
            ids.append(x("""INSERT INTO assets(asset_code,name,category,serial_no,cost,condition,status,purchase_date)
                            VALUES(?,?,?,?,?,?,'available',?)""",
                         (code, name, cat, request.form.get("new_serial", "").strip() or None,
                          to_float(request.form.get("new_cost"), 0), "New", date.today().isoformat())))
        if not ids:
            flash("Tick an asset from store, or type a new asset name.", "error")
            return redirect(back)
        done, errs = [], []
        for aid in ids:
            err = assign_asset(aid, eid, request.form.get("assigned_on") or date.today().isoformat(),
                               request.form.get("expected_return"), None, request.form.get("notes", "").strip() or None)
            a = q("SELECT asset_code, name FROM assets WHERE id=?", (aid,), one=True)
            (errs if err else done).append(f"{a['asset_code']} {a['name']}" + (f" ({err})" if err else ""))
        if done:
            flash(f"Given to {e['name']}: " + ", ".join(done), "ok")
        if errs:
            flash("Not assigned: " + "; ".join(errs), "error")
        return redirect(back)
    term = request.args.get("q", "").strip()
    grp = request.args.get("grp", "")
    held = q("""SELECT a.*, s.assigned_on, s.expected_return, s.employee_id FROM asset_assignments s JOIN assets a ON a.id=s.asset_id
                WHERE s.returned_on IS NULL ORDER BY a.category, a.name""")
    by_emp = {}
    for a in held:
        by_emp.setdefault(a["employee_id"], []).append(dict(a, grp=_asset_group(a["category"])))
    sql = """SELECT e.id, e.name, e.emp_code, e.designation, d.name dept FROM employees e LEFT JOIN departments d
             ON d.id=e.department_id WHERE e.status='active'"""
    args = []
    if term:
        sql += " AND (e.name LIKE ? OR e.emp_code LIKE ? OR e.designation LIKE ?)"
        args += [f"%{term}%"] * 3
    emps = []
    for e in q(sql + " ORDER BY e.name", args):
        mine = by_emp.get(e["id"], [])
        if grp and not any(a["grp"] == grp for a in mine):
            continue
        emps.append({"e": e, "assets": mine, "value": sum(a["cost"] or 0 for a in mine)})
    store = [dict(a, grp=_asset_group(a["category"])) for a in
             q("SELECT * FROM assets WHERE status='available' ORDER BY category, name")]
    allassets = q("SELECT category, status FROM assets")
    groups = []
    for name, ic, _c in ASSET_GROUPS:
        mine = [a for a in allassets if _asset_group(a["category"]) == name]
        groups.append({"name": name, "ic": ic, "total": len(mine), "given": sum(1 for a in mine if a["status"] == "assigned"),
                       "store": sum(1 for a in mine if a["status"] == "available")})
    return render_template("asset_management.html", emps=emps, store=store, groups=groups, term=term, grp=grp,
                           cats=asset_categories(), with_assets=sum(1 for x_ in emps if x_["assets"]))


init_db()

if __name__ == "__main__":
    port = int(os.environ.get("HR_PORT", "5000"))
    ip = lan_address()
    print("=" * 60)
    print("  HR System is running")
    print(f"  On this PC:            http://localhost:{port}")
    print(f"  Other PCs / phones:    http://{ip}:{port}")
    host = socket.gethostname()
    if host and host.lower() not in ("localhost",):
        print(f"  Other PCs (by name):   http://{host}:{port}   <- never changes")
    print("  IP keeps changing? Right-click FIX-IP.bat > Run as administrator (one time).")
    print("  First login: admin / admin123  (you'll be asked to change it)")
    print("  Keep this window open. Press Ctrl+C to stop.")
    print("=" * 60)
    try:
        from waitress import serve
        serve(app, host="0.0.0.0", port=port, threads=12)
    except ImportError:
        app.run(host="0.0.0.0", port=port)
